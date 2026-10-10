"""素材链接 Excel 模版与解析；只读取数据，不获取图片。"""
from io import BytesIO
from ipaddress import ip_address
import re
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 1000


def validate_image_link(value: object, row_number: int) -> str:
    if not isinstance(value, str):
        raise ValueError(f"第 {row_number} 行「SKU图」须为 HTTP(S) 图片链接")
    url = value.strip()
    try:
        parsed = urlsplit(url)
        valid = parsed.scheme in {"http", "https"} and bool(parsed.hostname)
        hostname = parsed.hostname or ""
        try:
            ip_address(hostname)
        except ValueError:
            hostname = hostname.rstrip(".").encode("idna").decode("ascii")
            valid = valid and len(hostname) <= 253 and all(
                re.fullmatch(r"[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?", label)
                for label in hostname.split(".")
            )
        # 读取端口以识别无效或超出范围的端口。
        parsed.port
        valid = valid and not parsed.username and not parsed.password
        valid = valid and not any(char.isspace() or ord(char) < 32 for char in url)
        valid = valid and not any(char in parsed.netloc for char in "\\<>")
    except (ValueError, UnicodeError):
        valid = False
    if not valid:
        raise ValueError(f"第 {row_number} 行「SKU图」须为带有效主机的 HTTP(S) 图片链接")
    if len(url) > 500:
        raise ValueError(f"第 {row_number} 行「SKU图」链接超过 500 字符")
    return url


def build_import_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "素材导入"
    sheet.append(["SKU图"])
    sheet["A1"].font = Font(bold=True)
    sheet.column_dimensions["A"].width = 80
    sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse_import_workbook(data: bytes) -> list[str]:
    if len(data) > MAX_IMPORT_BYTES:
        raise ValueError("Excel 文件不能超过 10MB")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("无法读取 XLSX 文件，请检查文件格式") from exc
    try:
        sheet = workbook.worksheets[0]
        # 不依赖文件中可能失真的工作表尺寸。
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, ())
        matches = [index for index, value in enumerate(headers) if isinstance(value, str) and value.strip() == "SKU图"]
        if not matches:
            raise ValueError("第一行缺少必需列：SKU图")
        if len(matches) != 1:
            raise ValueError("第一行只能有一个「SKU图」列")
        column = matches[0]
        links = {}
        nonempty = 0
        for row_number, row in enumerate(rows, start=2):
            value = row[column] if column < len(row) else None
            if value is None or isinstance(value, str) and not value.strip():
                continue
            nonempty += 1
            if nonempty > MAX_IMPORT_ROWS:
                raise ValueError(f"第 {row_number} 行超限：单次最多导入 1,000 条非空图片记录")
            links.setdefault(validate_image_link(value, row_number), None)
        if not links:
            raise ValueError("「SKU图」列没有可导入的图片链接")
        return list(links)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("无法读取 XLSX 工作表，请检查文件格式") from exc
    finally:
        workbook.close()
