from __future__ import annotations

from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook


DATA_START_ROW = 7
DATA_END_ROW = 1007

FIELD_CATEGORY = "ps_category"
FIELD_PRODUCT_NAME = "ps_product_name"
FIELD_PRODUCT_DESCRIPTION = "ps_product_description"
FIELD_PARENT_SKU = "ps_sku_parent_short"
FIELD_DANGEROUS_GOODS = "ps_dangerous_goods"
FIELD_VARIATION_INTEGRATION_NO = "et_title_variation_integration_no"
FIELD_VARIATION_NAME_1 = "et_title_variation_1"
FIELD_VARIATION_VALUE_1 = "et_title_option_for_variation_1"
FIELD_VARIATION_IMAGE = "et_title_image_per_variation"
FIELD_VARIATION_NAME_2 = "et_title_variation_2"
FIELD_VARIATION_VALUE_2 = "et_title_option_for_variation_2"
FIELD_PRICE = "ps_price"
FIELD_STOCK = "ps_stock"
FIELD_SKU = "ps_sku_short"
FIELD_SIZE_CHART_IMAGE = "et_title_size_chart"
FIELD_COVER_IMAGE = "ps_item_cover_image"
FIELD_WEIGHT = "ps_weight"
FIELD_LENGTH = "ps_length"
FIELD_WIDTH = "ps_width"
FIELD_HEIGHT = "ps_height"

REQUIRED_FIELDS = {
    FIELD_CATEGORY,
    FIELD_PRODUCT_NAME,
    FIELD_PRODUCT_DESCRIPTION,
    FIELD_VARIATION_INTEGRATION_NO,
    FIELD_VARIATION_NAME_1,
    FIELD_VARIATION_VALUE_1,
    FIELD_VARIATION_IMAGE,
    FIELD_VARIATION_NAME_2,
    FIELD_VARIATION_VALUE_2,
    FIELD_PRICE,
    FIELD_STOCK,
    FIELD_SKU,
    FIELD_COVER_IMAGE,
    FIELD_WEIGHT,
}


def _clean(value) -> str:
    return str(value or "").strip()


def _machine_field(value) -> str:
    """Shopee 会在机器字段后追加必填/类型标记，解析时只取稳定字段名。"""
    return _clean(value).split("|", 1)[0]


def _normalize_shopee_ooxml(template_bytes: bytes) -> bytes:
    """修正 Shopee 生成器写出的非标准 activePane 值，原文件保持不变。"""
    output = BytesIO()
    try:
        with ZipFile(BytesIO(template_bytes), "r") as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                    content = content.replace(b'activePane="bottom_left"', b'activePane="bottomLeft"')
                    content = content.replace(b'activePane="top_left"', b'activePane="topLeft"')
                    content = content.replace(b'activePane="bottom_right"', b'activePane="bottomRight"')
                    content = content.replace(b'activePane="top_right"', b'activePane="topRight"')
                target.writestr(info, content)
    except Exception as exc:
        raise ValueError("无法读取 Shopee XLSX 模板") from exc
    return output.getvalue()


def _field_columns(sheet) -> dict[str, int]:
    return {
        field: column
        for column in range(1, sheet.max_column + 1)
        if (field := _machine_field(sheet.cell(1, column).value))
    }


def _template_sheet_path(template_bytes: bytes) -> str:
    """从工作簿关系中定位 Template 工作表，避免依赖固定 sheet2.xml。"""
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with ZipFile(BytesIO(template_bytes), "r") as archive:
        workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relation_id = None
        for sheet in workbook.findall(f"{{{main_ns}}}sheets/{{{main_ns}}}sheet"):
            if sheet.attrib.get("name") == "Template":
                relation_id = sheet.attrib.get(f"{{{rel_ns}}}id")
                break
        if not relation_id:
            raise ValueError("Shopee 模板缺少 Template 工作表")
        relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        for relation in relationships.findall(f"{{{package_ns}}}Relationship"):
            if relation.attrib.get("Id") == relation_id:
                target = relation.attrib["Target"].lstrip("/")
                return target if target.startswith("xl/") else "xl/" + target
    raise ValueError("无法定位 Shopee Template 工作表")


def _load_workbook(template_bytes: bytes):
    try:
        return load_workbook(BytesIO(_normalize_shopee_ooxml(template_bytes)), data_only=False, read_only=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("无法读取 Shopee XLSX 模板") from exc


def parse_listing_options(template_bytes: bytes, template_type: str | None = None) -> dict:
    if template_type and template_type != "shopee_basic":
        raise ValueError("所选店铺类型与上传模板不匹配，识别结果为 Shopee")
    workbook = _load_workbook(template_bytes)
    required_sheets = {"Template", "Pre-order DTS Range"}
    missing_sheets = sorted(required_sheets - set(workbook.sheetnames))
    if missing_sheets:
        raise ValueError(f"Shopee 模板缺少工作表：{', '.join(missing_sheets)}")
    sheet = workbook["Template"]
    fields = _field_columns(sheet)
    missing_fields = sorted(REQUIRED_FIELDS - set(fields))
    if missing_fields:
        raise ValueError(f"Shopee 模板缺少导出字段：{', '.join(missing_fields)}")
    if _clean(sheet.cell(2, 1).value).lower() != "basic":
        raise ValueError("当前仅支持 Shopee basic 批量上传模板")

    category_sheet = workbook["Pre-order DTS Range"]
    categories = []
    seen_ids = set()
    for row in range(7, category_sheet.max_row + 1):
        category_id = _clean(category_sheet.cell(row, 2).value)
        if not category_id or category_id in seen_ids:
            continue
        raw_name = _clean(category_sheet.cell(row, 1).value)
        name = raw_name.split("-", 1)[1] if raw_name.startswith(f"{category_id}-") else raw_name
        categories.append({
            "id": category_id,
            "name": name or category_id,
            "pre_order_dts_range": _clean(category_sheet.cell(row, 3).value),
        })
        seen_ids.add(category_id)

    shipping_channels = [
        {
            "field": field,
            "id": field.split(".", 1)[1] if "." in field else field,
            "name": _clean(sheet.cell(3, column).value) or field,
        }
        for field, column in fields.items()
        if field.startswith("channel_id.")
    ]
    if not shipping_channels:
        raise ValueError("Shopee 模板没有可用物流渠道")

    return {
        "template_version": "basic",
        "template_type": "shopee_basic",
        "template_type_label": "Shopee",
        "capabilities": {
            "supports_cod": False,
            "supports_hubstudio_submit": False,
            "supports_category_attributes": False,
        },
        "categories": categories,
        "shipping_channels": shipping_channels,
        "price_limits": {"min": 0.10, "max": 1_000_000_000},
        "quantity_limits": {"min": 0, "max": 10_000_000},
        "title_limits": {"min": 20, "max": 120},
        "description_limits": {"min": 20, "max": 3000},
        "variation_limits": {"one_tier": 20, "two_tier": 50},
        "attributes_by_category": {},
        "base_requirements_by_category": {},
    }


def _preserve_template_package(generated_bytes: bytes, template_bytes: bytes) -> bytes:
    template_path = _template_sheet_path(template_bytes)
    output = BytesIO()
    with ZipFile(BytesIO(template_bytes), "r") as source, ZipFile(BytesIO(generated_bytes), "r") as generated, ZipFile(output, "w", ZIP_DEFLATED) as target:
        generated_template_xml = generated.read(template_path)
        generated_styles_xml = generated.read("xl/styles.xml")
        for info in source.infolist():
            if info.filename == template_path:
                content = generated_template_xml
            elif info.filename == "xl/styles.xml":
                content = generated_styles_xml
            else:
                content = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                content = content.replace(b'activePane="bottom_left"', b'activePane="bottomLeft"')
                content = content.replace(b'activePane="top_left"', b'activePane="topLeft"')
                content = content.replace(b'activePane="bottom_right"', b'activePane="bottomRight"')
                content = content.replace(b'activePane="top_right"', b'activePane="topRight"')
            target.writestr(info, content)
    return output.getvalue()


def build_workbook(
    *, template, category_id: str, shipping_channels: list[str], dangerous_goods: str,
    products: list[dict], template_bytes: bytes,
) -> bytes:
    options = parse_listing_options(template_bytes, "shopee_basic")
    category_ids = {item["id"] for item in options["categories"]}
    if category_id not in category_ids:
        raise ValueError("Shopee 类目不存在，请重新选择")
    valid_channel_fields = {item["field"] for item in options["shipping_channels"]}
    selected_channels = list(dict.fromkeys(shipping_channels))
    if not selected_channels or any(field not in valid_channel_fields for field in selected_channels):
        raise ValueError("请至少选择一个模板支持的 Shopee 物流渠道")

    workbook = _load_workbook(template_bytes)
    sheet = workbook["Template"]
    field_columns = _field_columns(sheet)
    for row in sheet.iter_rows(min_row=DATA_START_ROW, max_row=max(sheet.max_row, DATA_END_ROW), min_col=1, max_col=sheet.max_column):
        for cell in row:
            cell.value = None

    sizes = [
        _clean(size)
        for size in ((template.sku_specifications or {}).get("size", {}).get("options", []))
        if _clean(size)
    ] or ["Default"]
    row_number = DATA_START_ROW
    for product in products:
        sku_images = product["sku_images"]
        if len(sku_images) * len(sizes) > 50:
            raise ValueError(f"商品草稿 #{product['draft_id']} 的 Shopee 二级变体组合不能超过 50 个")
        gallery = list(dict.fromkeys(product["image_urls"]))[:9]
        parent_sku = f"HT-{product['draft_id']}"
        integration_no = f"HT-{product['draft_id']}"
        for sku_image in sku_images:
            base_sku = sku_image["sku"]
            if len(base_sku) > 20:
                raise ValueError(f"商品草稿 #{product['draft_id']} 的 Color 变体值不能超过 20 个字符：{base_sku}")
            for size in sizes:
                if row_number > DATA_END_ROW:
                    raise ValueError(f"导出数据超过 Shopee 模板 {DATA_END_ROW - DATA_START_ROW + 1} 行限制")
                seller_sku = base_sku if size == "Default" else f"{base_sku}-{size}"
                if len(size) > 20:
                    raise ValueError(f"商品草稿 #{product['draft_id']} 的 Size 变体值不能超过 20 个字符：{size}")
                if len(seller_sku) > 100:
                    raise ValueError(f"商品草稿 #{product['draft_id']} 的 SKU 不能超过 100 个字符")
                values = {
                    FIELD_CATEGORY: category_id,
                    FIELD_PRODUCT_NAME: product["title"],
                    FIELD_PRODUCT_DESCRIPTION: product["description"],
                    FIELD_PARENT_SKU: parent_sku,
                    FIELD_DANGEROUS_GOODS: dangerous_goods,
                    FIELD_VARIATION_INTEGRATION_NO: integration_no,
                    FIELD_VARIATION_NAME_1: "Color",
                    FIELD_VARIATION_VALUE_1: base_sku,
                    FIELD_VARIATION_IMAGE: sku_image["image_url"],
                    FIELD_VARIATION_NAME_2: None if size == "Default" else "Size",
                    FIELD_VARIATION_VALUE_2: None if size == "Default" else size,
                    FIELD_PRICE: product["price"],
                    FIELD_STOCK: product["quantity"],
                    FIELD_SKU: seller_sku,
                    FIELD_SIZE_CHART_IMAGE: product.get("size_chart_url"),
                    FIELD_COVER_IMAGE: gallery[0],
                    FIELD_WEIGHT: template.package_weight,
                    FIELD_LENGTH: template.package_length,
                    FIELD_WIDTH: template.package_width,
                    FIELD_HEIGHT: template.package_height,
                }
                for index, image_url in enumerate(gallery[1:9], start=1):
                    values[f"ps_item_image_{index}"] = image_url
                for channel_field in valid_channel_fields:
                    values[channel_field] = "On" if channel_field in selected_channels else "Off"
                for field, value in values.items():
                    column = field_columns.get(field)
                    if column:
                        sheet.cell(row_number, column).value = value
                row_number += 1

    output = BytesIO()
    workbook.save(output)
    return _preserve_template_package(output.getvalue(), template_bytes)
