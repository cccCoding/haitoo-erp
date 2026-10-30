"""导入只读取数据；不执行表格公式和标题中的指令。"""
import csv
import io
import re
import zipfile
from datetime import datetime

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from .title_models import TitleEvidence, TitleKeyword, TitleMarket, TitleImportWord
from .title_service import TAGS, ask_json, digest, normalize
from .title_dictionary import name_key, tag_entry

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 10000
TITLE_COLUMN_NAMES = {'原标题': '商品标题', '商品品类': '商品分类'}
KEYWORD_COLUMN_NAMES = {'马来语词语': '马来语词语', '中文释义': '中文释义', '适用品类': '产品分类', '属性分类': '标签'}
IMPORT_OPTIONAL_COLUMNS = ('商品ID', '链接')
IGNORED_METRIC_COLUMNS = {'销量', '搜索热度', '竞争度', '搜索曝光', '搜索点击', '搜索转化', '搜索来源', '搜索周期'}


def parse_file(content, filename, kind, *, allowed_tags=None):
    allowed_tags = TAGS if allowed_tags is None else allowed_tags
    if not content or len(content) > MAX_BYTES:
        raise HTTPException(400, "文件为空或超过 20MB")
    try:
        if filename.lower().endswith('.csv'):
            reader = csv.reader(io.StringIO(content.decode('utf-8-sig')))
            rows = list_row_limit(reader)
        elif filename.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 60 * 1024 * 1024:
                    raise ValueError("文件解压后过大")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                rows = list_row_limit(book.active.iter_rows(values_only=True))
            finally:
                book.close()
        else:
            raise ValueError("仅支持 UTF-8 CSV 或 XLSX 文件")
        if not rows:
            raise ValueError("文件没有数据")
        headers = [str(c or '').strip() for c in rows[0]]
        # 保持任务快照内的字段稳定，同时接受新列名和历史模板列名。
        column_names = TITLE_COLUMN_NAMES if kind == 'titles' else KEYWORD_COLUMN_NAMES
        aliases = {label: field for field, label in column_names.items()}
        headers = [aliases.get(header, header) for header in headers]
        required = list(column_names)
        missing = set(required) - set(headers)
        if missing:
            raise ValueError("缺少列：" + '、'.join(column_names[field] for field in required if field in missing))
        if len(set(headers)) != len(headers):
            raise ValueError("表头不能重复")
        parsed = []
        for row_number, row in enumerate(rows[1:], 2):
            values = {h: str(v).strip() if v is not None else '' for h, v in zip(headers, row) if h not in IGNORED_METRIC_COLUMNS}
            if not any(values.values()):
                continue
            if any(not values.get(k) for k in required):
                raise ValueError(f"第 {row_number} 行缺少必填信息")
            if any(len(v) > 2000 for v in values.values()):
                raise ValueError(f"第 {row_number} 行字段超过 2000 字符")
            if kind == 'keywords':
                if not split_values(values['适用品类']):
                    raise ValueError(f"第 {row_number} 行请至少填写一个产品分类")
                tag_names = {normalize(tag): tag for tag in allowed_tags}
                if re.search(r'[、;；]', values['属性分类']):
                    raise ValueError(f"第 {row_number} 行只能填写一个标签，请将多个标签拆分为多行")
                if values['属性分类'] == '面料':
                    values['属性分类'] = '材质'
                if normalize(values['属性分类']) not in tag_names:
                    raise ValueError(f"第 {row_number} 行标签不存在或已停用，请先在标签管理中维护")
                values['属性分类'] = tag_names[normalize(values['属性分类'])]
                if len(values['马来语词语']) > 160 or len(values['中文释义']) > 255:
                    raise ValueError(f"第 {row_number} 行词语或释义过长")
                if re.search(r'[\u3400-\u9fff]', values['马来语词语']):
                    raise ValueError(f"第 {row_number} 行马来语词语不能包含中文")
            if len(values.get('商品品类', values.get('适用品类', ''))) > 255:
                raise ValueError(f"第 {row_number} 行品类过长")
            parsed.append({'row': row_number, 'values': values, 'metrics': {}})
        if not parsed:
            raise ValueError('文件没有有效数据行')
        return parsed
    except HTTPException:
        raise
    except (ValueError, TypeError, UnicodeError, zipfile.BadZipFile, KeyError, OSError, csv.Error) as exc:
        raise HTTPException(400, f"导入文件无效：{exc}") from exc


def list_row_limit(reader):
    result = []
    for row in reader:
        if len(result) >= MAX_ROWS + 1:
            raise ValueError(f'单次最多导入 {MAX_ROWS} 行')
        result.append(row)
    return result


def split_values(value):
    return list(dict.fromkeys(v.strip() for v in re.split(r'[,，;；、]', value or '') if v.strip()))


def validate_word(word, row, kind, allowed_tags=None):
    allowed_tags = TAGS if allowed_tags is None else allowed_tags
    if not isinstance(word, dict):
        raise ValueError('提词结果须为对象')
    phrase = word.get('phrase', '')
    translation = word.get('translation', '')
    tag = word.get('tag', '待分类')
    if not isinstance(phrase, str) or not phrase.strip() or len(phrase) > 160 or re.search(r'[\u3400-\u9fff]', phrase):
        raise ValueError('词语无效')
    if not isinstance(translation, str) or not translation.strip() or len(translation) > 255 or tag not in allowed_tags:
        raise ValueError('释义或属性分类无效')
    original = row['values'].get('原标题', phrase)
    matched = re.search(re.escape(phrase), original, re.I)
    if not matched:
        raise ValueError('提取词语不在原标题中')
    return {'phrase': matched.group(), 'translation': translation.strip(), 'tag': tag,
        'categories': [row['values']['商品品类']] if kind == 'titles' else split_values(row['values']['适用品类']),
        'is_borrowed': word.get('is_borrowed') is True, 'position': matched.start(),
        'synonyms': [], 'conflicts': [], 'row': row['row']}


async def extract_row(row, kind, api_key=None, *, allowed_tags=None, tag_semantics=None):
    allowed_tags = TAGS if allowed_tags is None else allowed_tags
    if kind == 'keywords':
        v = row['values']
        word = validate_word({'phrase': v['马来语词语'], 'translation': v['中文释义'], 'tag': v['属性分类'],
                             'is_borrowed': v.get('是否借词') in ('是', 'Y', 'true')}, row, kind, allowed_tags)
        word['synonyms'] = split_values(v.get('同义表达'))
        return [word]
    answer = await ask_json(api_key,
        'Extract all independently meaningful original Malay keyword phrases from one MY clothing title. '
        'Return {"language":"ms or other", "keywords":[{"phrase":"EXACT contiguous original substring", "translation":"accurate Chinese meaning", '
        '"tag":"one allowed tag", "is_borrowed":false}]}. Keep complete phrases, reduplication and hyphens. '
        'Do not translate English into invented Malay. Mark English loan expressions as borrowed. '
        'Distinguish brands, shop names and promotional claims using 品牌 or 营销修饰. '
        'Set language to ms only for a Malay title (Malay with English loan phrases is allowed). '
        'Set other for English-only or other-language titles. '
        'Never infer attributes, search heat or competition from title. Do not extract grammar words.',
        {'title': row['values']['原标题'], 'category': row['values']['商品品类'], 'allowed_tags': allowed_tags,
         'tag_semantics': tag_semantics or {tag: tag for tag in allowed_tags}}, max_tokens=3000)
    words = answer.get('keywords')
    if answer.get('language') != 'ms':
        raise ValueError('原标题未识别为马来语，请核对语言；不得把英文标题翻译后当作原文提词')
    if not isinstance(words, list) or not words or len(words) > 60:
        raise ValueError('模型未返回有效提词列表')
    return [validate_word(w, row, kind, allowed_tags) for w in words]


def confirm_import(db, task, selections):
    market = db.scalar(select(TitleMarket).where(TitleMarket.id == task.market_id).with_for_update())
    result = task.result
    selected_ids = set(selections)
    words = db.scalars(select(TitleImportWord).where(TitleImportWord.task_id == task.id, TitleImportWord.id.in_(selected_ids))
        .with_for_update().execution_options(populate_existing=True)).all()
    if len(words) != len(selected_ids):
        raise HTTPException(400, '选择的提词结果不存在')
    confirmed_ids = set(result.get('selections', []))
    words = [word for word in words if word.id not in confirmed_ids]
    if not words:
        return result.get('confirmed', {'created': 0, 'merged': 0, 'evidence': 0})
    from .title_categories import materialize_categories
    categories_by_key = materialize_categories(db, task, {c for candidate in words for c in candidate.word['categories']})
    rows = {r['row']: r for r in task.snapshot['rows']}
    counts = {'created': 0, 'merged': 0, 'evidence': 0}
    for candidate in words:
        word = {**candidate.word, 'categories': list(dict.fromkeys(categories_by_key[name_key(c)] for c in candidate.word['categories']))}
        tag_entry(db, market, word['tag'])
        identity = digest([normalize(word['phrase']), word['tag']])
        keyword = db.scalar(select(TitleKeyword).where(TitleKeyword.market_id == market.id, TitleKeyword.identity_hash == identity)
            .with_for_update().execution_options(populate_existing=True))
        while keyword and keyword.merged_into_id:
            keyword = db.get(TitleKeyword, keyword.merged_into_id, with_for_update=True, populate_existing=True)
        if not keyword:
            keyword = TitleKeyword(market_id=market.id, code=f'KW{market.next_code:06d}', identity_hash=identity,
                phrase=word['phrase'], translation=word['translation'], tag=word['tag'], categories=word['categories'],
                synonyms=word['synonyms'], conflicts=[], is_borrowed=word['is_borrowed'],
                status='enabled', priority='normal')
            market.next_code += 1
            db.add(keyword)
            db.flush()
            counts['created'] += 1
        else:
            keyword.categories = sorted(set(keyword.categories) | set(word['categories']))
            counts['merged'] += 1
        keyword.status = 'enabled'
        keyword.updated_at = datetime.utcnow()
        candidate.word = {**candidate.word, 'imported_categories': word['categories']}
        row = rows[word['row']]
        v = row['values']
        source_product = v.get('商品ID') or v.get('链接') or digest([
            normalize(v.get('原标题', word['phrase'])), sorted(candidate.word['categories'])])
        observation = digest([normalize(task.snapshot['source_name']), source_product,
            task.snapshot['period'], identity])
        if not db.scalar(select(TitleEvidence.id).where(TitleEvidence.market_id == market.id,
                TitleEvidence.observation_hash == observation).with_for_update()):
            db.add(TitleEvidence(market_id=market.id, keyword_id=keyword.id, import_id=task.id,
                observation_hash=observation, original_title=v.get('原标题', word['phrase']), original_phrase=word['phrase'],
                source_name=task.snapshot['source_name'], source_product=source_product[:255], source_url=v.get('链接') or None,
                period=task.snapshot['period'], source_kind=task.snapshot['kind'], position=word['position'], metrics=row['metrics']))
            db.flush()
            counts['evidence'] += 1
    market.revision += 1
    market.updated_at = datetime.utcnow()
    counts = {key: result.get('confirmed', {}).get(key, 0) + value for key, value in counts.items()}
    task.result = {**result, 'confirmed': counts, 'selections': sorted(confirmed_ids | selected_ids)}
    task.state = 'completed'
    return counts
