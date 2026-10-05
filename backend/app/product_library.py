"""产品库历史订单表格解析。表格内容仅作为数据读取。"""

from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from urllib.parse import urlsplit

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


IMPORT_HEADERS = (
    "店铺名称", "站点", "平台", "订单编号", "数量", "下单时间", "标题", "平台SKU", "产品图片链接", "产品ID",
)
LOCAL_TIMEZONE = timezone(timedelta(hours=8))
MAX_IMPORT_ROWS = 50000
ALLOWED_SITES = ("泰国", "越南", "菲律宾", "马来西亚", "新加坡", "印度尼西亚")
ALLOWED_PLATFORMS = ("TikTok", "Shopee", "Temu")


@dataclass(frozen=True)
class ImportedOrderProduct:
    row_number: int
    shop_name: str
    site: str
    platform: str
    order_number: str
    ordered_at: datetime  # UTC naive，与数据库 DateTime 保持一致
    title: str
    sku: str
    image_url: str
    external_product_id: str
    platform_sku: str
    quantity: int


def _identifier(value: object, row_number: int, column: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        raise ValueError(f"第 {row_number} 行「{column}」无效")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            raise ValueError(f"第 {row_number} 行「{column}」须为文本或整数")
        result = str(int(value))
        if len(result) > 15:
            raise ValueError(f"第 {row_number} 行「{column}」超过 Excel 数值精度，请将该列设为文本后重新填写")
        return result
    return str(value).strip()


def _order_time(value: object, row_number: int, epoch: datetime) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = from_excel(value, epoch)
        except (ValueError, OverflowError) as exc:
            raise ValueError(f"第 {row_number} 行「下单时间」无效") from exc
        if not isinstance(parsed, datetime):
            raise ValueError(f"第 {row_number} 行「下单时间」无效")
    elif isinstance(value, str) and value.strip():
        raw = value.strip().replace("年", "-").replace("月", "-").replace("日", " ").replace("/", "-")
        try:
            parsed = datetime.fromisoformat(raw)
        except ValueError as exc:
            raise ValueError(f"第 {row_number} 行「下单时间」格式无效，请使用日期或 YYYY-MM-DD HH:MM:SS") from exc
    else:
        raise ValueError(f"第 {row_number} 行缺少「下单时间」")
    if not 2000 <= parsed.year <= 2100:
        raise ValueError(f"第 {row_number} 行「下单时间」超出支持范围")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=LOCAL_TIMEZONE)
    return parsed.astimezone(timezone.utc).replace(tzinfo=None, microsecond=0)


def parse_order_workbook(data: bytes) -> list[ImportedOrderProduct]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 XLSX 文件，请检查文件格式") from exc
    try:
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers = {str(cell.value).strip(): index for index, cell in enumerate(header_row) if cell.value is not None}
        missing = [name for name in IMPORT_HEADERS if name not in headers]
        if missing:
            raise ValueError(f"缺少必需列：{'、'.join(missing)}")
        inherited: dict[str, str] = {}
        inherited_time: datetime | None = None
        items: list[ImportedOrderProduct] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = {name: row[index] if index < len(row) else None for name, index in headers.items() if name in IMPORT_HEADERS}
            if not any(value is not None and str(value).strip() for value in values.values()):
                continue
            if len(items) >= MAX_IMPORT_ROWS:
                raise ValueError(f"单次最多导入 {MAX_IMPORT_ROWS} 行数据")
            sku_value = _identifier(values["平台SKU"], row_number, "平台SKU")
            product_id = _identifier(values["产品ID"], row_number, "产品ID")
            if not sku_value or not product_id:
                raise ValueError(f"第 {row_number} 行缺少「{'平台SKU' if not sku_value else '产品ID'}」")
            for name in ("店铺名称", "站点", "平台", "订单编号"):
                current = _identifier(values[name], row_number, name)
                if current:
                    inherited[name] = current
                elif name not in inherited:
                    raise ValueError(f"第 {row_number} 行缺少「{name}」")
            for name, allowed in (("站点", ALLOWED_SITES), ("平台", ALLOWED_PLATFORMS)):
                if inherited[name] not in allowed:
                    raise ValueError(f"第 {row_number} 行「{name}」值「{inherited[name]}」无效，可填值：{'、'.join(allowed)}")
            raw_time = values["下单时间"]
            if raw_time is not None and str(raw_time).strip():
                inherited_time = _order_time(raw_time, row_number, workbook.epoch)
            elif values["订单编号"] is not None and str(values["订单编号"]).strip():
                raise ValueError(f"第 {row_number} 行缺少「下单时间」")
            if inherited_time is None:
                raise ValueError(f"第 {row_number} 行缺少「下单时间」")
            raw_quantity = values["数量"]
            try:
                if isinstance(raw_quantity, bool) or raw_quantity is None:
                    raise ValueError
                quantity = int(raw_quantity)
                if isinstance(raw_quantity, float) and not raw_quantity.is_integer():
                    raise ValueError
                if isinstance(raw_quantity, str) and not raw_quantity.strip().isdecimal():
                    raise ValueError
                if quantity <= 0 or quantity > 2147483647:
                    raise ValueError
            except (ValueError, TypeError, OverflowError) as exc:
                raise ValueError(f"第 {row_number} 行「数量」须为正整数（不超过 2147483647）") from exc
            if len(sku_value) > 120:
                raise ValueError(f"第 {row_number} 行「平台SKU」不能超过 120 个字符")
            sku = sku_value.split("-", 1)[0].strip()
            if not sku:
                raise ValueError(f"第 {row_number} 行「平台SKU」去除尺码后为空")
            title = _identifier(values["标题"], row_number, "标题")
            image_url = _identifier(values["产品图片链接"], row_number, "产品图片链接")
            for name, value in (("标题", title), ("产品图片链接", image_url)):
                if not value:
                    raise ValueError(f"第 {row_number} 行缺少「{name}」，请填写该字段")
            image_parts = urlsplit(image_url)
            if image_url and (image_parts.scheme.lower() not in {"http", "https"} or not image_parts.netloc):
                raise ValueError(f"第 {row_number} 行「产品图片链接」须为 HTTP(S) 地址")
            limits = {"店铺名称": 160, "站点": 80, "平台": 80, "订单编号": 120, "产品ID": 120, "平台SKU": 120, "标题": 500}
            fields = {**inherited, "产品ID": product_id, "平台SKU": sku, "标题": title}
            for name, limit in limits.items():
                if len(fields[name]) > limit:
                    raise ValueError(f"第 {row_number} 行「{name}」不能超过 {limit} 个字符")
            items.append(ImportedOrderProduct(
                row_number=row_number, shop_name=inherited["店铺名称"], site=inherited["站点"],
                platform=inherited["平台"], order_number=inherited["订单编号"], ordered_at=inherited_time,
                title=title, sku=sku, image_url=image_url, external_product_id=product_id,
                platform_sku=sku_value, quantity=quantity,
            ))
        if not items:
            raise ValueError("表格没有可导入的数据行")
        return items
    finally:
        workbook.close()
