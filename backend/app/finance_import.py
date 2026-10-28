"""财务 XLSX 解析；只读表格数据，不执行公式或其他文档内容。"""
from datetime import date, datetime, timezone, timedelta
from decimal import Decimal, InvalidOperation, localcontext
from io import BytesIO
import re
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50000
LOCAL_TIMEZONE = timezone(timedelta(hours=8))
BILL_HEADERS = ("交易类型", "订单ID/调整单ID", "货币", "结算总金额", "总收入")
CROSS_BORDER_BILL_HEADERS = (*BILL_HEADERS, "结算单ID", "SKUID")
ERP_HEADERS = ("订单编号", "平台SKU", "商品数量", "下单时间")
HEADER_ALIASES = {
    "订单ID/调整单ID": ("订单ID/调整单ID", "订单ID", "订单编号", "订单号"),
    "订单编号": ("订单编号", "订单ID", "订单号"),
}
MAX_HEADER_ROWS = 50


def normalize_header(value):
    return re.sub(r'[\s\ufeff\u200b]+', '', str(value or '')).casefold()


def find_headers(workbook, required, sheet_name, optional=()):
    matches, best = [], None
    if sheet_name and sheet_name not in workbook.sheetnames:
        raise ValueError('所选工作表不存在')
    for sheet in workbook:
        if sheet_name and sheet.title != sheet_name:
            continue
        sheet.reset_dimensions()
        for number, header in enumerate(sheet.iter_rows(max_row=MAX_HEADER_ROWS, values_only=True), 1):
            normalized = [normalize_header(v) for v in header]
            found = {name: [i for i, value in enumerate(normalized)
                if value in {normalize_header(alias) for alias in HEADER_ALIASES.get(name, (name,))}]
                for name in required}
            missing = [name for name, indexes in found.items() if not indexes]
            if best is None or len(missing) < len(best[2]):
                best = (sheet.title, number, missing)
            if missing:
                continue
            duplicates = [name for name, indexes in found.items() if len(indexes) > 1]
            if duplicates:
                raise ValueError(f"工作表「{sheet.title}」第 {number} 行存在重复必需列：{'、'.join(duplicates)}")
            columns = {name: indexes[0] for name, indexes in found.items()}
            for name in optional:
                indexes = [i for i, value in enumerate(normalized) if value == normalize_header(name)]
                if len(indexes) == 1:
                    columns[name] = indexes[0]
            matches.append((sheet, columns, number))
            break
    if not matches:
        message = f"未找到可导入的工作表。必需列：{'、'.join(required)}"
        if best and len(best[2]) < len(required):
            message += f"；工作表「{best[0]}」第 {best[1]} 行缺少：{'、'.join(best[2])}"
        raise ValueError(message)
    if len(matches) != 1:
        raise SheetChoice([entry[0].title for entry in matches])
    return matches[0]


class SheetChoice(ValueError):
    def __init__(self, sheets):
        self.sheets = sheets
        super().__init__("存在多个符合条件的工作表，请选择")


def identifier(value, row, column):
    if value is None or isinstance(value, bool):
        raise ValueError(f"第 {row} 行「{column}」不能为空，且须为文本或整数")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            raise ValueError(f"第 {row} 行「{column}」须为文本或整数")
        value = str(int(value))
        if len(value.lstrip('-')) > 15:
            raise ValueError(f"第 {row} 行「{column}」超过 Excel 数值精度，请改为文本后重新导出")
    text = str(value).strip()
    if not text or text.startswith('=') or len(text) > 120:
        raise ValueError(f"第 {row} 行「{column}」须为 1–120 字符的原始文本，不能使用公式")
    return text


def decimal_value(value, label, places=6, maximum=Decimal('999999999999999999.999999'), positive=False):
    try:
        if value is None or isinstance(value, bool) or str(value).strip() == '':
            raise InvalidOperation
        result = Decimal(str(value).strip())
        if not result.is_finite() or abs(result) > maximum or result.as_tuple().exponent < -places:
            raise InvalidOperation
        if positive and result <= 0:
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f"{label}须为有效{'正' if positive else ''}数字，最多 {places} 位小数，且不能超出支持范围") from None


def ordered_time(value, row, epoch):
    try:
        if isinstance(value, bool) or value is None:
            raise ValueError
        if isinstance(value, (int, float)):
            value = from_excel(value, epoch)
        if isinstance(value, date) and not isinstance(value, datetime):
            value = datetime.combine(value, datetime.min.time())
        if isinstance(value, str):
            value = datetime.fromisoformat(value.strip().replace('/', '-'))
        if not isinstance(value, datetime) or not 2000 <= value.year <= 2100:
            raise ValueError
        if value.tzinfo is None:
            value = value.replace(tzinfo=LOCAL_TIMEZONE)
        return value.astimezone(timezone.utc).replace(tzinfo=None, microsecond=0)
    except (ValueError, TypeError, OverflowError):
        raise ValueError(f"第 {row} 行「下单时间」无效，请使用 YYYY-MM-DD HH:MM:SS") from None


def parse_workbook(data: bytes, kind: str, sheet_name: str | None = None, shop_type: str = 'cross_border'):
    if shop_type not in {'cross_border', 'local'}:
        raise ValueError('店铺类型无效')
    if len(data) > MAX_BYTES:
        raise ValueError("Excel 文件不能超过 20MB")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 150 * 1024 * 1024:
                raise ValueError("表格解压后不能超过 150MB")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("无法读取 XLSX 文件，请检查文件格式") from exc
    try:
        required = (CROSS_BORDER_BILL_HEADERS if shop_type == 'cross_border' else BILL_HEADERS) if kind == 'platform' else ERP_HEADERS
        sheet, columns, header_row = find_headers(workbook, required, sheet_name)
        bills = {}
        bill_sources = {}
        orders = {}
        previous_order = None
        previous_time = None
        source_rows = ignored = 0
        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if not any(v is not None and str(v).strip() for v in row):
                previous_order = previous_time = None
                continue
            source_rows += 1
            if source_rows > MAX_ROWS:
                raise ValueError("单次最多导入 50,000 条数据")
            values = {name: row[index] if index < len(row) else None for name, index in columns.items()}
            if kind == 'platform':
                if str(values['交易类型'] or '').strip() != '订单':
                    ignored += 1
                    continue
                oid = identifier(values['订单ID/调整单ID'], row_number, '订单 ID')
                settlement_id = identifier(values['结算单ID'], row_number, '结算单 ID') if shop_type == 'cross_border' else ''
                sku_id = identifier(values['SKUID'], row_number, 'SKU ID') if shop_type == 'cross_border' else ''
                currency = str(values['货币'] or '').strip().upper()
                currency = 'CNY' if currency == 'RMB' else currency
                if not re.fullmatch('[A-Z]{3}', currency):
                    raise ValueError(f"第 {row_number} 行「货币」须为 MYR、CNY 等三位货币代码")
                item = (currency,
                    decimal_value(values['结算总金额'], f'第 {row_number} 行「结算总金额」'),
                    decimal_value(values['总收入'], f'第 {row_number} 行「总收入」'))
                key = (oid, settlement_id, sku_id)
                if key in bills:
                    previous = bills[key]
                    if previous[0] != currency:
                        if shop_type == 'cross_border':
                            raise ValueError(f"第 {row_number} 行同一订单、结算单和 SKU 的货币不一致，无法汇总")
                        raise ValueError(f"第 {row_number} 行同一订单的货币不一致，无法汇总")
                    with localcontext() as context:
                        context.prec = 64
                        bills[key] = (currency, previous[1] + item[1], previous[2] + item[2])
                else:
                    bills[key] = item
                bill_sources.setdefault(key, []).append({'settlement_id': settlement_id or None, 'sku_id': sku_id or None,
                    'currency': currency, 'settlement': str(item[1]), 'revenue': str(item[2])})
            else:
                raw_order = values['订单编号']
                oid = identifier(raw_order, row_number, '订单编号') if raw_order is not None and str(raw_order).strip() else previous_order
                if not oid:
                    raise ValueError(f"第 {row_number} 行缺少「订单编号」，无法继承上一订单")
                raw_time = values['下单时间']
                if raw_time is not None and str(raw_time).strip():
                    ordered_at = ordered_time(raw_time, row_number, workbook.epoch)
                elif oid == previous_order:
                    ordered_at = previous_time
                else:
                    raise ValueError(f"第 {row_number} 行新订单缺少「下单时间」")
                sku = identifier(values['平台SKU'], row_number, '平台SKU')
                base = sku.split('-', 1)[0].strip()
                if not base:
                    raise ValueError(f"第 {row_number} 行平台 SKU 去除尺码后为空")
                qty = decimal_value(values['商品数量'], f'第 {row_number} 行「商品数量」', places=0,
                    maximum=Decimal(2147483647), positive=True)
                if oid not in orders:
                    orders[oid] = {'ordered_at': ordered_at, 'items': {}}
                if orders[oid]['ordered_at'] != ordered_at:
                    raise ValueError(f"第 {row_number} 行同订单「下单时间」冲突")
                items = orders[oid]['items']
                quantity = items.get(sku, (base, 0))[1] + int(qty)
                if quantity > 2147483647:
                    raise ValueError(f"第 {row_number} 行同订单 SKU 数量合计超限")
                items[sku] = (base, quantity)
                previous_order, previous_time = oid, ordered_at
        if not bills and not orders:
            raise ValueError("没有可导入的订单数据")
        for currency, settlement, revenue in bills.values():
            decimal_value(settlement, '同订单同结算单的结算总金额合计')
            decimal_value(revenue, '同订单同结算单的总收入合计')
        return {'bills': bills, 'bill_sources': bill_sources, 'orders': orders,
                'source_rows': source_rows, 'ignored': ignored, 'sheet': sheet.title}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("无法读取 XLSX 工作表，请检查文件内容") from exc
    finally:
        workbook.close()
