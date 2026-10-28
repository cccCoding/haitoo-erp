from dataclasses import dataclass
from copy import copy, deepcopy
from functools import lru_cache
from io import BytesIO
from pathlib import Path
import posixpath
import re
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange


TEMPLATE_PATH = Path(__file__).parent / "resources" / "tiktok_seller_batch_upload_zh_v5_0_2.xlsx"
DATA_START_ROW = 7
DATA_END_ROW = 5000
MAX_DATA_ROWS = DATA_END_ROW - DATA_START_ROW + 1
BASE_FIELD_START_COLUMN = 3  # C
OPTIONS_PARSER_VERSION = 3

FIELD_CATEGORY = "category"
FIELD_PRODUCT_NAME = "product_name"
FIELD_PRODUCT_DESCRIPTION = "product_description"
FIELD_MAIN_IMAGE = "main_image"
FIELD_PROPERTY_NAME_1 = "property_name_1"
FIELD_PROPERTY_VALUE_1 = "property_value_1"
FIELD_PROPERTY_1_IMAGE = "property_1_image"
FIELD_PROPERTY_NAME_2 = "property_name_2"
FIELD_PROPERTY_VALUE_2 = "property_value_2"
FIELD_PARCEL_WEIGHT = "parcel_weight"
FIELD_PARCEL_LENGTH = "parcel_length"
FIELD_PARCEL_WIDTH = "parcel_width"
FIELD_PARCEL_HEIGHT = "parcel_height"
FIELD_DELIVERY = "delivery"
FIELD_PRICE = "price"
FIELD_QUANTITY = "quantity"
FIELD_SELLER_SKU = "seller_sku"
FIELD_SIZE_CHART = "size_chart"
FIELD_COD = "cod"
FIELD_SHIPPING_INSURANCE = "shipping_insurance"
FIELD_SPECIAL_PRODUCT_LISTING_TYPE = "special_product_listing_type"
FIELD_AUCTION_STARTING_PRICE = "auction_starting_price"

COMMON_REQUIRED_EXPORT_FIELDS = {
    FIELD_CATEGORY, FIELD_PRODUCT_NAME, FIELD_PRODUCT_DESCRIPTION, FIELD_MAIN_IMAGE,
    FIELD_PROPERTY_NAME_1, FIELD_PROPERTY_VALUE_1, FIELD_PROPERTY_1_IMAGE,
    FIELD_PROPERTY_NAME_2, FIELD_PROPERTY_VALUE_2, FIELD_PARCEL_WEIGHT,
    FIELD_PARCEL_LENGTH, FIELD_PARCEL_WIDTH, FIELD_PARCEL_HEIGHT,
    FIELD_PRICE, FIELD_QUANTITY, FIELD_SELLER_SKU, FIELD_SIZE_CHART,
}


@dataclass(frozen=True)
class TiktokTemplateAdapter:
    key: str
    label: str
    required_fields: frozenset[str]
    supports_cod: bool
    supports_hubstudio_submit: bool


TIKTOK_LOCAL_ADAPTER = TiktokTemplateAdapter(
    key="tiktok_local", label="tk本土店",
    required_fields=frozenset(COMMON_REQUIRED_EXPORT_FIELDS | {FIELD_DELIVERY, FIELD_COD}),
    supports_cod=True, supports_hubstudio_submit=True,
)
TIKTOK_CROSS_BORDER_ADAPTER = TiktokTemplateAdapter(
    key="tiktok_cross_border", label="tk跨境店",
    required_fields=frozenset(COMMON_REQUIRED_EXPORT_FIELDS),
    supports_cod=False, supports_hubstudio_submit=False,
)
TIKTOK_TEMPLATE_ADAPTERS = {
    TIKTOK_LOCAL_ADAPTER.key: TIKTOK_LOCAL_ADAPTER,
    TIKTOK_CROSS_BORDER_ADAPTER.key: TIKTOK_CROSS_BORDER_ADAPTER,
}


def _clean(value) -> str:
    return str(value or "").strip()


def _template_source(template_bytes: bytes | None = None):
    return BytesIO(template_bytes) if template_bytes is not None else TEMPLATE_PATH


def _field_columns(sheet) -> dict[str, int]:
    """按模板第 1 行的机器字段名定位列，避免依赖不同语言模板的固定列号。"""
    return {
        _clean(sheet.cell(1, column).value): column
        for column in range(1, sheet.max_column + 1)
        if _clean(sheet.cell(1, column).value)
    }


def _attribute_fields(sheet) -> list[tuple[int, str]]:
    """按机器字段名查找类目属性，兼容不同模板版本中的起始列。"""
    return [
        (column, field)
        for column in range(1, sheet.max_column + 1)
        if (field := _clean(sheet.cell(1, column).value)).startswith(("product_property/", "qualification/"))
    ]


def _template_version(sheet) -> str | None:
    """本土模板将版本放在 A2，跨境模板则将语言提示放在 A2、版本放在 B2。"""
    values = [_clean(sheet.cell(2, column).value) for column in range(1, sheet.max_column + 1)]
    return next((value for value in values if value.upper().startswith("V") and any(char.isdigit() for char in value)), values[0][:40] or None)


def _shipping_insurance_options(workbook, fields: dict[str, int]) -> list[str]:
    if FIELD_SHIPPING_INSURANCE not in fields:
        return []
    if "ShippingInsurance" not in workbook.sheetnames:
        raise ValueError("双平台模板缺少 ShippingInsurance 工作表")
    options = list(dict.fromkeys(_clean(row[0]) for row in workbook["ShippingInsurance"].iter_rows(values_only=True)
                                if row and _clean(row[0])))
    if "Optional" not in options:
        raise ValueError("运输保险选项缺少默认值 Optional")
    return options


def resolve_shipping_insurance(value: str | None, options: list[str]) -> str | None:
    """双平台默认选择 Optional；不含保险字段的模板不使用该参数。"""
    if not options:
        return None
    selected = "Optional" if value is None else value
    if selected not in options:
        raise ValueError("运输保险必须从模板选项中选择")
    return selected


def tiktok_template_adapter(template_fields: set[str], expected_type: str | None = None) -> TiktokTemplateAdapter:
    """根据机器字段识别本土/跨境模板，并拒绝与用户选择不一致的文件。"""
    if {FIELD_DELIVERY, FIELD_COD}.issubset(template_fields):
        adapter = TIKTOK_LOCAL_ADAPTER
    elif (
        {FIELD_SPECIAL_PRODUCT_LISTING_TYPE, FIELD_AUCTION_STARTING_PRICE}.issubset(template_fields)
        or (
            COMMON_REQUIRED_EXPORT_FIELDS.issubset(template_fields)
            and FIELD_DELIVERY not in template_fields
            and FIELD_COD not in template_fields
        )
    ):
        # 新版跨境模板不再包含拍卖相关字段，属性列也可能从 AA 开始。
        adapter = TIKTOK_CROSS_BORDER_ADAPTER
    else:
        raise ValueError("无法识别 TikTok 模板类型：未找到本土店或跨境店的特征字段")
    if expected_type and adapter.key != expected_type:
        raise ValueError(f"所选{TIKTOK_TEMPLATE_ADAPTERS[expected_type].label if expected_type in TIKTOK_TEMPLATE_ADAPTERS else '店铺类型'}与上传模板不匹配，识别结果为{adapter.label}")
    return adapter


def _default_attribute_input_mode(field: dict) -> str:
    if not field.get("options"):
        return "text"
    if field.get("allow_custom"):
        return "select_or_text"
    return "select"


def attribute_input_mode(field: dict) -> str:
    """返回属性的输入方式，并兼容尚未保存 input_mode 的旧类目库。"""
    if field.get("input_mode") in {"text", "select", "select_or_text"}:
        return field["input_mode"]
    if field.get("input_mode") in {"single", "multiple"}:
        return "select"
    return _default_attribute_input_mode(field)


def _attribute_options_by_column(template_sheet, attribute_sheet, attribute_fields) -> dict[int, dict[str, list[str]]]:
    """按数据验证中的实际引用定位选项，属性列和 HiddenAttr 列对不一定一一对应。"""
    references: dict[int, tuple[int, int, int]] = {}
    columns = {column for column, _ in attribute_fields}
    hidden_sheet = r"(?:'HiddenAttr'|HiddenAttr)!"
    reference_pattern = re.compile(
        rf"^=?\s*OFFSET\(\s*{hidden_sheet}\$?([A-Z]+)\$?1\s*,\s*"
        rf"MATCH\([^,]+,\s*{hidden_sheet}\$?([A-Z]+)\$?1:\$?\2\$?(\d+)\s*,\s*0\s*\)",
        re.I,
    )
    for validation in template_sheet.data_validations.dataValidation:
        if validation.type != "list":
            continue
        targets = {column for area in validation.sqref.ranges
                   if area.min_row <= DATA_START_ROW <= area.max_row
                   for column in columns if area.min_col <= column <= area.max_col}
        if not targets:
            continue
        match = reference_pattern.match(_clean(validation.formula1))
        if not match:
            raise ValueError(f"无法解析 TikTok 属性列 {', '.join(get_column_letter(column) for column in sorted(targets))} 的下拉引用")
        value_column, key_column = (column_index_from_string(value.upper()) for value in match.group(1, 2))
        end_row = int(match.group(3))
        if max(key_column, value_column) > attribute_sheet.max_column or end_row > attribute_sheet.max_row:
            raise ValueError("TikTok 属性下拉引用超出 HiddenAttr 数据范围")
        reference = (key_column, value_column, end_row)
        for column in targets:
            if column in references and references[column] != reference:
                raise ValueError(f"TikTok 属性列 {get_column_letter(column)} 存在冲突的下拉引用")
            references[column] = reference

    result = {}
    for column, (key_column, value_column, end_row) in references.items():
        options: dict[str, list[str]] = {}
        for row in range(1, end_row + 1):
            category = _clean(attribute_sheet.cell(row, key_column).value)
            value = _clean(attribute_sheet.cell(row, value_column).value)
            if category and value:
                options.setdefault(category, []).append(value)
        result[column] = {category: list(dict.fromkeys(values)) for category, values in options.items()}
    return result


def _template_sheet_part(package: ZipFile) -> str:
    """上传模板允许工作表重排，不能假定 Template 一定位于 sheet1.xml。"""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    workbook = ElementTree.fromstring(package.read("xl/workbook.xml"))
    sheets = workbook.findall("m:sheets/m:sheet", ns)
    template_sheet = next((sheet for sheet in sheets if sheet.attrib.get("name") == "Template"), None)
    if template_sheet is None:
        raise ValueError("TikTok 模板缺少 Template 工作表")
    relationship_id = template_sheet.attrib.get(f"{{{ns['r']}}}id")
    relationships = ElementTree.fromstring(package.read("xl/_rels/workbook.xml.rels"))
    relationship = next((item for item in relationships if item.attrib.get("Id") == relationship_id), None)
    if relationship is None or relationship.attrib.get("TargetMode") == "External":
        raise ValueError("无法安全定位 TikTok 模板数据工作表")
    target = relationship.attrib.get("Target", "")
    part = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target))
    if not part.startswith("xl/worksheets/") or part not in package.namelist():
        raise ValueError("TikTok 模板数据工作表路径不合法")
    return part


def _preserve_template_package(generated_bytes: bytes, template_bytes: bytes | None = None) -> bytes:
    """仅替换数据工作表 XML，避免 openpyxl 丢弃说明页中的绘图和媒体。"""
    generated_stream = BytesIO(generated_bytes)
    output = BytesIO()
    with ZipFile(_template_source(template_bytes), "r") as source_zip, ZipFile(generated_stream, "r") as generated_zip, ZipFile(output, "w", ZIP_DEFLATED) as output_zip:
        source_part = _template_sheet_part(source_zip)
        generated_part = _template_sheet_part(generated_zip)
        generated_template_xml = generated_zip.read(generated_part)
        for info in source_zip.infolist():
            content = generated_template_xml if info.filename == source_part else source_zip.read(info.filename)
            output_zip.writestr(info, content)
    return output.getvalue()


def parse_listing_options(template_bytes: bytes | None = None, template_type: str | None = None) -> dict:
    """解析 TikTok 模板中的类目、字段状态和属性选项。"""
    try:
        workbook = load_workbook(_template_source(template_bytes), data_only=False, read_only=False)
    except Exception as exc:
        raise ValueError("无法读取 TikTok XLSX 模板") from exc
    required_sheets = {"Category", "Template", "HiddenStyle", "HiddenAttr"}
    missing_sheets = sorted(required_sheets - set(workbook.sheetnames))
    if missing_sheets:
        raise ValueError(f"TikTok 模板缺少工作表：{', '.join(missing_sheets)}")
    category_sheet = workbook["Category"]
    template_sheet = workbook["Template"]
    style_sheet = workbook["HiddenStyle"]
    attribute_sheet = workbook["HiddenAttr"]
    template_fields = _field_columns(template_sheet)
    adapter = tiktok_template_adapter(set(template_fields), template_type)
    missing_fields = sorted(adapter.required_fields - set(template_fields))
    if missing_fields:
        raise ValueError(f"TikTok 模板缺少导出字段：{', '.join(missing_fields)}")
    attribute_fields = _attribute_fields(template_sheet)
    attribute_options = _attribute_options_by_column(template_sheet, attribute_sheet, attribute_fields)
    insurance_options = _shipping_insurance_options(workbook, template_fields)

    categories = [
        {"name": _clean(category_sheet.cell(row, 1).value), "id": category_sheet.cell(row, 2).value}
        for row in range(1, category_sheet.max_row + 1)
        if _clean(category_sheet.cell(row, 1).value)
    ]
    category_rows = {
        _clean(style_sheet.cell(row, 1).value): row
        for row in range(1, style_sheet.max_row + 1)
        if _clean(style_sheet.cell(row, 1).value)
    }
    attributes_by_category: dict[str, list[dict]] = {}
    base_requirements_by_category: dict[str, dict[str, str]] = {}
    for category in categories:
        category_name = category["name"]
        style_row = category_rows.get(category_name)
        fields = []
        if style_row is None:
            attributes_by_category[category_name] = fields
            base_requirements_by_category[category_name] = {}
            continue
        base_requirements_by_category[category_name] = {
            field: _clean(style_sheet.cell(style_row, column).value)
            for column in range(BASE_FIELD_START_COLUMN, template_sheet.max_column + 1)
            if not (field := _clean(template_sheet.cell(1, column).value)).startswith(("product_property/", "qualification/"))
            and _clean(style_sheet.cell(style_row, column).value)
            and _clean(style_sheet.cell(style_row, column).value) != "Forbid"
        }
        for column, field in attribute_fields:
            status = _clean(style_sheet.cell(style_row, column).value)
            if not status or status == "Forbid":
                continue
            label = _clean(template_sheet.cell(3, column).value)
            description = _clean(template_sheet.cell(5, column).value)
            options = attribute_options.get(column, {}).get(category_name, [])
            allow_custom = "custom value" in description.lower() or "输入自定义值" in description
            if not options:
                input_mode = "text"
            elif allow_custom:
                input_mode = "select_or_text"
            else:
                input_mode = "select"
            fields.append({
                "field": field,
                "column": get_column_letter(column),
                "label": label,
                "status": status,
                "required": status == "Mandatory",
                "options": list(dict.fromkeys(options)),
                "allow_custom": allow_custom,
                "input_type": "url" if field.startswith("qualification/") else "select",
                "input_mode": input_mode,
            })
        attributes_by_category[category_name] = fields

    return {
        "parser_version": OPTIONS_PARSER_VERSION,
        "template_version": _template_version(template_sheet),
        "template_type": adapter.key,
        "template_type_label": adapter.label,
        "capabilities": {
            "supports_cod": adapter.supports_cod,
            "supports_hubstudio_submit": adapter.supports_hubstudio_submit,
            "supports_shipping_insurance": bool(insurance_options),
        },
        "categories": categories,
        "attributes_by_category": attributes_by_category,
        "base_requirements_by_category": base_requirements_by_category,
        "cod_options": ["Y", "N"],
        "shipping_insurance_options": insurance_options,
        "shipping_insurance_default": "Optional" if insurance_options else None,
        "price_limits": {"min": 0.01, "max": 999999},
        "quantity_limits": {"min": 0, "max": 999999},
        "multiple_value_separator": ",",
    }


@lru_cache(maxsize=1)
def listing_options() -> dict:
    """兼容系统默认模板的只读选项。"""
    return parse_listing_options()


@lru_cache(maxsize=8)
def _reparsed_listing_options(template_bytes: bytes, template_type: str) -> dict:
    return parse_listing_options(template_bytes, template_type)


def catalog_listing_options(catalog) -> dict:
    """读取旧类目库时从原始模板修正选项，不要求重新上传，也不在只读请求中写库。"""
    options = catalog.parsed_options or {}
    if catalog.template_type not in TIKTOK_TEMPLATE_ADAPTERS or options.get("parser_version") == OPTIONS_PARSER_VERSION:
        return options
    refreshed = deepcopy(_reparsed_listing_options(bytes(catalog.template_blob), catalog.template_type))
    for category, fields in refreshed["attributes_by_category"].items():
        previous = {field["field"]: field for field in options.get("attributes_by_category", {}).get(category, [])}
        for field in fields:
            old = previous.get(field["field"])
            if old and (old.get("input_mode_overridden") or attribute_input_mode(old) != _default_attribute_input_mode(old)):
                mode = attribute_input_mode(old)
                field["input_mode"] = mode if field["options"] else "text"
                field["input_mode_overridden"] = True
    return refreshed


def category_attributes(category: str, options: dict | None = None) -> list[dict]:
    options = options or listing_options()
    if category not in options["attributes_by_category"]:
        raise ValueError("TikTok 类目不存在，请重新选择")
    return options["attributes_by_category"][category]


def category_base_requirements(category: str, options: dict | None = None) -> dict[str, str]:
    options = options or listing_options()
    if category not in options["base_requirements_by_category"]:
        raise ValueError("TikTok 类目不存在，请重新选择")
    return options["base_requirements_by_category"][category]


def validate_attributes(category: str, attributes: dict, options: dict | None = None) -> dict[str, str]:
    options = options or listing_options()
    fields = {item["field"]: item for item in category_attributes(category, options)}
    cleaned: dict[str, str] = {}
    for key, raw_value in attributes.items():
        field = fields.get(key)
        input_mode = attribute_input_mode(field) if field else "select"
        if isinstance(raw_value, list):
            values = list(dict.fromkeys(_clean(item) for item in raw_value if _clean(item)))
            if not field or input_mode != "select":
                raise ValueError(f"{field['label'] if field else key} 不支持填写多个值")
        else:
            raw_text = _clean(raw_value)
            if field and input_mode == "select":
                values = list(dict.fromkeys(_clean(item) for item in raw_text.replace("，", ",").split(",") if _clean(item)))
            else:
                values = [raw_text] if raw_text else []
        value = options.get("multiple_value_separator", ",").join(values) if field and input_mode == "select" else (values[0] if values else "")
        if value:
            cleaned[key] = value
    unknown = sorted(set(cleaned) - set(fields))
    if unknown:
        raise ValueError(f"所选类目不支持属性：{', '.join(unknown)}")
    for key, value in cleaned.items():
        field = fields[key]
        if len(value) > 500:
            raise ValueError(f"{field['label']} 不能超过 500 个字符")
        if field["input_type"] == "url" and not value.lower().startswith(("http://", "https://")):
            raise ValueError(f"{field['label']} 必须填写 http:// 或 https:// 开头的 URL")
        input_mode = attribute_input_mode(field)
        selected_values = value.split(options.get("multiple_value_separator", ",")) if input_mode == "select" else [value]
        if input_mode == "select" and field["options"] and any(item not in field["options"] for item in selected_values):
            raise ValueError(f"{field['label']} 必须从模板选项中选择")
    missing = [item["label"] for item in fields.values() if item["required"] and not cleaned.get(item["field"])]
    if missing:
        raise ValueError(f"请填写必填类目属性：{', '.join(missing)}")
    return cleaned


def _shift_row_formula(formula: str | None, source_row: int, target_row: int) -> str | None:
    if not formula:
        return formula
    leading_equals = formula.startswith("=")
    tokens = Tokenizer(formula if leading_equals else f"={formula}")
    for token in tokens.items:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        reference = token.value
        if "!" in reference:
            sheet_name, reference = reference.rsplit("!", 1)
            if sheet_name.strip("'") != "Template":
                continue
            prefix = f"{sheet_name}!"
        else:
            prefix = ""
        token.value = prefix + re.sub(rf"(?<![A-Za-z0-9_])(\$?[A-Z]{{1,3}}\$?){source_row}(?!\d)", rf"\g<1>{target_row}", reference)
    result = tokens.render()
    return result if leading_equals else result[1:]


def _expand_range(ranges) -> MultiCellRange:
    expanded = []
    for original in ranges:
        item = copy(original)
        if item.min_row == DATA_START_ROW and DATA_START_ROW < item.max_row < DATA_END_ROW:
            item.max_row = DATA_END_ROW
        expanded.append(item)
    return MultiCellRange(expanded)


def _prepare_data_capacity(sheet) -> None:
    """将可安全扩展的模板数据区、样式和验证规则统一到第 5000 行。"""
    if sheet.max_column > 512 or sheet.max_row > 50000 or sheet.tables:
        raise ValueError("TikTok 模板数据区结构无法安全扩展")
    if any(item.max_row >= DATA_START_ROW for item in sheet.merged_cells.ranges):
        raise ValueError("TikTok 模板数据区包含合并单元格，无法安全扩展")
    source_row = max(DATA_START_ROW, min(sheet.max_row, DATA_END_ROW))
    styles = {
        column: copy(sheet.cell(source_row, column)._style)
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(source_row, column).has_style
    }
    # 使用已有单元格，避免由异常 dimension 声明分配大量空单元格。
    for cell in list(sheet._cells.values()):
        if cell.row >= DATA_START_ROW:
            cell.value = None
            if not cell.has_style and not cell.comment and not cell.hyperlink:
                del sheet._cells[(cell.row, cell.column)]
    for row_number in range(source_row + 1, DATA_END_ROW + 1):
        for column, style in styles.items():
            sheet.cell(row_number, column)._style = copy(style)
        if source_row in sheet.row_dimensions:
            dimension = copy(sheet.row_dimensions[source_row])
            dimension.index = row_number
            sheet.row_dimensions[row_number] = dimension

    singleton_rules: dict[int, dict[int, object]] = {}
    for validation in list(sheet.data_validations.dataValidation):
        ranges = list(validation.sqref.ranges)
        if len(ranges) == 1:
            item = ranges[0]
            if item.min_row == item.max_row and item.min_col == item.max_col and DATA_START_ROW <= item.min_row <= DATA_END_ROW:
                singleton_rules.setdefault(item.min_col, {})[item.min_row] = validation
        validation.sqref = _expand_range(ranges)
    # 部分类目属性使用每一行独立的绝对引用；不能简单延长范围共用 $A$7。
    for column, rules in singleton_rules.items():
        rows = sorted(rules)
        if rows[0] != DATA_START_ROW or rows != list(range(DATA_START_ROW, rows[-1] + 1)) or rows[-1] >= DATA_END_ROW:
            continue
        prototype = rules[rows[-1]]
        first = rules[DATA_START_ROW]
        if any(_shift_row_formula(getattr(prototype, field), rows[-1], DATA_START_ROW) != getattr(first, field) for field in ("formula1", "formula2")):
            raise ValueError("TikTok 模板的逐行属性验证规则无法安全扩展")
        for row_number in range(rows[-1] + 1, DATA_END_ROW + 1):
            validation = copy(prototype)
            validation.sqref = MultiCellRange([CellRange(min_col=column, max_col=column, min_row=row_number, max_row=row_number)])
            validation.formula1 = _shift_row_formula(prototype.formula1, rows[-1], row_number)
            validation.formula2 = _shift_row_formula(prototype.formula2, rows[-1], row_number)
            sheet.add_data_validation(validation)
    conditional_rules = list(sheet.conditional_formatting._cf_rules.items())
    sheet.conditional_formatting._cf_rules.clear()
    for conditional, rules in conditional_rules:
        conditional.sqref = _expand_range(conditional.sqref.ranges)
        for rule in rules:
            sheet.conditional_formatting.add(conditional, rule)


def _set_literal_value(cell, value) -> None:
    cell.value = value
    if isinstance(value, str):
        # 强制 inlineStr，不让用户或 AI 文本被 openpyxl 推断为公式。
        cell.data_type = "s"


def build_workbook(*, template, category: str, cod: str, attributes: dict[str, str], products: list[dict], template_bytes: bytes | None = None, template_type: str | None = None, shipping_insurance: str | None = None) -> bytes:
    workbook = load_workbook(_template_source(template_bytes), data_only=False, read_only=False)
    sheet = workbook["Template"]
    field_columns = _field_columns(sheet)
    adapter = tiktok_template_adapter(set(field_columns), template_type)
    missing_fields = sorted(adapter.required_fields - set(field_columns))
    if missing_fields:
        raise ValueError(f"TikTok 模板缺少导出字段：{', '.join(missing_fields)}")
    insurance_value = resolve_shipping_insurance(shipping_insurance, _shipping_insurance_options(workbook, field_columns))
    sizes = [
        _clean(size)
        for size in ((template.sku_specifications or {}).get("size", {}).get("options", []))
        if _clean(size)
    ] or ["Default"]
    projected_rows = sum(len(product.get("sku_images") or list(dict.fromkeys(product["image_urls"]))) * len(sizes) for product in products)
    if projected_rows > MAX_DATA_ROWS:
        raise ValueError(f"导出数据超过 TikTok 模板 {MAX_DATA_ROWS} 行限制")
    _prepare_data_capacity(sheet)
    data_end_row = DATA_END_ROW
    row_number = DATA_START_ROW
    attribute_columns = {field: column for column, field in _attribute_fields(sheet)}
    weight_grams = round(float(template.package_weight) * 1000, 3)
    for product in products:
        image_urls = list(dict.fromkeys(product["image_urls"]))
        gallery = image_urls[:9]
        sku_images = product.get("sku_images") or [
            {"image_url": image_url, "sku": product["base_sku_by_image"][image_url]}
            for image_url in image_urls
        ]
        for color_index, sku_image in enumerate(sku_images, start=1):
            image_url = sku_image["image_url"]
            base_sku = sku_image["sku"]
            color_name = f"Color {color_index}"
            for size in sizes:
                if row_number > data_end_row:
                    raise ValueError(f"导出数据超过 TikTok 模板 {data_end_row - DATA_START_ROW + 1} 行限制")
                seller_sku = base_sku if size == "Default" else f"{base_sku}-{size}"
                values = {
                    FIELD_CATEGORY: category,
                    FIELD_PRODUCT_NAME: product["title"],
                    FIELD_PRODUCT_DESCRIPTION: product["description"],
                    FIELD_PROPERTY_NAME_1: "Color",
                    FIELD_PROPERTY_VALUE_1: color_name,
                    # 模板 P 列 property_1_image 的中文表头是“主要变体图片 1”。
                    FIELD_PROPERTY_1_IMAGE: image_url,
                    FIELD_PROPERTY_NAME_2: "Size",
                    FIELD_PROPERTY_VALUE_2: size,
                    FIELD_PARCEL_WEIGHT: weight_grams,
                    FIELD_PARCEL_LENGTH: template.package_length,
                    FIELD_PARCEL_WIDTH: template.package_width,
                    FIELD_PARCEL_HEIGHT: template.package_height,
                    FIELD_PRICE: product["price"],
                    FIELD_QUANTITY: product["quantity"],
                    FIELD_SELLER_SKU: seller_sku,
                    FIELD_SIZE_CHART: product.get("size_chart_url"),
                }
                if adapter.key == TIKTOK_LOCAL_ADAPTER.key:
                    values[FIELD_DELIVERY] = None
                    values[FIELD_COD] = cod
                else:
                    for field in (FIELD_SPECIAL_PRODUCT_LISTING_TYPE, FIELD_AUCTION_STARTING_PRICE):
                        if field in field_columns:
                            values[field] = None
                if insurance_value is not None:
                    values[FIELD_SHIPPING_INSURANCE] = insurance_value
                for index, gallery_url in enumerate(gallery):
                    values[FIELD_MAIN_IMAGE if index == 0 else f"image_{index + 1}"] = gallery_url
                for field, value in values.items():
                    _set_literal_value(sheet.cell(row_number, field_columns[field]), value)
                for field, value in attributes.items():
                    _set_literal_value(sheet.cell(row_number, attribute_columns[field]), value)
                row_number += 1

    output = BytesIO()
    workbook.save(output)
    return _preserve_template_package(output.getvalue(), template_bytes)
