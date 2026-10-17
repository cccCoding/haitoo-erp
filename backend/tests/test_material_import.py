import asyncio
from datetime import datetime, timedelta, timezone
from io import BytesIO
import unittest
from unittest.mock import patch
import httpx

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import main
from app.database import Base
from app.material_import import build_import_template, parse_import_workbook, validate_image_link, MAX_IMPORT_BYTES
from app.models import Company, MaterialAsset, ProductTemplate, Role, User
from app.schemas import MaterialImportConfirmInput


def xlsx(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class MaterialImportTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        with self.sessions() as db:
            db.add_all([
                Company(id=1, name='A'), Company(id=2, name='B'),
                User(id=1, company_id=1, email='a@example.com', name='A', password_hash='x', user_code='AA', role=Role.MEMBER),
                ProductTemplate(id=1, company_id=1, name='Y1', cover_url='https://img.example/cover'),
                ProductTemplate(id=2, company_id=2, name='Y2', cover_url='https://img.example/cover'),
                MaterialAsset(company_id=1, template_id=2, claimed_by=99, name='old', url='https://img.example/existing', usage_status='used', sku='OLD'),
                MaterialAsset(company_id=2, claimed_by=99, name='foreign', url='https://img.example/foreign', sku='FOREIGN'),
            ])
            db.commit()
            from tests.image_file_helpers import index_urls
            files=index_urls(db,1,["https://img.example/existing"])
            db.scalar(select(MaterialAsset).where(MaterialAsset.url=="https://img.example/existing")).image_file_id=files[0].id
            db.commit()

    def tearDown(self):
        self.engine.dispose()

    def preview(self, data, name='import.xlsx'):
        with self.sessions() as db:
            return asyncio.run(main.preview_material_import(UploadFile(filename=name, file=BytesIO(data)), user=db.get(User, 1), db=db))

    def test_template_and_parser(self):
        template = main.download_material_import_template().body
        workbook = load_workbook(BytesIO(template))
        self.assertEqual(workbook.active['A1'].value, 'SKU图')
        workbook.close()
        with self.assertRaisesRegex(ValueError, '没有可导入'):
            parse_import_workbook(template)
        data = xlsx([['备注', 'SKU图'], ['a', ' https://img.example/a?q=A '], ['b', None], ['c', 'https://img.example/a?q=A'], ['d', 'https://img.example/a?q=a']])
        self.assertEqual(parse_import_workbook(data), ['https://img.example/a?q=A', 'https://img.example/a?q=a'])

    def test_parser_errors(self):
        cases = [([['图片']], '缺少'), ([['SKU图', 'SKU图']], '只能有一个'), ([['SKU图'], ['ftp://example.com/a']], '第 2 行'), ([['SKU图'], [123]], '第 2 行'), ([['SKU图'], ['=HYPERLINK("https://example.com","a")']], '第 2 行'), ([['SKU图']] + [['https://img.example/a']] * 1001, '超限')]
        for rows, message in cases:
            with self.subTest(message=message), self.assertRaisesRegex(ValueError, message):
                parse_import_workbook(xlsx(rows))
        for data, message in [(b'broken', '无法读取'), (b'x' * (MAX_IMPORT_BYTES + 1), '10MB')]:
            with self.assertRaisesRegex(ValueError, message):
                parse_import_workbook(data)
        for url in ['https:///a', 'https://example.com:bad/a', 'https://example.com/a b', 'https://example.com/' + 'a' * 500]:
            with self.assertRaises(ValueError):
                validate_image_link(url, 2)
        with self.assertRaises(HTTPException):
            self.preview(xlsx([['SKU图'], ['https://img.example/a']]), 'a.xls')

    def test_company_wide_preview_is_read_only(self):
        result = self.preview(xlsx([['SKU图'], ['https://img.example/existing'], ['https://img.example/existing'], ['https://img.example/foreign']]))
        self.assertEqual(result['items'], [dict(url='https://img.example/existing', exists=True), dict(url='https://img.example/foreign', exists=False)])
        with self.sessions() as db:
            self.assertEqual(db.scalar(select(func.count()).select_from(MaterialAsset)), 2)



    def test_permissions_and_missing_code(self):
        with self.sessions() as db:
            user = db.get(User, 1)
            for template_id in (2, 999):
                with self.assertRaises(HTTPException) as error:
                    main.confirm_material_import(MaterialImportConfirmInput(request_key="import-test-key",template_id=template_id, urls=['https://img.example/a']), user=user, db=db)
                self.assertEqual(error.exception.status_code, 404)
            user.user_code = None
            with self.assertRaises(HTTPException) as error:
                main.confirm_material_import(MaterialImportConfirmInput(request_key="import-test-key",template_id=1, urls=['https://img.example/a']), user=user, db=db)
            self.assertIn('用户代码', error.exception.detail)


    def test_material_type_filter_includes_legacy_sources_and_scopes(self):
        with self.sessions() as db:
            db.add_all([
                MaterialAsset(company_id=1, claimed_by=1, template_id=1, source_task_id=123, name='AI', url='https://img.example/ai', sku='AI'),
                MaterialAsset(company_id=1, claimed_by=1, template_id=1, name='local', url='https://img.example/local', sku='LOCAL'),
                MaterialAsset(company_id=1, claimed_by=1, template_id=1, source_type='imported', name='imported', url='https://img.example/imported', sku='IMPORTED'),
                MaterialAsset(company_id=2, claimed_by=1, source_type='imported', name='foreign', url='https://img.example/foreign-imported', sku='FOREIGN_IMPORT'),
            ])
            db.commit()
            user = db.get(User, 1)
            for source, expected_sku in [('ai_created', 'AI'), ('local_upload', 'LOCAL'), ('imported', 'IMPORTED')]:
                result = main.list_material_assets(page=1, page_size=20, creator_id=None, template_id=1, usage_status='unused', source_type=source, user=user, db=db)
                self.assertEqual(result['total'], 1)
                self.assertEqual([item['sku'] for item in result['items']], [expected_sku])
                self.assertEqual(result['items'][0]['source_type'], source)
            with self.assertRaises(HTTPException) as error:
                main.list_material_assets(source_type='unsupported', user=user, db=db)
            self.assertEqual(error.exception.status_code, 422)

    def test_creation_time_filter_uses_utc_inclusive_boundaries(self):
        with self.sessions() as db:
            for index, instant in enumerate([datetime(2026, 10, 10, 0), datetime(2026, 10, 10, 1), datetime(2026, 10, 10, 2)]):
                db.add(MaterialAsset(company_id=1, claimed_by=1, name=str(index), url=f'https://img.example/date-{index}', sku=f'DATE{index}', created_at=instant))
            db.commit()
            user = db.get(User, 1)
            hk = timezone(timedelta(hours=8))
            def query(**dates):
                return main.list_material_assets(page=1, page_size=20, creator_id=None, template_id=None, user=user, db=db, **dates)
            result = query(created_from=datetime(2026, 10, 10, 8, tzinfo=hk), created_to=datetime(2026, 10, 10, 9, tzinfo=hk))
            self.assertEqual([item['sku'] for item in result['items']], ['DATE1', 'DATE0'])
            self.assertEqual(result['total'], 2)
            self.assertEqual(query(created_from=datetime(2026, 10, 10, 2))['total'], 1)
            self.assertEqual(query(created_to=datetime(2026, 10, 10, 0))['total'], 1)
            with self.assertRaises(HTTPException) as error:
                query(created_from=datetime(2026, 10, 10, 2), created_to=datetime(2026, 10, 10, 9, tzinfo=hk))
            self.assertEqual(error.exception.status_code, 400)



if __name__ == '__main__':
    unittest.main()
