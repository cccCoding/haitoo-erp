from io import BytesIO
from copy import copy
from types import SimpleNamespace
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.tiktok_export import DATA_END_ROW, MAX_DATA_ROWS, TEMPLATE_PATH, TIKTOK_LOCAL_ADAPTER, build_workbook


class TiktokCapacityTests(unittest.TestCase):
    def setUp(self):
        self.template = SimpleNamespace(sku_specifications={"size": {"options": ["S"]}}, package_weight=0.2, package_length=20, package_width=10, package_height=2)
        workbook = Workbook()
        workbook.active.title = "Instruction"
        workbook.active["A1"] = "Keep this sheet and package unchanged"
        sheet = workbook.create_sheet("Template")
        self.fields = list(sorted(TIKTOK_LOCAL_ADAPTER.required_fields)) + [f"image_{i}" for i in range(2, 10)]
        for column, field in enumerate(self.fields, 1):
            sheet.cell(1, column, field)
        sheet["A7"].fill = PatternFill(fill_type="solid", fgColor="ABCDEF")
        sheet["A20"].fill = copy(sheet["A7"].fill)
        validation = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="255")
        validation.add("A7:A20")
        sheet.add_data_validation(validation)
        for row in range(7, 21):
            validation = DataValidation(type="custom", formula1=f'LEN($B${row})>0')
            validation.add(f"B{row}")
            sheet.add_data_validation(validation)
        sheet.conditional_formatting.add("A7:A20", CellIsRule(operator="greaterThan", formula=["0"]))
        output = BytesIO()
        workbook.save(output)
        self.uploaded_template = output.getvalue()

    def product(self, count=1, **overrides):
        product = {"title": "A valid descriptive product title", "description": "Product details", "image_urls": ["https://images.example/a.png"],
            "sku_images": [{"image_url": "https://images.example/a.png", "sku": f"SOURCE{index}"} for index in range(count)],
            "size_chart_url": "https://images.example/sizes.png", "price": 10, "quantity": 999}
        product.update(overrides)
        return product

    def export(self, product, template_bytes=None, attributes=None):
        return build_workbook(template=self.template, category="Clothing", cod="Y", attributes=attributes or {}, products=[product],
            template_bytes=template_bytes or self.uploaded_template, template_type="tiktok_local")

    def test_exact_capacity_and_next_row_rejected_for_uploaded_template(self):
        result = self.export(self.product(MAX_DATA_ROWS))
        sheet = load_workbook(BytesIO(result))["Template"]
        self.assertEqual(sheet.max_row, DATA_END_ROW)
        sku_column = self.fields.index("seller_sku") + 1
        self.assertEqual(sheet.cell(5000, sku_column).value, f"SOURCE{MAX_DATA_ROWS - 1}-S")
        with self.assertRaisesRegex(ValueError, "4994 行限制"):
            self.export(self.product(MAX_DATA_ROWS + 1))

    def test_uploaded_template_reordering_preserves_other_package_parts(self):
        result = self.export(self.product())
        with ZipFile(BytesIO(self.uploaded_template)) as original, ZipFile(BytesIO(result)) as generated:
            self.assertEqual(original.namelist(), generated.namelist())
            for name in original.namelist():
                if name != "xl/worksheets/sheet2.xml":
                    self.assertEqual(original.read(name), generated.read(name), name)
        sheet = load_workbook(BytesIO(result))["Template"]
        self.assertEqual(sheet.cell(7, self.fields.index("product_name") + 1).value, self.product()["title"])
        self.assertEqual(sheet["A5000"].fill.fgColor.rgb, "00ABCDEF")
        self.assertIn("A7:A5000", [str(item.sqref) for item in sheet.data_validations.dataValidation])
        last_validation = next(item for item in sheet.data_validations.dataValidation if str(item.sqref) == "B5000")
        self.assertEqual(last_validation.formula1, "LEN($B$5000)>0")
        self.assertIn("A7:A5000", [str(item.sqref) for item in sheet.conditional_formatting])

    def test_formula_looking_product_text_is_literal(self):
        values = {"title": '=HYPERLINK("https://evil.example","click")', "description": "+SUM(1,2)",
            "sku_images": [{"image_url": "https://images.example/a.png", "sku": "=BAD()"}]}
        result = self.export(self.product(**values))
        sheet = load_workbook(BytesIO(result), data_only=False)["Template"]
        for field, expected in (("product_name", values["title"]), ("product_description", values["description"]), ("seller_sku", "=BAD()-S")):
            cell = sheet.cell(7, self.fields.index(field) + 1)
            self.assertEqual(cell.value, expected)
            self.assertEqual(cell.data_type, "s")

    def test_bundled_template_expands_per_row_validation_beyond_old_limit(self):
        result = self.export(self.product(2000), template_bytes=TEMPLATE_PATH.read_bytes())
        sheet = load_workbook(BytesIO(result))["Template"]
        self.assertEqual(sheet["Z2006"].value, "SOURCE1999-S")
        validation = next(item for item in sheet.data_validations.dataValidation if str(item.sqref) == "AJ5000")
        self.assertIn("$A$5000", validation.formula1)
        self.assertIn("HiddenAttr!$P$1", validation.formula1)

    def test_unsafe_merged_data_area_is_rejected(self):
        workbook = load_workbook(BytesIO(self.uploaded_template))
        workbook["Template"].merge_cells("A7:B7")
        output = BytesIO()
        workbook.save(output)
        with self.assertRaisesRegex(ValueError, "合并单元格"):
            self.export(self.product(), template_bytes=output.getvalue())
