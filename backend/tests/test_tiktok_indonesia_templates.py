import asyncio
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import main
from app.database import Base
from app.models import Company, ProductDraft, ProductTemplate, Role, TiktokCategoryCatalog, User
from app.schemas import TiktokCategoryCatalogUpdate, TiktokDraftExportInput
from app.tiktok_export import OPTIONS_PARSER_VERSION, build_workbook, parse_listing_options, validate_attributes


FIXTURE_DIRECTORY = Path(__file__).parent / "fixtures" / "tiktok" / "indonesia"
CASES = (
    {
        "file": "menswear_20261026.xlsx", "name": "印尼男装", "count": 25,
        "category": "Men's Tops/Waistcoats & Gilets", "id": 601194,
        "pattern": "Plain", "forbidden": "product_property/100395",
        "season_category": "Men's Tops/Waistcoats & Gilets", "season_column": "AI",
    },
    {
        "file": "muslim_fashion_20261026.xlsx", "name": "印尼穆斯林", "count": 47,
        "category": "Hijabs/Underscarves", "id": 601305,
        "pattern": "Flowers", "forbidden": "product_property/100347",
        "season_category": "Hijabs/Underscarves", "season_column": "AK",
    },
    {
        "file": "womenswear_20261026.xlsx", "name": "印尼女装", "count": 41,
        "category": "Women's Underwear/Panties", "id": 601247,
        "pattern": "Plain", "forbidden": "product_property/100395",
        "season_category": "Women's Underwear/Bras", "season_column": "AJ",
    },
)


class TiktokIndonesiaTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.templates = {case["file"]: (FIXTURE_DIRECTORY / case["file"]).read_bytes() for case in CASES}
        cls.options = {name: parse_listing_options(content, "tiktok_local") for name, content in cls.templates.items()}

    def test_english_categories_attributes_and_shifted_columns(self) -> None:
        for case in CASES:
            with self.subTest(template=case["file"]):
                options = self.options[case["file"]]
                self.assertEqual(options["template_type"], "tiktok_local")
                self.assertEqual(options["template_version"], "V5.0.2")
                self.assertEqual(len(options["categories"]), case["count"])
                self.assertEqual(options["categories"][0], {"name": case["category"], "id": case["id"]})
                self.assertTrue(options["capabilities"]["supports_cod"])
                self.assertTrue(options["capabilities"]["supports_hubstudio_submit"])
                fields = options["attributes_by_category"][case["category"]]
                self.assertEqual(fields[0]["column"], "AE")
                pattern = next(field for field in fields if field["field"] == "product_property/100198")
                self.assertEqual(pattern["label"], "Pattern")
                self.assertEqual(pattern["options"][0], case["pattern"])
                self.assertEqual(pattern["status"], "Optional")
                self.assertFalse(pattern["required"])
                self.assertNotIn(case["forbidden"], {field["field"] for field in fields})
                self.assertEqual(options["base_requirements_by_category"][case["category"]]["price"], "Mandatory")

    def test_custom_english_values_and_forbidden_attributes(self) -> None:
        for case in CASES:
            with self.subTest(template=case["file"]):
                options = self.options[case["file"]]
                pattern = next(field for field in options["attributes_by_category"][case["category"]]
                               if field["field"] == "product_property/100198")
                self.assertTrue(pattern["allow_custom"])
                self.assertEqual(pattern["input_mode"], "select_or_text")
                for value in (case["pattern"], "Custom Indonesian floral pattern"):
                    self.assertEqual(validate_attributes(case["category"], {pattern["field"]: value}, options),
                                     {pattern["field"]: value})
                with self.assertRaisesRegex(ValueError, "所选类目不支持属性"):
                    validate_attributes(case["category"], {case["forbidden"]: "Unsupported value"}, options)

    def test_selection_validates_english_options_in_each_template(self) -> None:
        for case in CASES:
            with self.subTest(template=case["file"]):
                options = self.options[case["file"]]
                season = next(field for field in options["attributes_by_category"][case["season_category"]]
                              if field["field"] == "product_property/100397")
                self.assertEqual(season["label"], "Season")
                self.assertEqual(season["column"], case["season_column"])
                self.assertEqual(season["input_mode"], "select")
                self.assertFalse(season["allow_custom"])
                self.assertEqual(season["options"][:2], ["Spring", "Summer"])
                self.assertEqual(validate_attributes(case["season_category"], {season["field"]: "Spring， Summer"}, options),
                                 {season["field"]: "Spring,Summer"})
                with self.assertRaisesRegex(ValueError, "必须从模板选项中选择"):
                    validate_attributes(case["season_category"], {season["field"]: "Custom season"}, options)

    def test_washing_and_waist_options_follow_their_own_dropdown_references(self) -> None:
        washing_options = ["Hand wash only", "Dry clean", "Do not wash", "Do not dry clean", "Machine washable", "Cold wash"]
        waist_options = ["Low waist", "Mid waist", "High waist"]
        for filename, category, expected in (
            ("menswear_20261026.xlsx", "Men's Bottoms/Men's Shorts", {
                "product_property/100401": ("Washing instructions", washing_options),
                "product_property/100403": ("Waist height", waist_options),
            }),
            ("womenswear_20261026.xlsx", "Women's Bottoms/Skirts", {
                "product_property/100403": ("Waist height", waist_options),
            }),
        ):
            with self.subTest(template=filename):
                options = self.options[filename]
                fields = {field["field"]: field for field in options["attributes_by_category"][category]}
                for field_name, (label, values) in expected.items():
                    self.assertEqual(fields[field_name]["label"], label)
                    self.assertEqual(fields[field_name]["input_mode"], "select")
                    self.assertEqual(fields[field_name]["options"], values)
                    self.assertEqual(validate_attributes(category, {field_name: values[0]}, options), {field_name: values[0]})
                if "product_property/100401" in expected:
                    with self.assertRaisesRegex(ValueError, "必须从模板选项中选择"):
                        validate_attributes(category, {"product_property/100401": "Low waist"}, options)

        for filename, options in self.options.items():
            for category, fields in options["attributes_by_category"].items():
                for field in fields:
                    with self.subTest(template=filename, category=category, field=field["field"]):
                        if field["field"] == "product_property/100401":
                            self.assertCountEqual(field["options"], washing_options)
                        elif field["field"] == "product_property/100403":
                            self.assertCountEqual(field["options"], waist_options)

    def test_export_maps_variants_and_preserves_original_package(self) -> None:
        template = SimpleNamespace(
            package_weight=0.28, package_length=30, package_width=16, package_height=2,
            sku_specifications={"size": {"options": ["S", "M"]}},
        )
        product = {
            "title": "Indonesian English template regression product", "description": "English product description",
            "image_urls": ["https://img.example/one.jpg", "https://img.example/two.jpg"],
            "base_sku_by_image": {"https://img.example/one.jpg": "ID001", "https://img.example/two.jpg": "ID002"},
            "price": 199000, "quantity": 12, "size_chart_url": "https://img.example/sizes.jpg",
        }
        for case in CASES:
            with self.subTest(template=case["file"]):
                content = self.templates[case["file"]]
                options = self.options[case["file"]]
                fields = options["attributes_by_category"][case["category"]]
                attributes = {field["field"]: field["options"][0] if field["options"] else "Manual English value"
                              for field in fields}
                attributes["product_property/100198"] = "Custom Indonesian floral pattern"
                attributes = validate_attributes(case["category"], attributes, options)
                result = build_workbook(
                    template=template, category=case["category"], cod="Y", attributes=attributes,
                    products=[product], template_bytes=content, template_type="tiktok_local",
                )
                workbook = load_workbook(BytesIO(result))
                original = load_workbook(BytesIO(content))
                try:
                    sheet = workbook["Template"]
                    columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
                    for header_row in original["Template"].iter_rows(min_row=1, max_row=6):
                        for cell in header_row:
                            # openpyxl 保存时会将原模板中的空字符串规范化为空单元格。
                            expected_header = None if cell.value == "" else cell.value
                            self.assertEqual(sheet[cell.coordinate].value, expected_header, cell.coordinate)
                    expected = {
                        "category": case["category"], "product_name": product["title"],
                        "product_description": product["description"], "main_image": product["image_urls"][0],
                        "image_2": product["image_urls"][1], "property_name_1": "Color", "property_name_2": "Size",
                        "parcel_weight": 280, "parcel_length": 30, "parcel_width": 16, "parcel_height": 2,
                        "price": 199000, "quantity": 12, "cod": "Y", "size_chart": product["size_chart_url"],
                        "delivery": None, "pre_order_time": None, "minimum_order_quantity": None, **attributes,
                    }
                    for row, sku, size, color, image in (
                        (7, "ID001-S", "S", "Color 1", product["image_urls"][0]),
                        (8, "ID001-M", "M", "Color 1", product["image_urls"][0]),
                        (9, "ID002-S", "S", "Color 2", product["image_urls"][1]),
                        (10, "ID002-M", "M", "Color 2", product["image_urls"][1]),
                    ):
                        values = {**expected, "seller_sku": sku, "property_value_2": size,
                                  "property_value_1": color, "property_1_image": image}
                        for field, value in values.items():
                            self.assertEqual(sheet.cell(row, columns[field]).value, value, (row, field))
                    self.assertIsNone(sheet.cell(11, columns["seller_sku"]).value)
                finally:
                    workbook.close()
                    original.close()
                with ZipFile(BytesIO(content)) as source, ZipFile(BytesIO(result)) as generated:
                    self.assertEqual(source.namelist(), generated.namelist())
                    # 三份原始样本的第一张工作表均为 Template，其余包内容必须逐字节保留。
                    changed_parts = [name for name in source.namelist() if source.read(name) != generated.read(name)]
                    self.assertEqual(changed_parts, ["xl/worksheets/sheet1.xml"])


class TiktokIndonesiaCatalogTests(unittest.TestCase):
    def setUp(self) -> None:
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.session_factory = sessionmaker(bind=self.engine)
        with self.session_factory() as db:
            db.add_all([
                Company(id=1, name="Indonesia template test company"),
                User(id=1, company_id=1, email="admin@example.com", name="Admin",
                     password_hash="x", role=Role.COMPANY_ADMIN),
            ])
            db.commit()

    def tearDown(self) -> None:
        self.engine.dispose()

    def test_upload_read_and_change_input_mode_preserves_english_metadata(self) -> None:
        with self.session_factory() as db:
            user = db.get(User, 1)
            for case in CASES:
                with self.subTest(template=case["file"]):
                    content = (FIXTURE_DIRECTORY / case["file"]).read_bytes()
                    result = asyncio.run(main.create_tiktok_category_catalog(
                        case["name"], UploadFile(filename=case["file"], file=BytesIO(content)),
                        "tiktok_local", user=user, db=db,
                    ))
                    self.assertEqual(result["name"], case["name"])
                    self.assertEqual(result["source_filename"], case["file"])
                    self.assertEqual(result["template_type"], "tiktok_local")
                    self.assertEqual(result["template_version"], "V5.0.2")
                    self.assertEqual(result["category_count"], case["count"])
                    db.expire_all()
                    catalog = db.get(TiktokCategoryCatalog, result["id"])
                    self.assertEqual(catalog.template_blob, content)
                    self.assertEqual(catalog.parsed_options, result["options"])
                    options = main.get_tiktok_export_options(catalog.id, user=user, db=db)
                    self.assertEqual(options["attributes_by_category"], result["options"]["attributes_by_category"])
                    pattern = next(field for field in options["attributes_by_category"][case["category"]]
                                   if field["field"] == "product_property/100198")
                    main.update_tiktok_category_catalog(catalog.id, TiktokCategoryCatalogUpdate(attribute_input_modes=[{
                        "category": case["category"], "field": pattern["field"], "input_mode": "select",
                    }]), user=user, db=db)
                    db.expire_all()
                    updated = main.get_tiktok_export_options(catalog.id, user=user, db=db)
                    updated_pattern = next(field for field in updated["attributes_by_category"][case["category"]]
                                           if field["field"] == pattern["field"])
                    self.assertEqual(updated_pattern, {**pattern, "input_mode": "select", "input_mode_overridden": True})
                    self.assertEqual(db.get(TiktokCategoryCatalog, catalog.id).template_blob, content)
                    with self.assertRaisesRegex(ValueError, "必须从模板选项中选择"):
                        validate_attributes(case["category"], {pattern["field"]: "Custom pattern"}, updated)
            self.assertEqual(len(main.list_tiktok_category_catalogs(user=user, db=db)), 3)

    def legacy_catalog(self, db):
        """复现已导入类目库：洗涤选项串入腰高、腰高被误判为手动输入。"""
        content = (FIXTURE_DIRECTORY / "menswear_20261026.xlsx").read_bytes()
        options = parse_listing_options(content, "tiktok_local")
        options.pop("parser_version")
        fields = {field["field"]: field for field in options["attributes_by_category"]["Men's Bottoms/Men's Shorts"]}
        fields["product_property/100401"].update(options=["Low waist", "Mid waist", "High waist"], input_mode="select")
        fields["product_property/100403"].update(options=[], input_mode="text")
        # 模拟管理员曾将默认允许自定义的图案改成仅选择。
        fields["product_property/100198"]["input_mode"] = "select"
        catalog = TiktokCategoryCatalog(
            company_id=1, name="旧印尼男装", source_filename="menswear.xlsx", template_type="tiktok_local",
            template_blob=content, parsed_options=options, created_by=1,
        )
        db.add(catalog)
        db.commit()
        return catalog

    def test_legacy_catalog_options_are_corrected_without_reupload_or_read_time_writes(self) -> None:
        from app import supply_catalog

        with self.session_factory() as db:
            catalog = self.legacy_catalog(db)
            previous = deepcopy(catalog.parsed_options)
            user = db.get(User, 1)
            options = main.get_tiktok_export_options(catalog.id, user=user, db=db)
            fields = {field["field"]: field for field in options["attributes_by_category"]["Men's Bottoms/Men's Shorts"]}
            self.assertEqual(options["parser_version"], OPTIONS_PARSER_VERSION)
            self.assertEqual(fields["product_property/100401"]["options"][0], "Hand wash only")
            self.assertEqual(fields["product_property/100403"]["options"], ["Low waist", "Mid waist", "High waist"])
            self.assertEqual(fields["product_property/100403"]["input_mode"], "select")
            self.assertEqual(fields["product_property/100198"]["input_mode"], "select")
            self.assertEqual(main.serialize_tiktok_catalog(catalog, include_options=True)["options"]["attributes_by_category"],
                             options["attributes_by_category"])
            self.assertEqual(supply_catalog.catalog_options(catalog.id, user=user, db=db)["attributes_by_category"],
                             options["attributes_by_category"])
            self.assertFalse(db.dirty)
            db.expire_all()
            self.assertEqual(db.get(TiktokCategoryCatalog, catalog.id).parsed_options, previous)

    def test_editing_legacy_catalog_saves_corrected_options_and_explicit_input_modes(self) -> None:
        with self.session_factory() as db:
            catalog = self.legacy_catalog(db)
            catalog_id = catalog.id
            user = db.get(User, 1)
            main.update_tiktok_category_catalog(catalog_id, TiktokCategoryCatalogUpdate(attribute_input_modes=[{
                "category": "Men's Bottoms/Men's Shorts", "field": "product_property/100401", "input_mode": "text",
            }]), user=user, db=db)
            db.expire_all()
            catalog = db.get(TiktokCategoryCatalog, catalog_id)
            self.assertEqual(catalog.parsed_options["parser_version"], OPTIONS_PARSER_VERSION)
            # 再次触发兼容解析时，也要保留管理员明确选择的手动输入方式。
            previous = deepcopy(catalog.parsed_options)
            previous.pop("parser_version")
            catalog.parsed_options = previous
            db.commit()
            options = main.get_tiktok_export_options(catalog_id, user=user, db=db)
            washing = next(field for field in options["attributes_by_category"]["Men's Bottoms/Men's Shorts"]
                           if field["field"] == "product_property/100401")
            self.assertEqual(washing["options"][0], "Hand wash only")
            self.assertEqual(washing["input_mode"], "text")
            self.assertTrue(washing["input_mode_overridden"])

    def test_export_and_supply_validation_use_corrected_legacy_options(self) -> None:
        from app.supply_common import validate_strategy

        with self.session_factory() as db:
            catalog = self.legacy_catalog(db)
            user = db.get(User, 1)
            db.add_all([
                ProductTemplate(id=1, company_id=1, name="ID", product_description="English product description",
                    supply_public=True, supply_fallback_title="Printed Long Sleeve Menswear Shirt",
                    size_chart_url="https://img.example/sizes.jpg", package_weight=.28,
                    package_length=30, package_width=16, package_height=2,
                    sku_specifications={"size": {"options": ["S"]}}),
                ProductDraft(id=1, company_id=1, template_id=1, title="Indonesian menswear export regression product",
                    image_urls=["https://img.example/one.jpg"],
                    sku_items=[{"image_url": "https://img.example/one.jpg", "sku": "ID001"}], created_by=1, updated_by=1),
            ])
            db.commit()
            attributes = {"product_property/100401": "Hand wash only", "product_property/100403": "Low waist"}
            parameters = {"template_id": 1, "category_catalog_id": catalog.id, "category_name": "Men's Bottoms/Men's Shorts",
                          "attributes": dict(attributes), "price": "199000", "cod": "Y", "sku_count": 1, "link_count": 1}
            validate_strategy(db, user, parameters)
            self.assertEqual(parameters["attributes"], attributes)
            payload = TiktokDraftExportInput(draft_ids=[1], category_catalog_id=catalog.id,
                category="Men's Bottoms/Men's Shorts", default_price=199000, default_quantity=12, cod="Y", attributes=attributes)
            with self.assertRaisesRegex(HTTPException, "必须从模板选项中选择"):
                main.export_drafts_to_tiktok(payload.model_copy(update={"attributes": {"product_property/100401": "Low waist"}}),
                                            user=user, db=db)
            with patch("app.main.resolve_export_images"):
                result = main.export_drafts_to_tiktok(payload, user=user, db=db)
            workbook = load_workbook(BytesIO(result.body))
            try:
                self.assertEqual(workbook["Template"]["AM7"].value, "Hand wash only")
                self.assertEqual(workbook["Template"]["AN7"].value, "Low waist")
            finally:
                workbook.close()

    def test_cross_border_upload_rejects_indonesia_local_templates(self) -> None:
        with self.session_factory() as db:
            for case in CASES:
                with self.subTest(template=case["file"]):
                    upload = UploadFile(filename=case["file"], file=BytesIO((FIXTURE_DIRECTORY / case["file"]).read_bytes()))
                    with self.assertRaisesRegex(HTTPException, "不匹配") as error:
                        asyncio.run(main.create_tiktok_category_catalog(
                            case["name"], upload, "tiktok_cross_border", user=db.get(User, 1), db=db,
                        ))
                    self.assertEqual(error.exception.status_code, 400)
                    self.assertEqual(db.query(TiktokCategoryCatalog).count(), 0)


if __name__ == "__main__":
    unittest.main()
