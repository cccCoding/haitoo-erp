import asyncio
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import UploadFile
from openpyxl import load_workbook
from pydantic import ValidationError

from app import main
from app.models import Company, HubUploadTask, ProductDraft, ProductTemplate, TiktokCategoryCatalog, User
from app.schemas import HubUploadTaskCreate, TiktokDraftExportInput
from app.supply_catalog import StrategyInput
from app.tiktok_export import OPTIONS_PARSER_VERSION, build_workbook, catalog_listing_options, parse_listing_options
from tests import test_tiktok_indonesia_templates as indonesia


FIXTURES = Path(__file__).parent / "fixtures" / "tiktok" / "indonesia"
CASES = [("bothplatforms_" + case["file"], case) for case in indonesia.CASES]


class BothplatformsTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.templates = {filename: (FIXTURES / filename).read_bytes() for filename, _ in CASES}

    def export(self, content, insurance=None):
        template = SimpleNamespace(package_weight=.28, package_length=30, package_width=16, package_height=2,
                                   sku_specifications={"size": {"options": ["S", "M"]}})
        category = parse_listing_options(content, "tiktok_local")["categories"][0]["name"]
        return build_workbook(template=template, category=category, cod="Y",
            attributes={"product_property/100198": "Embroidery"}, products=[{
                "title": "Both platforms English export regression product", "description": "Product description",
                "image_urls": ["https://img.example/one.jpg", "https://img.example/two.jpg"],
                "base_sku_by_image": {"https://img.example/one.jpg": "ID001", "https://img.example/two.jpg": "ID002"},
                "price": 199000, "quantity": 12, "size_chart_url": "https://img.example/sizes.jpg",
            }], template_bytes=content, template_type="tiktok_local", shipping_insurance=insurance)

    def test_all_three_templates_expose_insurance_and_keep_correct_attribute_options(self):
        for filename, case in CASES:
            with self.subTest(template=filename):
                options = parse_listing_options(self.templates[filename], "tiktok_local")
                self.assertEqual(options["template_type"], "tiktok_local")
                self.assertEqual(len(options["categories"]), case["count"])
                self.assertTrue(options["capabilities"]["supports_shipping_insurance"])
                self.assertEqual(options["shipping_insurance_options"], ["Optional", "Mandatory"])
                self.assertEqual(options["shipping_insurance_default"], "Optional")
                self.assertEqual(options["attributes_by_category"][case["category"]][0]["column"], "AF")
                for category, fields in options["attributes_by_category"].items():
                    self.assertEqual(options["base_requirements_by_category"][category]["shipping_insurance"], "Mandatory")
                    for field in fields:
                        if field["field"] == "product_property/100401":
                            self.assertIn("Hand wash only", field["options"])
                            self.assertNotIn("Low waist", field["options"])
                        if field["field"] == "product_property/100403":
                            self.assertCountEqual(field["options"], ["Low waist", "Mid waist", "High waist"])

    def test_every_sku_gets_default_or_selected_insurance_and_package_is_preserved(self):
        for filename, _ in CASES:
            for insurance in (None, "Mandatory"):
                with self.subTest(template=filename, insurance=insurance):
                    content = self.templates[filename]
                    output = self.export(content, insurance)
                    workbook = load_workbook(BytesIO(output))
                    try:
                        sheet = workbook["Template"]
                        columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
                        for row, sku in enumerate(["ID001-S", "ID001-M", "ID002-S", "ID002-M"], 7):
                            self.assertEqual(sheet.cell(row, columns["shipping_insurance"]).value, insurance or "Optional")
                            self.assertEqual(sheet.cell(row, columns["seller_sku"]).value, sku)
                            self.assertEqual(sheet.cell(row, columns["product_property/100198"]).value, "Embroidery")
                            self.assertEqual(sheet.cell(row, columns["price"]).value, 199000)
                            self.assertEqual(sheet.cell(row, columns["quantity"]).value, 12)
                        self.assertEqual(sheet["H2"].value, "TikTok Shop and Tokopedia (0;1)")
                    finally:
                        workbook.close()
                    with ZipFile(BytesIO(content)) as source, ZipFile(BytesIO(output)) as generated:
                        self.assertEqual(source.namelist(), generated.namelist())
                        self.assertEqual([name for name in source.namelist() if source.read(name) != generated.read(name)],
                                         ["xl/worksheets/sheet1.xml"])

    def test_invalid_insurance_is_rejected_by_api_inputs_and_workbook_generator(self):
        for model, payload in (
            (TiktokDraftExportInput, {"draft_ids": [1], "category_catalog_id": 1, "category": "Clothing", "default_price": 10}),
            (HubUploadTaskCreate, {"draft_ids": [1], "category_catalog_id": 1, "category": "Clothing", "default_price": 10}),
            (StrategyInput, {"name": "Strategy", "category_catalog_id": 1, "category_name": "Clothing", "template_id": 1,
                             "price": "10", "sku_count": 5, "link_count": 10}),
        ):
            with self.subTest(model=model.__name__), self.assertRaises(ValidationError):
                model(**payload, shipping_insurance="Invalid")
        with self.assertRaisesRegex(ValueError, "运输保险必须从模板选项中选择"):
            self.export(self.templates[CASES[0][0]], "Invalid")

    def test_missing_insurance_sheet_or_default_is_rejected(self):
        for missing_sheet in (True, False):
            with self.subTest(missing_sheet=missing_sheet):
                workbook = load_workbook(BytesIO(self.templates[CASES[0][0]]))
                if missing_sheet:
                    del workbook["ShippingInsurance"]
                else:
                    workbook["ShippingInsurance"]["A1"] = "Unsupported"
                output = BytesIO()
                workbook.save(output)
                workbook.close()
                with self.assertRaisesRegex(ValueError, "ShippingInsurance|Optional"):
                    parse_listing_options(output.getvalue(), "tiktok_local")

    def test_old_parsed_catalogs_refresh_insurance_and_preserve_admin_input_modes(self):
        for filename, case in CASES:
            with self.subTest(template=filename):
                content = self.templates[filename]
                previous = parse_listing_options(content, "tiktok_local")
                previous["parser_version"] = 2
                previous["capabilities"].pop("supports_shipping_insurance")
                previous.pop("shipping_insurance_options")
                previous.pop("shipping_insurance_default")
                field = next(field for field in previous["attributes_by_category"][case["category"]]
                             if field["field"] == "product_property/100198")
                field.update(input_mode="text", input_mode_overridden=True)
                original = deepcopy(previous)
                options = catalog_listing_options(SimpleNamespace(template_type="tiktok_local", template_blob=content, parsed_options=previous))
                self.assertEqual(options["parser_version"], OPTIONS_PARSER_VERSION)
                self.assertEqual(options["shipping_insurance_default"], "Optional")
                restored = next(field for field in options["attributes_by_category"][case["category"]]
                                if field["field"] == "product_property/100198")
                self.assertEqual(restored["input_mode"], "text")
                self.assertEqual(previous, original)

    def test_single_platform_templates_do_not_enable_or_write_insurance(self):
        for _, case in CASES:
            with self.subTest(template=case["file"]):
                content = (FIXTURES / case["file"]).read_bytes()
                options = parse_listing_options(content, "tiktok_local")
                self.assertFalse(options["capabilities"]["supports_shipping_insurance"])
                self.assertEqual(options["shipping_insurance_options"], [])
                self.assertIsNone(options["shipping_insurance_default"])
                workbook = load_workbook(BytesIO(self.export(content, "Mandatory")))
                try:
                    self.assertNotIn("shipping_insurance", [cell.value for cell in workbook["Template"][1]])
                finally:
                    workbook.close()


class BothplatformsCatalogFlowTests(unittest.TestCase):
    setUp = indonesia.TiktokIndonesiaCatalogTests.setUp
    tearDown = indonesia.TiktokIndonesiaCatalogTests.tearDown

    def test_upload_download_and_hub_file_include_insurance(self):
        with self.session_factory() as db:
            user = db.get(User, 1)
            company = db.get(Company, 1)
            company.hubstudio_app_id = "test-app"
            company.hubstudio_secret_encrypted = "configured"
            company.hubstudio_group_code = "test-group"
            db.add_all([
                ProductTemplate(id=1, company_id=1, name="ID", product_description="English product description",
                    package_weight=.28, package_length=30, package_width=16, package_height=2,
                    size_chart_url="https://img.example/sizes.jpg", sku_specifications={"size": {"options": ["S", "M"]}}),
                ProductDraft(id=1, company_id=1, template_id=1, title="Both platforms menswear export regression product",
                    image_urls=["https://img.example/one.jpg"],
                    sku_items=[{"image_url": "https://img.example/one.jpg", "sku": "ID001"}], created_by=1, updated_by=1),
            ])
            db.commit()
            for filename, case in CASES:
                with self.subTest(template=filename):
                    content = (FIXTURES / filename).read_bytes()
                    result = asyncio.run(main.create_tiktok_category_catalog(
                        "双平台" + case["name"], UploadFile(filename=filename, file=BytesIO(content)), "tiktok_local", user=user, db=db))
                    self.assertEqual(db.get(TiktokCategoryCatalog, result["id"]).template_blob, content)
                    self.assertEqual(main.get_tiktok_export_options(result["id"], user=user, db=db)["shipping_insurance_default"], "Optional")
                    payload = TiktokDraftExportInput(draft_ids=[1], category_catalog_id=result["id"],
                        category=case["category"], default_price=199000, shipping_insurance=None)
                    with patch("app.main.resolve_export_images"):
                        response = main.export_drafts_to_tiktok(payload, user=user, db=db)
                        with patch("app.main.upload_hub_export", return_value="https://r2.invalid/private/test.xlsx") as upload:
                            task = main.create_hubstudio_upload_task(
                                HubUploadTaskCreate(**{**payload.model_dump(), "shipping_insurance": "Mandatory"}), user=user, db=db)
                    self.assertEqual(db.get(HubUploadTask, task["id"]).parameters["shipping_insurance"], "Mandatory")
                    for data, expected in ((response.body, "Optional"), (upload.call_args.args[0], "Mandatory")):
                        workbook = load_workbook(BytesIO(data))
                        try:
                            self.assertEqual(workbook["Template"]["AE7"].value, expected)
                            self.assertEqual(workbook["Template"]["AE8"].value, expected)
                            self.assertEqual(workbook["Template"]["H2"].value, "TikTok Shop and Tokopedia (0;1)")
                        finally:
                            workbook.close()
