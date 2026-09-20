from io import BytesIO
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import main
from app.database import Base
from app.models import ProductDraft, ProductTemplate, Role, TiktokCategoryCatalog, User
from app.schemas import TiktokDraftExportInput
from app.tiktok_export import TEMPLATE_PATH, build_workbook, listing_options, parse_listing_options, validate_attributes


CROSS_BORDER_TEMPLATE_PATH = TEMPLATE_PATH.parent / "tiktok_seller_cross_border_zh.xlsx"


class TiktokExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.session_factory = sessionmaker(bind=self.engine)
        with self.session_factory() as db:
            options = listing_options()
            db.add_all([
                TiktokCategoryCatalog(id=1, company_id=1, name="穆斯林服装", source_filename=TEMPLATE_PATH.name, template_version=options["template_version"], template_blob=TEMPLATE_PATH.read_bytes(), parsed_options=options, created_by=1),
                ProductTemplate(id=1, company_id=1, name="Y1", product_description="Template description", size_chart_url="https://img.example/size-chart.jpg", package_weight=0.28, package_length=30, package_width=16, package_height=2, sku_specifications={"size": {"name": "尺码", "options": ["S", "M"]}}),
                ProductTemplate(id=2, company_id=1, name="Y2", product_description="Other template", size_chart_url="https://img.example/size-chart-2.jpg", package_weight=0.3, package_length=31, package_width=17, package_height=3, sku_specifications={"size": {"name": "尺码", "options": ["M"]}}),
                User(id=1, company_id=1, email="operator@example.com", name="Operator", user_code="AA", password_hash="x", role=Role.MEMBER),
                User(id=2, company_id=1, email="other@example.com", name="Other", user_code="BB", password_hash="x", role=Role.MEMBER),
                ProductDraft(id=1, company_id=1, template_id=1, title="First unique TikTok product title", product_description="Draft description", size_chart_url="https://img.example/size-chart.jpg", image_urls=["https://img.example/one.jpg", "https://img.example/two.jpg"], sku_items=[{"image_url": "https://img.example/one.jpg", "size": None, "sku": "Y1AA000001"}, {"image_url": "https://img.example/two.jpg", "size": None, "sku": "Y1AA000002"}], created_by=1, updated_by=1),
                ProductDraft(id=2, company_id=1, template_id=1, title="Second unique TikTok product title", image_urls=["https://img.example/three.jpg"], sku_items=[{"image_url": "https://img.example/three.jpg", "size": None, "sku": "Y1AA000003"}], created_by=1, updated_by=1),
                ProductDraft(id=3, company_id=1, template_id=2, title="Third unique TikTok product title", image_urls=["https://img.example/four.jpg"], sku_items=[{"image_url": "https://img.example/four.jpg", "size": None, "sku": "Y2AA000004"}], created_by=1, updated_by=1),
                ProductDraft(id=4, company_id=1, template_id=1, title="Other user unique TikTok product title", image_urls=["https://img.example/five.jpg"], sku_items=[{"image_url": "https://img.example/five.jpg", "size": None, "sku": "Y1BB000005"}], created_by=2, updated_by=2),
            ])
            db.commit()

    def tearDown(self) -> None:
        self.engine.dispose()

    def payload(self) -> TiktokDraftExportInput:
        return TiktokDraftExportInput(draft_ids=[1, 2], category_catalog_id=1, category="女士上装/女士衬衫", default_price=10, default_quantity=999, cod="Y", attributes={"product_property/100198": "花朵"}, product_overrides=[{"draft_id": 1, "price": 12.5, "quantity": 88}])

    def test_export_request_accepts_at_most_twenty_drafts(self) -> None:
        values = self.payload().model_dump()
        values["draft_ids"] = list(range(1, 22))
        with self.assertRaisesRegex(ValueError, "at most 20"):
            TiktokDraftExportInput(**values)

    def test_options_are_read_from_bundled_template(self) -> None:
        options = listing_options()
        self.assertEqual(options["template_version"], "V5.0.2")
        self.assertEqual(len(options["categories"]), 41)
        blouse_fields = {item["field"]: item for item in options["attributes_by_category"]["女士上装/女士衬衫"]}
        self.assertIn("product_property/100198", blouse_fields)
        self.assertTrue(blouse_fields["product_property/100198"]["allow_custom"])
        self.assertEqual(blouse_fields["product_property/100198"]["input_mode"], "select_or_text")
        season = next(item for item in blouse_fields.values() if item["label"] == "季节")
        self.assertEqual(season["input_mode"], "select")
        self.assertEqual(options["base_requirements_by_category"]["女士上装/女士衬衫"]["size_chart"], "Mandatory")

    def test_cross_border_template_uses_its_own_adapter(self) -> None:
        options = parse_listing_options(CROSS_BORDER_TEMPLATE_PATH.read_bytes(), "tiktok_cross_border")
        self.assertEqual(options["template_type"], "tiktok_cross_border")
        self.assertEqual(options["template_version"], "V5.0.2")
        self.assertFalse(options["capabilities"]["supports_cod"])
        self.assertFalse(options["capabilities"]["supports_hubstudio_submit"])
        self.assertEqual(len(options["categories"]), 41)
        with self.assertRaisesRegex(ValueError, "不匹配"):
            parse_listing_options(CROSS_BORDER_TEMPLATE_PATH.read_bytes(), "tiktok_local")

    def test_select_accepts_supported_values_and_custom_input_rejects_lists(self) -> None:
        options = listing_options()
        blouse_fields = options["attributes_by_category"]["女士上装/女士衬衫"]
        season = next(item for item in blouse_fields if item["label"] == "季节")
        values = season["options"][:2]
        self.assertEqual(validate_attributes(
            "女士上装/女士衬衫", {season["field"]: f"{values[0]}， {values[1]}"}, options,
        )[season["field"]], ",".join(values))
        pattern = next(item for item in blouse_fields if item["label"] == "图案花纹")
        with self.assertRaisesRegex(ValueError, "不支持填写多个值"):
            validate_attributes("女士上装/女士衬衫", {pattern["field"]: pattern["options"][:2]}, options)

    def test_select_or_text_accepts_supported_or_custom_value(self) -> None:
        options = listing_options()
        pattern = next(item for item in options["attributes_by_category"]["女士上装/女士衬衫"] if item["label"] == "图案花纹")
        self.assertEqual(validate_attributes(
            "女士上装/女士衬衫", {pattern["field"]: pattern["options"][0]}, options,
        )[pattern["field"]], pattern["options"][0])
        self.assertEqual(validate_attributes(
            "女士上装/女士衬衫", {pattern["field"]: "My custom pattern"}, options,
        )[pattern["field"]], "My custom pattern")

    def test_export_preserves_template_maps_drafts_and_counts_exports(self) -> None:
        source = load_workbook(TEMPLATE_PATH, data_only=False)
        with self.session_factory() as db:
            response = main.export_drafts_to_tiktok(self.payload(), user=db.get(User, 1), db=db)
            self.assertEqual(db.get(ProductDraft, 1).export_count, 1)
            self.assertEqual(db.get(ProductDraft, 2).export_count, 1)
        exported = load_workbook(BytesIO(response.body), data_only=False)
        self.assertEqual(exported.sheetnames, source.sheetnames)
        self.assertEqual([sheet.sheet_state for sheet in exported], [sheet.sheet_state for sheet in source])
        self.assertEqual(exported["TemplateConfig"]["A2"].value, source["TemplateConfig"]["A2"].value)
        self.assertEqual(len(exported["Template"].data_validations.dataValidation), len(source["Template"].data_validations.dataValidation))
        self.assertEqual(len(exported["Template"].conditional_formatting), len(source["Template"].conditional_formatting))
        with ZipFile(TEMPLATE_PATH) as source_zip, ZipFile(BytesIO(response.body)) as exported_zip:
            source_media = {name for name in source_zip.namelist() if name.startswith("xl/media/")}
            self.assertTrue(source_media)
            self.assertTrue(source_media.issubset(exported_zip.namelist()))
            self.assertEqual(exported_zip.read("xl/worksheets/sheet3.xml"), source_zip.read("xl/worksheets/sheet3.xml"))
        sheet = exported["Template"]
        self.assertEqual([sheet.cell(row, 26).value for row in range(7, 13)], ["Y1AA000001-S", "Y1AA000001-M", "Y1AA000002-S", "Y1AA000002-M", "Y1AA000003-S", "Y1AA000003-M"])
        self.assertEqual(sheet["E7"].value, "https://img.example/one.jpg")
        self.assertEqual(sheet["F7"].value, "https://img.example/two.jpg")
        self.assertEqual(sheet["P9"].value, "https://img.example/two.jpg")
        self.assertEqual(sheet["N7"].value, "Color")
        self.assertEqual(sheet["Q7"].value, "Size")
        self.assertEqual(sheet["D7"].value, "Draft description")
        self.assertEqual(sheet["D11"].value, "Template description")
        self.assertEqual([sheet.cell(row, 24).value for row in range(7, 13)], [12.5, 12.5, 12.5, 12.5, 10, 10])
        self.assertEqual([sheet.cell(row, 25).value for row in range(7, 13)], [88, 88, 88, 88, 999, 999])
        self.assertTrue(all(sheet.cell(row, 28).value == "Y" for row in range(7, 13)))
        self.assertTrue(all(sheet.cell(row, 30).value == "花朵" for row in range(7, 13)))
        self.assertIsNone(sheet["A13"].value)

    def test_repeated_successful_export_increments_once_per_draft(self) -> None:
        with self.session_factory() as db:
            user = db.get(User, 1)
            main.export_drafts_to_tiktok(self.payload(), user=user, db=db)
            main.export_drafts_to_tiktok(self.payload(), user=user, db=db)
            self.assertEqual(db.get(ProductDraft, 1).export_count, 2)
            self.assertEqual(db.get(ProductDraft, 2).export_count, 2)

    def test_generated_main_and_carousel_are_gallery_while_sku_images_stay_variants(self) -> None:
        payload = self.payload().model_copy(update={"draft_ids": [1], "product_overrides": []})
        with self.session_factory() as db:
            draft = db.get(ProductDraft, 1)
            draft.carousel_items = [
                {"sku": None, "image_url": "https://img.example/main.jpg", "task_id": 11, "source_type": "main_image"},
                {"sku": "Y1AA000001", "image_url": "https://img.example/carousel.jpg", "task_id": 10},
            ]
            draft.image_urls = [item["image_url"] for item in draft.carousel_items]
            db.commit()
            response = main.export_drafts_to_tiktok(payload, user=db.get(User, 1), db=db)
        sheet = load_workbook(BytesIO(response.body), data_only=False)["Template"]
        self.assertEqual(sheet["E7"].value, "https://img.example/main.jpg")
        self.assertEqual(sheet["F7"].value, "https://img.example/carousel.jpg")
        self.assertEqual(sheet["P7"].value, "https://img.example/one.jpg")
        self.assertEqual(sheet["P9"].value, "https://img.example/two.jpg")

    def test_cross_border_workbook_leaves_auction_fields_empty(self) -> None:
        with self.session_factory() as db:
            draft = db.get(ProductDraft, 1)
            template = db.get(ProductTemplate, 1)
            workbook_bytes = build_workbook(
                template=template, category="女士上装/女士衬衫", cod="Y", attributes={},
                products=[{
                    "title": draft.title, "description": draft.product_description,
                    "image_urls": draft.image_urls, "sku_images": [{"image_url": "https://img.example/one.jpg", "sku": "Y1AA000001"}],
                    "size_chart_url": draft.size_chart_url, "price": 10, "quantity": 2,
                }], template_bytes=CROSS_BORDER_TEMPLATE_PATH.read_bytes(), template_type="tiktok_cross_border",
            )
        sheet = load_workbook(BytesIO(workbook_bytes), data_only=False)["Template"]
        self.assertEqual(sheet["N7"].value, "Color")
        self.assertEqual(sheet["Q7"].value, "Size")
        self.assertEqual(sheet["W7"].value, 10)
        self.assertIsNone(sheet["AA7"].value)
        self.assertIsNone(sheet["AB7"].value)

    def test_validation_and_generation_failures_do_not_increment(self) -> None:
        with self.session_factory() as db:
            user = db.get(User, 1)
            mixed = self.payload().model_copy(update={"draft_ids": [1, 3]})
            with self.assertRaisesRegex(HTTPException, "同一产品模板"):
                main.export_drafts_to_tiktok(mixed, user=user, db=db)
            with patch("app.main.build_tiktok_workbook", side_effect=ValueError("生成失败")), self.assertRaisesRegex(HTTPException, "生成失败"):
                main.export_drafts_to_tiktok(self.payload(), user=user, db=db)
            self.assertEqual(db.get(ProductDraft, 1).export_count, 0)

    def test_duplicate_product_names_are_rejected_without_incrementing_exports(self) -> None:
        with self.session_factory() as db:
            draft = db.get(ProductDraft, 2)
            draft.title = "  FIRST unique tiktok product title  "
            db.commit()
            with self.assertRaisesRegex(HTTPException, "商品名称必须不同"):
                main.export_drafts_to_tiktok(self.payload(), user=db.get(User, 1), db=db)
            self.assertEqual(db.get(ProductDraft, 1).export_count, 0)
            self.assertEqual(db.get(ProductDraft, 2).export_count, 0)

    def test_member_cannot_export_another_users_draft(self) -> None:
        payload = self.payload().model_copy(update={"draft_ids": [4]})
        with self.session_factory() as db, self.assertRaisesRegex(HTTPException, "无权导出"):
            main.export_drafts_to_tiktok(payload, user=db.get(User, 1), db=db)


if __name__ == "__main__":
    unittest.main()
