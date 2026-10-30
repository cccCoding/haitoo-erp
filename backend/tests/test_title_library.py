"""词源追溯、市场隔离、事实约束、预算和后台恢复的行为测试。"""
import asyncio
import copy
import io
import unittest
from datetime import datetime, timedelta
from unittest.mock import AsyncMock, patch

from alembic import command
from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import create_engine, func, inspect, select

from app import main
from app.ai_providers import ProviderError
from app.credentials import encrypt_secret
from app.models import ProductDraft, ProductTemplate, UserAIProviderCredential
from app.schema_version import alembic_config
from app.title_import import extract_row, parse_file
from app.title_models import TitleEvidence, TitleKeyword, TitleMarket, TitleRecord, TitleTask
from app.title_service import (adopt_record, create_record, digest, generate_record, get_market,
    normalized_title, profile_snapshot, similarity, validate_malay_title)
from app.title_worker import claim_task, run_task
from tests.test_supply_catalog import SupplyFixture


FACTS = {'attributes': [
    {'tag': '品类', 'value': '连衣裙', 'ms': 'gaun', 'evidence': 'manual category', 'source': 'manual'},
    {'tag': '结构设计', 'value': '长袖', 'ms': 'lengan panjang', 'evidence': 'manual sleeve', 'source': 'manual'},
    {'tag': '图案', 'value': '花卉', 'ms': 'corak bunga', 'evidence': 'visible floral print', 'source': 'image'},
    {'tag': '性别', 'value': '女', 'ms': 'wanita', 'evidence': 'confirmed template gender', 'source': 'manual'},
    {'tag': '年龄', 'value': '成人', 'ms': 'dewasa', 'evidence': 'confirmed template age', 'source': 'manual'},
], 'conflicts': []}
TITLE = 'Gaun Lengan Panjang Corak Bunga'
KEYWORDS = '马来语词语,中文释义,产品分类,标签,是否借词\ngaun,连衣裙,连衣裙,品类,否\nlengan panjang,长袖,连衣裙,结构设计,否\ncorak bunga,花卉,连衣裙,图案,否\n'


class TitleLibraryTests(SupplyFixture, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self.app.include_router(main.app.router)
        with self.sessions() as db:
            for uid in (1, 2):
                db.add(UserAIProviderCredential(company_id=1, user_id=uid, provider='deepseek', secret_encrypted=encrypt_secret('test-key')))
            db.commit()

    def enable(self, platform='TikTok'):
        response = self.request('PUT', '/title-market', json={'enabled': True, 'platform': platform})
        self.assertEqual(response.status_code, 200, response.text)

    def profile(self, **extra):
        response = self.request('PUT', '/title-profiles/1', json={'category': '连衣裙',
            'gender': '女', 'age_group': '成人',
            'attributes': {'结构设计': '长袖'}, 'fallback_title': 'Gaun Wanita Lengan Panjang Longgar',
            'fallback_approved': True, **extra})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def upload(self, content=KEYWORDS, filename='keywords.csv', period='2026-10-29', kind='keywords', platform='TikTok'):
        response = self.request('POST', '/title-imports', files={'file': (filename, content.encode() if isinstance(content, str) else content)},
            data={'kind': kind, 'source_name': '运营市场样本', 'period': period, 'platform': platform})
        self.assertEqual(response.status_code, 200, response.text)
        task = response.json()
        self.confirm_new_categories(task)
        return task

    def confirm_new_categories(self, task):
        """这些既有用例明确选择新增；陌生分类拦截由分类专项测试覆盖。"""
        resolution = task.get('category_resolution')
        if task['state'] != 'preview' or not resolution or resolution['ready']:
            return
        choices = [{'source_key': item['source_key'], 'target_id': item['target_id'],
            'new_name': None if item['target_id'] else item['new_name'] or item['source'], 'save_alias': True}
            for item in resolution['items']]
        result = self.request('PUT', f"/title-imports/{task['id']}/categories", json={'token': resolution['token'], 'choices': choices})
        self.assertEqual(result.status_code, 200, result.text)
        task['category_resolution'] = result.json()

    def run_worker(self):
        claimed = claim_task(self.sessions)
        self.assertIsNotNone(claimed)
        asyncio.run(run_task(*claimed, session_factory=self.sessions))

    def import_words(self, **kwargs):
        upload = self.upload(**kwargs)
        response = self.request('POST', f"/title-tasks/{upload['id']}/start", json={})
        self.assertEqual(response.status_code, 200, response.text)
        self.run_worker()
        task = self.request('GET', f"/title-tasks/{upload['id']}").json()
        self.assertEqual(task['state'], 'review', task)
        response = self.request('POST', f"/title-imports/{upload['id']}/confirm", json={
            'selections': [w['index'] for w in task['result']['words']]})
        self.assertEqual(response.status_code, 200, response.text)
        return upload['id']

    def model(self, title=TITLE):
        async def answer(key, instructions, data, **kwargs):
            return copy.deepcopy(FACTS) if 'manual' in data else {'title': title}
        return AsyncMock(side_effect=answer)

    def test_import_confirmation_enables_all_selected_words_even_for_legacy_false_flag(self):
        for index, legacy in enumerate(({}, {'enable': False}, {'enable': True})):
            with self.subTest(legacy=legacy):
                content = KEYWORDS + f'Acme {index},品牌,连衣裙,品牌\nPromosi {index},促销,连衣裙,营销修饰\n'
                job = self.upload(content=content, filename=f'import-{index}.csv', period=f'2026-10-{27+index}')
                self.request('POST', f"/title-tasks/{job['id']}/start", json={})
                self.run_worker()
                task = self.request('GET', f"/title-tasks/{job['id']}").json()
                selected = [word['index'] for word in task['result']['words']]
                response = self.request('POST', f"/title-imports/{job['id']}/confirm", json={'selections': selected, **legacy})
                self.assertEqual(response.status_code, 200, response.text)
                words = self.request('GET', '/title-keywords').json()['items']
                self.assertTrue(words)
                self.assertTrue(all(word['status'] == 'enabled' for word in words))
                brand = next(word for word in words if word['phrase'] == f'Acme {index}')
                disabled = self.request('POST', f"/title-keywords/{brand['id']}/status",
                    json={'status': 'disabled', 'review_token': brand['review_token']})
                self.assertEqual(disabled.status_code, 200, disabled.text)
                # 重放已完成的入库请求不能覆盖后续人工停用，也不能重复累计词源。
                with self.sessions() as db:
                    evidence_count = db.scalar(select(func.count(TitleEvidence.id)))
                    revision = get_market(db, 1).revision
                repeated = self.request('POST', f"/title-imports/{job['id']}/confirm", json={'selections': selected, **legacy})
                self.assertEqual(repeated.json(), response.json())
                with self.sessions() as db:
                    self.assertEqual(db.get(TitleKeyword, brand['id']).status, 'disabled')
                    self.assertEqual(db.scalar(select(func.count(TitleEvidence.id))), evidence_count)
                    self.assertEqual(get_market(db, 1).revision, revision)
                # 下一轮只检查当轮词条，先恢复前一轮的品牌状态。
                current = next(word for word in self.request('GET', '/title-keywords').json()['items'] if word['id'] == brand['id'])
                self.request('POST', f"/title-keywords/{brand['id']}/status",
                    json={'status': 'enabled', 'review_token': current['review_token']})

    def test_confirming_a_new_import_enables_merged_words_and_only_inserts_selected_results(self):
        self.import_words()
        words = self.request('GET', '/title-keywords').json()['items']
        self.request('POST', '/title-keywords/1/status', json={'status': 'disabled', 'review_token': words[0]['review_token']})
        with self.sessions() as db:
            db.get(TitleKeyword, 2).status = 'pending'  # 模拟迁移前已入库但待审核的历史词。
            db.commit()
        job = self.upload(content=KEYWORDS+'dewasa,成人,连衣裙,年龄\n', period='2026-10-30')
        self.request('POST', f"/title-tasks/{job['id']}/start", json={})
        self.run_worker()
        task = self.request('GET', f"/title-tasks/{job['id']}").json()
        selected = [word['index'] for word in task['result']['words'] if word['phrase'] != 'dewasa']
        response = self.request('POST', f"/title-imports/{job['id']}/confirm", json={'selections': selected})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['merged'], 3)
        words = self.request('GET', '/title-keywords').json()['items']
        self.assertEqual(len(words), 3)
        self.assertTrue(all(word['status'] == 'enabled' for word in words))
        self.assertTrue(all(word['evidence_count'] == 2 for word in words))

    def test_import_can_be_confirmed_in_batches_and_reopened_without_replaying_imported_words(self):
        job = self.upload()
        self.request('POST', f"/title-tasks/{job['id']}/start", json={})
        self.run_worker()
        path = f"/title-tasks/{job['id']}"
        words = self.request('GET', path).json()['result']['words']
        endpoint = f"/title-imports/{job['id']}/confirm"
        first = self.request('POST', endpoint, json={'selections': [words[0]['index']]})
        self.assertEqual(first.json(), {'created': 1, 'merged': 0, 'evidence': 1})
        reopened = self.request('GET', path).json()
        self.assertEqual(reopened['state'], 'completed')
        self.assertEqual(reopened['result']['imported_total'], 1)
        self.assertEqual([word['imported'] for word in reopened['result']['words']], [True, False, False])
        keyword = self.request('GET', '/title-keywords').json()['items'][0]
        self.request('POST', f"/title-keywords/{keyword['id']}/status",
            json={'status': 'disabled', 'review_token': keyword['review_token']})
        second = self.request('POST', endpoint, json={'selections': [words[0]['index'], words[1]['index']]})
        self.assertEqual(second.json(), {'created': 2, 'merged': 0, 'evidence': 2})
        with self.sessions() as db:
            revision = get_market(db, 1).revision
            self.assertEqual(db.get(TitleKeyword, keyword['id']).status, 'disabled')
        replay = self.request('POST', endpoint, json={'selections': [words[1]['index']]})
        self.assertEqual(replay.json(), second.json())
        with self.sessions() as db:
            self.assertEqual(get_market(db, 1).revision, revision)
            self.assertEqual(db.scalar(select(func.count(TitleEvidence.id))), 2)
        invalid = self.request('POST', endpoint, json={'selections': [words[2]['index'], 999999]})
        self.assertEqual(invalid.status_code, 400)
        self.assertEqual(self.request('GET', path).json()['result']['imported_total'], 2)
        final = self.request('POST', endpoint, json={'selections': [words[2]['index']]})
        self.assertEqual(final.json(), {'created': 3, 'merged': 0, 'evidence': 3})
        self.assertTrue(all(word['imported'] for word in self.request('GET', path).json()['result']['words']))

    def test_retrying_failed_rows_after_partial_import_keeps_imported_candidates(self):
        job = self.upload()
        self.request('POST', f"/title-tasks/{job['id']}/start", json={})

        async def fail_one_row(row, *args, **kwargs):
            if row['row'] == 3:
                raise ValueError('临时提词失败')
            return await extract_row(row, *args, **kwargs)

        with patch('app.title_worker.extract_row', new=fail_one_row):
            self.run_worker()
        path = f"/title-tasks/{job['id']}"
        before = self.request('GET', path).json()['result']
        selected = [before['words'][0]['index']]
        response = self.request('POST', f"/title-imports/{job['id']}/confirm", json={'selections': selected})
        self.assertEqual(response.status_code, 200, response.text)
        retry = self.request('POST', f"{path}/retry", json={})
        self.assertEqual(retry.status_code, 200, retry.text)
        self.run_worker()
        after = self.request('GET', path).json()['result']
        self.assertEqual((after['word_total'], after['imported_total'], after['failed_rows']), (3, 1, {}))
        self.assertEqual([word['index'] for word in after['words'] if word['imported']], selected)
        remaining = [word['index'] for word in after['words'] if not word['imported']]
        final = self.request('POST', f"/title-imports/{job['id']}/confirm", json={'selections': remaining})
        self.assertEqual(final.json(), {'created': 3, 'merged': 0, 'evidence': 3})

    def test_import_filters_all_candidates_before_pagination_and_preserves_options(self):
        content = KEYWORDS + 'baju,上衣,"上衣,连衣裙",品类\nseluar,裤子,裤子,品类\n'
        job = self.upload(content=content)
        self.request('POST', f"/title-tasks/{job['id']}/start", json={})
        self.run_worker()
        path = f"/title-tasks/{job['id']}"
        result = self.request('GET', path, params={'category': '连衣裙', 'tag': '品类', 'limit': 1, 'offset': 1}).json()['result']
        self.assertEqual((result['word_total'], result['filtered_total'], result['offset']), (5, 2, 1))
        self.assertEqual([word['phrase'] for word in result['words']], ['baju'])
        self.assertEqual(set(result['filter_categories']), {'连衣裙', '上衣', '裤子'})
        self.assertEqual(set(result['filter_tags']), {'品类', '结构设计', '图案'})
        clamped = self.request('GET', path, params={'category': '上衣', 'limit': 1, 'offset': 100}).json()['result']
        self.assertEqual((clamped['filtered_total'], clamped['offset'], clamped['words'][0]['phrase']), (1, 0, 'baju'))
        empty = self.request('GET', path, params={'category': '裤子', 'tag': '图案'}).json()['result']
        self.assertEqual((empty['filtered_total'], empty['words']), (0, []))
        self.assertEqual(empty['filter_categories'], result['filter_categories'])
        self.assertEqual(self.request('GET', path, uid=5, params={'category': '连衣裙'}).status_code, 404)

    def test_import_status_filter_combines_with_category_tag_and_pagination(self):
        content = KEYWORDS + 'baju,上衣,上衣,品类\nseluar,裤子,裤子,品类\n'
        job = self.upload(content=content)
        self.request('POST', f"/title-tasks/{job['id']}/start", json={})
        self.run_worker()
        path = f"/title-tasks/{job['id']}"
        before = self.request('GET', path, params={'imported': False}).json()['result']
        self.assertEqual(before['filtered_total'], 5)
        self.assertEqual(self.request('GET', path, params={'imported': True}).json()['result']['words'], [])
        selected = [word['index'] for word in before['words'][:2]]
        self.request('POST', f"/title-imports/{job['id']}/confirm", json={'selections': selected})
        imported = self.request('GET', path, params={'imported': True, 'offset': 1, 'limit': 1}).json()['result']
        self.assertEqual((imported['word_total'], imported['imported_total'], imported['filtered_total'], imported['offset']), (5, 2, 2, 1))
        self.assertEqual([word['index'] for word in imported['words']], selected[1:])
        self.assertTrue(all(word['imported'] for word in imported['words']))
        remaining = self.request('GET', path, params={'imported': False, 'offset': 99, 'limit': 2}).json()['result']
        self.assertEqual((remaining['filtered_total'], remaining['offset']), (3, 2))
        self.assertFalse(any(word['imported'] for word in remaining['words']))
        combined = self.request('GET', path, params={'category': '上衣', 'tag': '品类', 'imported': False}).json()['result']
        self.assertEqual((combined['filtered_total'], combined['words'][0]['phrase']), (1, 'baju'))
        empty = self.request('GET', path, params={'category': '上衣', 'tag': '品类', 'imported': True}).json()['result']
        self.assertEqual((empty['filtered_total'], empty['words']), (0, []))
        self.assertEqual(empty['filter_categories'], before['filter_categories'])
        listed = self.request('GET', '/title-tasks').json()['items'][0]['result']
        self.assertEqual((listed['imported_total'], listed['word_total']), (2, 5))
        self.assertEqual(self.request('GET', path, params={'imported': 'invalid'}).status_code, 422)
        self.request('POST', f"/title-imports/{job['id']}/confirm", json={
            'selections': [word['index'] for word in before['words']]})
        final = self.request('GET', path, params={'imported': False}).json()['result']
        self.assertEqual((final['filtered_total'], final['imported_total'], final['words']), (0, 5, []))

    def test_prompt_is_default_and_saves_english_even_when_library_configuration_exists(self):
        self.profile(); self.import_words()
        title = 'Floral Long Sleeve Everyday Dress'
        with patch.object(main, 'generate_draft_title', new=AsyncMock(return_value=title)) as prompt, \
             patch('app.title_service.ask_json', new=AsyncMock()) as library:
            for endpoint, body in (('/templates/1/generate-draft-title', {'image_url': 'https://img.example/1.png'}),
                                   ('/product-library/generate-draft-title', {'source_type': 'material', 'id': 1})):
                response = self.request('POST', endpoint, uid=2, json=body)
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json(), {'title': title})
            self.assertEqual(prompt.await_count, 2)
            library.assert_not_awaited()
        saved = self.request('POST', '/drafts/from-material-assets', uid=2,
            json={'template_id': 1, 'material_asset_ids': [1], 'title': title})
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertEqual(self.request('GET', '/title-statistics').json()['adopted'], 0)

    def test_explicit_library_generation_does_not_depend_on_retired_enable_flag(self):
        self.profile(); self.import_words()
        with self.sessions() as db:
            get_market(db, 1).enabled = False
            db.commit()
        with patch('app.title_service.ask_json', new=self.model()), \
             patch.object(main, 'generate_draft_title', new=AsyncMock()) as prompt:
            response = self.request('POST', '/templates/1/generate-draft-title', uid=2,
                json={'method': 'library', 'image_url': 'https://img.example/1.png'})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertTrue(response.json()['generation_id'])
            prompt.assert_not_awaited()
        job = self.request('POST', '/title-generation-jobs', uid=2, json={'template_id': 1, 'request_key': 'explicit-library-job',
            'groups': [{'key': 'one', 'source': {'source_type': 'material', 'id': 1}}]})
        self.assertEqual(job.status_code, 200, job.text)

    def test_library_missing_facts_does_not_silently_fallback_to_prompt(self):
        with patch.object(main, 'generate_draft_title', new=AsyncMock()) as prompt:
            response = self.request('POST', '/templates/1/generate-draft-title', uid=2,
                json={'method': 'library', 'image_url': 'https://img.example/1.png'})
            self.assertEqual(response.status_code, 400, response.text)
            self.assertIn('配置商品品类', response.json()['detail'])
            prompt.assert_not_awaited()

    def dictionary_items(self, kind):
        response = self.request('GET', '/title-dictionaries', params={'kind': kind})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()['items']

    def create_dictionary(self, kind, **payload):
        response = self.request('POST', f'/title-dictionaries/{kind}', json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_dictionary_crud_normalization_permissions_and_company_scope(self):
        self.enable()
        built_in = self.dictionary_items('tag')[0]
        self.assertTrue(built_in['built_in'])
        self.assertEqual(self.request('DELETE', f"/title-dictionaries/{built_in['id']}").status_code, 409)
        self.assertEqual(self.request('PATCH', f"/title-dictionaries/{built_in['id']}", json={'name': '改名'}).status_code, 409)
        category = self.create_dictionary('category', name=' Tops ')
        identity = category['id']
        self.assertEqual(category['name'], 'Tops')
        self.assertEqual(self.request('POST', '/title-dictionaries/category', json={'name': 'ＴＯＰＳ'}).status_code, 409)
        self.assertEqual(self.request('POST', '/title-dictionaries/tag', json={'name': '标签', 'semantic_tag': '伪造语义'}).status_code, 400)
        self.assertEqual(self.request('POST', '/title-dictionaries/tag', json={'name': 'x'*41, 'semantic_tag': '材质'}).status_code, 400)
        for uid in (2, 3, 4):
            self.assertEqual(self.request('POST', '/title-dictionaries/category', uid=uid, json={'name': 'Other'}).status_code, 403)
        self.assertEqual(self.request('GET', '/title-dictionaries', uid=2).status_code, 200)
        self.assertEqual(self.request('GET', '/title-dictionaries', uid=4).status_code, 403)
        self.assertEqual(self.request('GET', '/title-dictionaries', uid=5).json()['items'], [])
        self.assertEqual(self.request('PATCH', f'/title-dictionaries/{identity}', uid=5, json={'enabled': False}).status_code, 404)
        changed = self.request('PATCH', f'/title-dictionaries/{identity}', json={'name': '裙装', 'description': '适用裙装', 'enabled': False})
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertEqual(changed.json()['name'], '裙装')
        self.assertFalse(changed.json()['enabled'])
        self.assertEqual(self.request('DELETE', f'/title-dictionaries/{identity}').status_code, 200)
        self.assertEqual(self.request('DELETE', f'/title-dictionaries/{identity}').status_code, 404)

    def test_referenced_dictionary_blocks_identity_edits_and_disabled_category_blocks_new_generation(self):
        self.enable(); self.profile(); self.import_words()
        category = self.dictionary_items('category')[0]
        self.assertEqual((category['keyword_count'], category['template_count']), (3, 1))
        identity = category['id']
        self.assertEqual(self.request('PATCH', f'/title-dictionaries/{identity}', json={'name': '裙子'}).status_code, 409)
        self.assertEqual(self.request('DELETE', f'/title-dictionaries/{identity}').status_code, 409)
        before = self.request('GET', '/title-market').json()['revision']
        self.assertEqual(self.request('PATCH', f'/title-dictionaries/{identity}', json={'description': '已核对品类', 'enabled': False}).status_code, 200)
        after = self.request('GET', '/title-market').json()['revision']
        self.assertGreater(after, before)
        with patch('app.title_service.ask_json', new=self.model()) as model:
            response = self.request('POST', '/templates/1/generate-draft-title', uid=2, json={'method': 'library', 'image_url': 'https://img.example/1.png'})
            self.assertEqual(response.status_code, 400, response.text)
            model.assert_not_awaited()
        # 编辑其他字段不会丢失已停用的既有分类。
        self.assertEqual(self.request('PATCH', '/title-keywords/1', json={'translation': '裙装', 'categories': ['连衣裙']}).status_code, 200)
        upload = self.request('POST', '/title-imports', files={'file': ('words.csv', KEYWORDS.encode())},
            data={'kind': 'keywords', 'source_name': '新数据', 'period': '2026-10-30'})
        self.assertEqual(upload.status_code, 400)
        self.assertIn('停用', upload.json()['detail'])
        self.assertEqual(self.request('PATCH', f'/title-dictionaries/{identity}', json={'enabled': True}).status_code, 200)
        self.assertEqual(len(self.request('GET', '/title-keywords').json()['items']), 3)

    def test_gender_and_age_belong_to_template_and_category_no_longer_accepts_gender(self):
        category = self.create_dictionary('category', name='共享上衣')
        self.assertNotIn('gender', category)
        self.assertEqual(self.request('POST', '/title-dictionaries/category', json={'name': '旧设置', 'gender': '女'}).status_code, 422)
        self.assertEqual(self.request('PATCH', f"/title-dictionaries/{category['id']}", json={'gender': '女'}).status_code, 422)
        self.profile()
        for gender in ('男', '女', '中性'):
            for age in ('成人', '儿童'):
                result = self.profile(gender=gender, age_group=age)
                self.assertEqual((result['gender'], result['age_group']), (gender, age))
        for invalid in ({'gender': '儿童'}, {'gender': None}, {'age_group': '婴儿'}, {'age_group': None}):
            response = self.request('PUT', '/title-profiles/1', json={'category': '共享上衣', 'gender': '女', 'age_group': '成人', **invalid})
            self.assertEqual(response.status_code, 422, response.text)
        for uid in (2, 3, 5):
            self.assertIn(self.request('PUT', '/title-profiles/1', uid=uid, json={'category': '共享上衣', 'gender': '女', 'age_group': '成人'}).status_code, (403,404))

    def test_category_fields_are_separate_from_builtin_tags_and_core_words_remain_usable(self):
        self.enable(); self.profile(); self.import_words()
        data = self.request('GET', '/title-dictionaries').json()
        self.assertEqual(data['genders'], ['男', '女', '中性'])
        self.assertEqual(data['age_groups'], ['成人','儿童'])
        self.assertEqual(data['category_fields'], ['品类', '性别', '年龄'])
        for field in data['category_fields']:
            self.assertNotIn(field, data['semantic_tags'])
            self.assertNotIn(field, [row['name'] for row in data['items'] if row['kind'] == 'tag'])
            self.assertEqual(self.request('POST', '/title-dictionaries/tag', json={'name': field, 'semantic_tag': '待分类'}).status_code, 400)
            self.assertEqual(self.request('POST', '/title-dictionaries/tag', json={'name': f'{field}别名', 'semantic_tag': field}).status_code, 400)
        tag = self.dictionary_items('tag')[0]
        self.assertEqual(self.request('PATCH', f"/title-dictionaries/{tag['id']}", json={'gender': '女'}).status_code, 422)
        self.assertEqual(self.request('POST', '/title-dictionaries/tag', json={'name': '新标签', 'semantic_tag': '材质', 'gender': '女'}).status_code, 422)
        # 品类词保留来源及内部匹配语义，不因移除标签字典条目而丢失或无法编辑。
        core = next(word for word in self.request('GET', '/title-keywords').json()['items'] if word['tag'] == '品类')
        self.assertEqual(self.request('PATCH', f"/title-keywords/{core['id']}", json={'translation': '连衣裙'}).status_code, 200)
        with self.sessions() as db:
            from app.title_service import match_keywords
            matched = match_keywords(profile_snapshot(db, get_market(db, 1, create=True), 1), FACTS)
        self.assertTrue(any(word['id'] == core['id'] for word in matched))
        self.assertNotIn('品类', [row['name'] for row in self.dictionary_items('tag')])

    def test_site_category_name_is_editable_without_changing_identity_or_references(self):
        self.profile(); self.import_words()
        original = self.dictionary_items('category')[0]
        self.assertIsNone(original['localized_name'])
        identity = original['id']
        changed = self.request('PATCH', f'/title-dictionaries/{identity}', json={'localized_name': '  Gaun  '})
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertEqual(changed.json()['localized_name'], 'Gaun')
        self.assertEqual(changed.json()['name'], '连衣裙')
        self.assertEqual(changed.json()['reference_count'], original['reference_count'])
        self.assertEqual(self.request('GET', '/title-dictionaries', uid=2, params={'kind': 'category'}).json()['items'][0]['localized_name'], 'Gaun')
        for invalid in ('上衣', 'Baju\nGaun', 'Baju\x00'):
            self.assertEqual(self.request('PATCH', f'/title-dictionaries/{identity}', json={'localized_name': invalid}).status_code, 400)
        self.assertEqual(self.request('PATCH', f'/title-dictionaries/{identity}', json={'localized_name': 'a'*256}).status_code, 422)
        for uid in (2, 5):
            self.assertIn(self.request('PATCH', f'/title-dictionaries/{identity}', uid=uid, json={'localized_name': 'Baju'}).status_code, (403, 404))
        cleared = self.request('PATCH', f'/title-dictionaries/{identity}', json={'localized_name': ''})
        self.assertIsNone(cleared.json()['localized_name'])
        category = self.create_dictionary('category', name='上衣', localized_name='Baju')
        self.assertEqual(category['localized_name'], 'Baju')
        tag = self.dictionary_items('tag')[0]
        self.assertEqual(self.request('PATCH', f"/title-dictionaries/{tag['id']}", json={'localized_name': 'Nama'}).status_code, 400)
        self.assertEqual(self.request('POST', '/title-dictionaries/tag', json={'name': '额外标签', 'semantic_tag': '材质', 'localized_name': 'Nama'}).status_code, 400)

    def test_confirmed_site_name_is_used_for_category_and_invalidates_fact_cache(self):
        from app.title_service import identify_facts
        self.profile(); self.import_words()
        category = self.dictionary_items('category')[0]
        self.request('PATCH', f"/title-dictionaries/{category['id']}", json={'localized_name': 'Gaun'})
        with self.sessions() as db:
            snapshot = profile_snapshot(db, get_market(db, 1), 1)
            self.assertEqual(snapshot['category_localized_name'], 'Gaun')
            with patch('app.title_service.ask_json', new=self.model()) as model:
                first = asyncio.run(identify_facts(db, get_market(db, 1).id, snapshot, 'https://img.example/1.png', 'test-key'))
                self.assertEqual(next(f['ms'] for f in first['attributes'] if f['tag'] == '品类'), 'Gaun')

                cached = asyncio.run(identify_facts(db, get_market(db, 1).id, snapshot, 'https://img.example/1.png', 'test-key'))
                self.assertEqual(model.await_count, 1)
                self.assertEqual(first, cached)
                self.request('PATCH', f"/title-dictionaries/{category['id']}", json={'localized_name': 'Gaun Panjang'})
                latest = profile_snapshot(db, get_market(db, 1), 1)
                updated = asyncio.run(identify_facts(db, get_market(db, 1).id, latest, 'https://img.example/1.png', 'test-key'))
                self.assertEqual(model.await_count, 2)
                self.assertEqual(next(f['ms'] for f in updated['attributes'] if f['tag'] == '品类'), 'Gaun Panjang')
                self.assertEqual(snapshot['category_localized_name'], 'Gaun')
                self.assertEqual(next(f['ms'] for f in first['attributes'] if f['tag'] == '品类'), 'Gaun')

    def test_category_keyword_count_matches_filtered_words_after_merge_and_preserves_reference_guard(self):
        self.profile(); self.import_words(content=KEYWORDS+'gaun panjang,连衣裙,连衣裙,品类,否\n')
        words = self.request('GET', '/title-keywords', params={'category': '连衣裙'}).json()['items']
        target = next(word for word in words if word['phrase'] == 'gaun')
        source = next(word for word in words if word['phrase'] == 'gaun panjang')
        merged = self.request('POST', f"/title-keywords/{source['id']}/merge", json={'target_id': target['id']})
        self.assertEqual(merged.status_code, 200, merged.text)
        category = self.dictionary_items('category')[0]
        self.assertEqual(category['keyword_count'], 4)
        self.assertEqual(category['applicable_keyword_count'], 3)
        self.assertEqual(category['applicable_keyword_count'], self.request('GET', '/title-keywords', params={'category': category['name']}).json()['total'])
        self.assertEqual(self.request('DELETE', f"/title-dictionaries/{category['id']}").status_code, 409)

    def test_custom_tags_reach_import_filter_and_semantic_matching_without_bypassing_brand_rules(self):
        from app.title_service import match_keywords
        self.enable(); self.profile()
        sleeve = self.create_dictionary('tag', name='袖型', semantic_tag='结构设计')
        self.create_dictionary('tag', name='商标', semantic_tag='品牌')
        self.import_words(content=KEYWORDS+'lengan panjang,长袖,连衣裙,袖型,否\nAcme,商标,连衣裙,商标,否\n')
        filtered = self.request('GET', '/title-keywords', params={'tag': '袖型'}).json()['items']
        self.assertEqual(len(filtered), 1)
        brand = self.request('GET', '/title-keywords', params={'tag': '商标'}).json()['items'][0]
        self.assertEqual(brand['status'], 'enabled')
        with self.sessions() as db:
            snap = profile_snapshot(db, get_market(db, 1), 1)
        matched = match_keywords(snap, FACTS)
        self.assertTrue(any(row['id'] == filtered[0]['id'] and row['tag'] == '结构设计' and row['label_tag'] == '袖型' for row in matched))
        self.assertFalse(any(row['id'] == brand['id'] for row in matched))
        old_words = copy.deepcopy(snap['keywords'])
        self.request('PATCH', f"/title-dictionaries/{sleeve['id']}", json={'enabled': False})
        with self.sessions() as db:
            latest = profile_snapshot(db, get_market(db, 1), 1)
        self.assertFalse(any(row['id'] == filtered[0]['id'] for row in latest['keywords']))
        self.assertEqual(snap['keywords'], old_words)
        self.assertEqual(self.request('GET', '/title-keywords', params={'tag': '袖型'}).json()['items'][0]['status'], 'enabled')

    def test_template_list_includes_unconfigured_and_platform_templates_with_company_scope_and_pagination(self):
        with self.sessions() as db:
            db.add(ProductTemplate(id=200, name='平台共享模板', is_platform=True))
            db.add_all(ProductTemplate(id=i, name=f'待设置模板{i}', company_id=1) for i in range(10,36))
            db.commit()
        initial = self.request('GET', '/title-template-configurations', params={'page_size':100}).json()
        self.assertEqual(initial['total'], 28)
        expanded = self.request('GET', '/title-template-configurations', params={'page_size':1000}).json()
        self.assertEqual((len(expanded['items']),expanded['page_size']),(28,1000))
        clamped = self.request('GET', '/title-template-configurations', params={'page_size':20,'page':99}).json()
        self.assertEqual((clamped['page'],len(clamped['items'])),(2,8))
        self.assertEqual(self.request('GET','/title-template-configurations',params={'page_size':1001}).status_code,422)
        self.assertNotIn(2, [row['id'] for row in initial['items']])
        self.assertTrue(all(not row['ready'] for row in initial['items']))
        first = self.request('GET', '/title-template-configurations', params={'page_size':2}).json()
        second = self.request('GET', '/title-template-configurations', params={'page_size':2,'page':2}).json()
        self.assertEqual(len(first['items']), 2)
        self.assertFalse({row['id'] for row in first['items']} & {row['id'] for row in second['items']})
        self.profile()
        ready = self.request('GET', '/title-template-configurations', params={'state':'ready'}).json()
        self.assertEqual([row['id'] for row in ready['items']], [1])
        self.assertEqual((ready['items'][0]['gender'],ready['items'][0]['age_group']), ('女','成人'))
        self.assertEqual(self.request('GET', '/title-template-configurations', params={'q':'平台共享'}).json()['total'], 1)
        self.assertEqual(self.request('GET', '/title-template-configurations', params={'category':'连衣裙'}).json()['total'], 1)
        self.assertEqual(self.request('GET', '/title-template-configurations', uid=2).status_code, 200)
        self.assertEqual(self.request('GET', '/title-template-configurations', uid=4).status_code, 403)
        other = self.request('GET', '/title-template-configurations', uid=5).json()
        self.assertEqual({row['id'] for row in other['items']}, {2,200})
        self.assertTrue(all(not row['ready'] for row in other['items']))
        platform = self.request('PUT', '/title-profiles/200', json={'category':'上衣','gender':'中性','age_group':'成人'})
        self.assertEqual(platform.status_code, 200, platform.text)
        self.assertTrue(platform.json()['ready'])
        self.assertIsNone(self.request('GET', '/title-profiles/200', uid=5).json()['gender'])

    def test_incomplete_template_blocks_library_calls_but_not_prompt_generation(self):
        from app.title_models import TitleTemplateProfile
        self.profile()
        with self.sessions() as db:
            row = db.scalar(select(TitleTemplateProfile))
            row.gender = None; row.age_group = None
            db.commit()
        with patch('app.title_service.ask_json', new=AsyncMock()) as model:
            response = self.request('POST','/templates/1/generate-draft-title',uid=2,json={'method':'library','image_url':'https://img.example/1.png'})
            self.assertEqual(response.status_code,400,response.text)
            self.assertIn('模板配置',response.json()['detail'])
            batch = self.request('POST','/title-generation-jobs',uid=2,json={'template_id':1,'request_key':'missing-audience-test','groups':[{'key':'a','source':{'id':1}}]})
            self.assertEqual(batch.status_code,400,batch.text)
            model.assert_not_awaited()
        with patch.object(main,'generate_draft_title',new=AsyncMock(return_value='Floral Long Sleeve Everyday Dress')):
            ordinary = self.request('POST','/templates/1/generate-draft-title',uid=2,json={'image_url':'https://img.example/1.png'})
            self.assertEqual(ordinary.status_code,200,ordinary.text)

    def test_profile_material_options_use_enabled_material_keywords_in_current_market(self):
        from app.title_service import match_keywords
        self.create_dictionary('tag', name='布料', semantic_tag='材质')
        self.import_words(content=KEYWORDS +
            'kapas,纯棉,连衣裙,材质,否\ncotton,棉,连衣裙,材质,是\n'
            'poliester,涤纶,上衣,材质,否\nsatin,缎面,连衣裙,布料,否\n'
            'nilon,尼龙,连衣裙,材质,否\nlinen,亚麻,连衣裙,材质,否\n'
            'bulu,羊毛,连衣裙,材质,否\nvelvet,天鹅绒,连衣裙,材质,否\n')
        with self.sessions() as db:
            words = {word.phrase: word for word in db.scalars(select(TitleKeyword))}
            words['cotton'].tag = '面料'
            words['nilon'].status = 'disabled'
            words['linen'].priority = 'forbidden'
            words['bulu'].status = 'pending'
            words['velvet'].merged_into_id = words['kapas'].id
            db.commit()
        options = self.request('GET', '/title-profiles/1', uid=2).json()['material_options']
        self.assertEqual({item['phrase'] for item in options}, {'kapas', 'cotton', 'poliester', 'satin'})
        self.assertEqual(next(item for item in options if item['phrase'] == 'poliester')['categories'], ['上衣'])
        self.assertTrue(all(item['id'] and item['translation'] for item in options))
        full = self.request('GET', '/title-profiles/1').json()
        self.assertNotIn('velvet', {item['phrase'] for item in full['keyword_options']})
        self.assertTrue({'nilon', 'linen', 'bulu'} <= {item['phrase'] for item in full['keyword_options']})
        self.assertFalse({'nilon', 'linen', 'bulu', 'velvet'} & {item['phrase'] for item in full['attribute_keyword_options']})
        self.assertEqual(self.request('GET', '/title-profiles/1', uid=5).status_code, 404)
        self.assertEqual(self.request('GET', '/title-profiles/2', uid=5).json()['material_options'], [])
        self.assertEqual(self.request('GET', '/title-profiles/2', uid=5).json()['keyword_options'], [])
        saved = self.profile(attributes={'面料': 'kapas'})
        self.assertEqual(saved['material'], 'kapas')
        with self.sessions() as db:
            snapshot = profile_snapshot(db, get_market(db, 1), 1)
            matched = match_keywords(snapshot, {'attributes': [{'tag': '材质', 'value': 'kapas', 'ms': 'kapas'}]})
            self.assertIn('kapas', [item['phrase'] for item in matched])
        for tag in (item for item in self.dictionary_items('tag') if item['semantic_tag'] == '材质'):
            response = self.request('PATCH', f"/title-dictionaries/{tag['id']}", json={'enabled': False})
            self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.request('GET', '/title-profiles/1').json()['material_options'], [])

    def test_profile_attribute_keyword_choices_and_simplified_forbidden_settings(self):
        from app.title_service import match_keywords
        self.create_dictionary('tag', name='袖型', semantic_tag='结构设计')
        self.import_words(content=KEYWORDS + 'berlengan panjang,长袖,连衣裙,袖型,否\n')
        data = self.request('GET', '/title-profiles/1').json()
        custom = next(item for item in data['attribute_keyword_options'] if item['tag'] == '袖型')
        self.assertEqual((custom['phrase'], custom['translation']), ('berlengan panjang', '长袖'))
        words = {item['phrase']: item for item in data['keyword_options']}
        banned = words['corak bunga']['id']
        self.profile(locked_keyword_ids=[words['gaun']['id']], category_names=['旧平台类目'],
                     forbidden_keyword_ids=[banned])
        response = self.request('PUT', '/title-profiles/1', json={'category': '连衣裙', 'gender': '女',
            'age_group': '成人', 'attributes': {'袖型': custom['phrase']}, 'forbidden_keyword_ids': [banned]})
        self.assertEqual(response.status_code, 200, response.text)
        saved = response.json()
        self.assertEqual((saved['category_names'], saved['locked_keyword_ids'], saved['forbidden_keyword_ids']), ([], [], [banned]))
        self.assertEqual(saved['attributes'], {'袖型': 'berlengan panjang'})
        with self.sessions() as db:
            snapshot = profile_snapshot(db, get_market(db, 1), 1)
            matched = match_keywords(snapshot, FACTS)
            self.assertNotIn(banned, [item['id'] for item in matched])
        tag = next(item for item in self.dictionary_items('tag') if item['name'] == '袖型')
        self.request('PATCH', f"/title-dictionaries/{tag['id']}", json={'enabled': False})
        latest = self.request('GET', '/title-profiles/1').json()
        self.assertNotIn(custom['id'], [item['id'] for item in latest['attribute_keyword_options']])

    def test_material_and_fabric_share_one_setting_and_historical_words_still_match(self):
        from app.title_service import match_keywords
        result = self.profile(attributes={'面料':'纯棉'})
        self.assertEqual(result['attributes'],{'材质':'纯棉'})
        self.assertEqual(result['material'],'纯棉')
        self.assertNotIn('面料',result['tags'])
        self.assertNotIn('年龄',result['tags'])
        duplicate = self.profile(attributes={'材质':'纯棉','面料':' 纯棉 '})
        self.assertEqual(duplicate['attributes'],{'材质':'纯棉'})
        conflict = self.request('PUT','/title-profiles/1',json={'category':'连衣裙','gender':'女','age_group':'成人','attributes':{'材质':'棉','面料':'涤纶'}})
        self.assertEqual(conflict.status_code,400,conflict.text)
        self.import_words(content=KEYWORDS+'kapas,纯棉,连衣裙,面料,否\n')
        with self.sessions() as db:
            word = db.scalar(select(TitleKeyword).where(TitleKeyword.phrase=='kapas'))
            self.assertEqual(word.tag,'材质')
            word.tag='面料'; word.identity_hash=digest(['kapas','面料'])
            db.commit()
            snap = profile_snapshot(db,get_market(db,1),1)
            facts = {'attributes':[{'tag':'材质','value':'纯棉','ms':'kapas'}]}
            matched = match_keywords(snap,facts)
            self.assertTrue(any(item['id']==word.id and item['tag']=='材质' for item in matched))
            word_id = word.id
        material_tag = next(row for row in self.dictionary_items('tag') if row['name']=='材质')
        self.request('PATCH',f"/title-dictionaries/{material_tag['id']}",json={'enabled':False})
        with self.sessions() as db:
            latest = profile_snapshot(db,get_market(db,1),1)
            self.assertFalse(any(item['id']==word_id for item in latest['keywords']))
        self.assertNotIn('面料',[row['name'] for row in self.dictionary_items('tag')])

    def test_custom_template_tags_are_resolved_to_fact_semantics_and_unknown_tags_are_rejected(self):
        self.enable(); self.create_dictionary('tag',name='袖型',semantic_tag='结构设计')
        self.profile(attributes={'袖型':'长袖'})
        with self.sessions() as db:
            snapshot = profile_snapshot(db,get_market(db,1),1)
            self.assertEqual(snapshot['attributes'],{'结构设计':'长袖'})
            self.assertEqual((snapshot['gender'],snapshot['age_group']),('女','成人'))
        invalid = self.request('PUT','/title-profiles/1',json={'category':'连衣裙','gender':'女','age_group':'成人','attributes':{'伪造标签':'值'}})
        self.assertEqual(invalid.status_code,400,invalid.text)
        with patch('app.title_service.ask_json',new=self.model()):
            generated = self.request('POST','/templates/1/generate-draft-title',uid=2,json={'method':'library','image_url':'https://img.example/1.png'})
            self.assertEqual(generated.status_code,200,generated.text)
        adopted = self.request('POST','/drafts/from-material-assets',uid=2,json={'template_id':1,'material_asset_ids':[1],'title':TITLE,'generation_id':generated.json()['generation_id']})
        self.assertEqual(adopted.status_code,200,adopted.text)

    def test_gender_changes_after_generation_require_new_title_before_adoption(self):
        self.profile()
        with patch('app.title_service.ask_json',new=self.model()):
            response = self.request('POST','/templates/1/generate-draft-title',uid=2,json={'method':'library','image_url':'https://img.example/1.png'})
            self.assertEqual(response.status_code,200,response.text)
            generated = response.json()
        self.profile(gender='男')
        created = self.request('POST','/drafts/from-material-assets',uid=2,json={'template_id':1,'material_asset_ids':[1],'title':TITLE,'generation_id':generated['generation_id']})
        self.assertEqual(created.status_code,400,created.text)
        self.assertIn('属性已改变',created.json()['detail'])

    def test_pending_import_protects_dictionary_and_approval_checks_current_enabled_state(self):
        self.enable()
        tag = self.create_dictionary('tag', name='袖型', semantic_tag='结构设计')
        category = self.create_dictionary('category', name='连衣裙')
        job = self.upload(content='马来语词语,中文释义,适用品类,属性分类\nlengan panjang,长袖,连衣裙,袖型\n')
        self.assertEqual(next(t for t in self.dictionary_items('tag') if t['id'] == tag['id'])['import_count'], 1)
        for identity in (tag['id'], category['id']):
            self.assertEqual(self.request('DELETE', f'/title-dictionaries/{identity}').status_code, 409)
        self.request('POST', f"/title-tasks/{job['id']}/start", json={})
        self.request('PATCH', f"/title-dictionaries/{tag['id']}", json={'enabled': False})
        self.run_worker()  # 提词使用提交时的标签快照。
        task = self.request('GET', f"/title-tasks/{job['id']}").json()
        self.assertEqual(task['state'], 'review', task)
        confirmation = {'selections': [task['result']['words'][0]['index']], 'enable': True}
        rejected = self.request('POST', f"/title-imports/{job['id']}/confirm", json=confirmation)
        self.assertEqual(rejected.status_code, 400)
        self.assertEqual(self.request('GET', '/title-keywords').json()['total'], 0)
        self.request('PATCH', f"/title-dictionaries/{tag['id']}", json={'enabled': True})
        self.assertEqual(self.request('POST', f"/title-imports/{job['id']}/confirm", json=confirmation).status_code, 200)
        with self.sessions() as db:
            self.assertEqual(db.scalar(select(func.count(TitleKeyword.id))), 1)

    def test_keyword_import_only_accepts_one_tag_per_row(self):
        headers = ['马来语词语', '中文释义', '产品分类', '标签']
        for tag in ('材质、功能', '材质;功能', '材质；功能'):
            book = Workbook()
            book.active.append(headers)
            book.active.append(['kapas', '棉', '上衣', tag])
            output = io.BytesIO()
            book.save(output)
            csv_content = (','.join(headers) + f'\nkapas,棉,上衣,{tag}\n').encode()
            for filename, content in (('tags.csv', csv_content), ('tags.xlsx', output.getvalue())):
                with self.subTest(tag=tag, filename=filename), self.assertRaises(HTTPException) as invalid:
                    parse_file(content, filename, 'keywords')
                self.assertIn('第 2 行只能填写一个标签', invalid.exception.detail)
        rows = parse_file('马来语词语,中文释义,产品分类,标签\nkapas,棉,上衣,面料\n'.encode(), 'single.csv', 'keywords')
        words = asyncio.run(extract_row(rows[0], 'keywords'))
        self.assertEqual(len(words), 1)
        self.assertEqual(words[0]['tag'], '材质')

    def test_scope_and_role_boundaries(self):
        self.assertTrue(self.request('GET', '/title-market').json()['enabled'])
        for uid in (2, 3, 4):
            self.assertEqual(self.request('PUT', '/title-market', uid=uid, json={'enabled': True}).status_code, 403)
        for query in ({'site': 'ID'}, {'language': 'en'}, {'platform': 'Temu'}):
            self.assertEqual(self.request('GET', '/title-market', params=query).status_code, 400)
        self.enable()
        self.assertTrue(self.request('GET', '/title-market', uid=5).json()['enabled'])
        self.assertTrue(self.request('GET', '/title-market', params={'platform': 'Shopee'}).json()['enabled'])
        self.import_words()
        self.assertEqual(self.request('GET', '/title-keywords', uid=2).json()['total'], 3)
        self.assertEqual(self.request('GET', '/title-keywords', uid=5).json()['total'], 0)
        self.assertEqual(self.request('GET', '/title-keywords', uid=4).status_code, 403)
        self.assertEqual(self.request('GET', '/title-keywords', params={'platform': 'Shopee'}).json()['total'], 3)
        self.assertEqual(self.request('GET', '/title-keywords/1/evidence', uid=5).status_code, 404)
        self.assertEqual(self.request('PATCH', '/title-keywords/1', uid=2, json={'status': 'disabled'}).status_code, 403)

    def test_platform_alias_shares_profiles_imports_and_settings(self):
        self.enable('Shopee')
        self.profile(platform='Shopee', attributes={'材质': '纯棉'})
        profile = self.request('GET', '/title-profiles/1').json()
        self.assertEqual(profile['attributes'], {'材质': '纯棉'})
        identity = self.import_words(platform='Shopee')
        self.assertEqual(self.upload()['id'], identity)
        self.assertEqual(self.request('GET', '/title-tasks').json()['items'][0]['id'], identity)
        self.request('PUT', '/title-market', json={'platform': 'Shopee', 'enabled': False})
        self.assertTrue(self.request('GET', '/title-market').json()['enabled'])
        with self.sessions() as db:
            self.assertEqual(db.scalar(select(func.count(TitleMarket.id)).where(TitleMarket.company_id == 1)), 1)

    def test_keyword_status_is_admin_scoped_and_checks_current_content(self):
        self.import_words()
        self.request('PATCH','/title-keywords/1',json={'status':'disabled'})
        def word():
            return next(row for row in self.request('GET','/title-keywords').json()['items'] if row['id']==1)
        initial = word()
        payload = {'status':'enabled','review_token':initial['review_token']}
        self.assertEqual(self.request('POST','/title-keywords/1/status',uid=2,json=payload).status_code,403)
        self.assertEqual(self.request('POST','/title-keywords/1/status',uid=5,json=payload).status_code,404)
        edited = self.request('PATCH','/title-keywords/1',json={'translation':'连衣裙（已核对）'})
        self.assertEqual(edited.status_code,200,edited.text)
        self.assertEqual(word()['status'],'disabled')
        stale = self.request('POST','/title-keywords/1/status',json=payload)
        self.assertEqual(stale.status_code,409,stale.text)
        current = word()
        # 生成选用次数变化不影响管理员正在核对的词条内容。
        with self.sessions() as db:
            db.get(TitleKeyword,1).selected_count += 1
            db.commit()
        approved = self.request('POST','/title-keywords/1/status',json={'status':'enabled','review_token':current['review_token']})
        self.assertEqual(approved.status_code,200,approved.text)
        self.assertEqual(word()['status'],'enabled')
        rejected = self.request('POST','/title-keywords/1/status',json={'status':'disabled','review_token':word()['review_token']})
        self.assertEqual(rejected.status_code,200,rejected.text)
        self.assertEqual(word()['status'],'disabled')
        self.assertEqual(self.request('POST','/title-keywords/1/status',json={'status':'pending','review_token':word()['review_token']}).status_code,422)
        self.assertEqual(self.request('POST','/title-keywords/1/status',json={'status':'enabled'}).status_code,422)
        self.assertEqual(self.request('PATCH','/title-keywords/1',json={'status':'pending'}).status_code,422)
        # 旧客户端的审核 URL 保持兼容，行为与启停接口相同。
        compatible = self.request('POST','/title-keywords/1/review',json={'status':'enabled','review_token':word()['review_token']})
        self.assertEqual(compatible.status_code,200,compatible.text)

    def test_fabric_semantic_alias_accepts_new_name_without_rewriting_historical_fields(self):
        self.enable()
        created = self.request('POST','/title-dictionaries/tag',json={'name':'面料说明','semantic_tag':'面料'})
        self.assertEqual(created.status_code,200,created.text)
        self.assertEqual(created.json()['semantic_tag'],'材质')
        updated = self.request('PATCH',f"/title-dictionaries/{created.json()['id']}",json={'semantic_tag':'面料','description':'运营确认的面料信息'})
        self.assertEqual(updated.status_code,200,updated.text)
        profile = self.profile(attributes={'面料':'纯棉'})
        self.assertEqual(profile['material'],'纯棉')

    def test_keyword_status_does_not_enable_disabled_dictionary_entries_or_merged_words(self):
        self.import_words(content=KEYWORDS+'gaun panjang,长连衣裙,连衣裙,品类,否\n')
        categories = self.request('GET','/title-dictionaries',params={'kind':'category'}).json()['items']
        category = next(row for row in categories if row['name']=='连衣裙')
        self.request('PATCH',f"/title-dictionaries/{category['id']}",json={'enabled':False})
        words = self.request('GET','/title-keywords').json()['items']
        blocked = self.request('POST','/title-keywords/1/status',json={'status':'enabled','review_token':words[0]['review_token']})
        self.assertEqual(blocked.status_code,400,blocked.text)
        disabled = self.request('POST','/title-keywords/1/status',json={'status':'disabled','review_token':words[0]['review_token']})
        self.assertEqual(disabled.status_code,200,disabled.text)
        self.request('PATCH',f"/title-dictionaries/{category['id']}",json={'enabled':True})
        token = next(row for row in self.request('GET','/title-keywords').json()['items'] if row['id']==4)['review_token']
        self.assertEqual(self.request('POST','/title-keywords/4/merge',json={'target_id':1}).status_code,200)
        merged = self.request('POST','/title-keywords/4/status',json={'status':'enabled','review_token':token})
        self.assertEqual(merged.status_code,400,merged.text)

    def test_keywords_and_evidence_support_page_sizes_and_clamp_empty_or_out_of_range_pages(self):
        self.enable()
        with self.sessions() as db:
            market = get_market(db,1)
            for index in range(25):
                db.add(TitleKeyword(market_id=market.id,code=f'KW{index+1:06d}',phrase=f'baju {index}',translation='上衣',tag='品类',categories=['上衣'],identity_hash=digest([f'baju {index}','品类'])))
            db.flush()
            for index in range(25):
                db.add(TitleEvidence(market_id=market.id,keyword_id=1,import_id=100+index,observation_hash=digest(['evidence',index]),original_title=f'Baju Wanita {index}',original_phrase='baju',source_name='手动导入',source_product=str(index),period='2026-10-30',source_kind='titles',position=0,metrics={}))
            db.commit()
        for endpoint in ('/title-keywords','/title-keywords/1/evidence'):
            with self.subTest(endpoint=endpoint):
                first = self.request('GET',endpoint,params={'page_size':20}).json()
                second = self.request('GET',endpoint,params={'page':2,'page_size':20}).json()
                self.assertEqual((first['total'],len(first['items']),first['page'],first['page_size']),(25,20,1,20))
                self.assertEqual((second['total'],len(second['items']),second['page']),(25,5,2))
                large = self.request('GET',endpoint,params={'page_size':1000}).json()
                self.assertEqual(len(large['items']),25)
                clamped = self.request('GET',endpoint,params={'page':99,'page_size':20}).json()
                self.assertEqual((clamped['page'],clamped['items']),(2,second['items']))
                for params in ({'page_size':0},{'page_size':1001},{'page':0}):
                    self.assertEqual(self.request('GET',endpoint,params=params).status_code,422)
        empty = self.request('GET','/title-keywords',params={'search':'没有匹配的词','page':9,'page_size':20}).json()
        self.assertEqual((empty['total'],empty['items'],empty['page']),(0,[],1))
        missing_evidence = self.request('GET','/title-keywords/2/evidence',params={'page':9,'page_size':20}).json()
        self.assertEqual((missing_evidence['total'],missing_evidence['items'],missing_evidence['page']),(0,[],1))

    def test_configuration_timestamp_ignores_upload_preview_and_advances_on_merge(self):
        self.import_words(content=KEYWORDS+'gaun panjang,长连衣裙,连衣裙,品类,否\n')
        saved_at = datetime(2026, 1, 1)
        with self.sessions() as db:
            market = get_market(db, 1)
            market.updated_at = saved_at
            version = market.revision
            db.commit()
        self.upload(filename='preview.csv', period='2026-10-30')
        before = self.request('GET', '/title-market').json()
        self.assertEqual(before['updated_at'], saved_at.isoformat()+'Z')
        response = self.request('POST', '/title-keywords/4/merge', json={'target_id': 1})
        self.assertEqual(response.status_code, 200, response.text)
        after = self.request('GET', '/title-market').json()
        self.assertGreater(after['updated_at'], before['updated_at'])
        self.assertEqual(after['revision'], version+1)
        self.assertEqual(self.request('GET', '/title-keywords').json()['total'], 3)

    def test_import_deduplicates_files_evidence_and_keeps_codes(self):
        identity = self.import_words()
        duplicate = self.upload()
        self.assertEqual((duplicate['id'], duplicate['duplicate']), (identity, True))
        self.import_words(content=KEYWORDS+'\n')
        words = self.request('GET', '/title-keywords').json()['items']
        self.assertEqual([w['code'] for w in words], ['KW000001', 'KW000002', 'KW000003'])
        self.assertTrue(all(w['evidence_count'] == 1 and w['title_frequency'] == 0 for w in words))
        self.assertTrue(all(w['associated_sales'] is None and w['search_metrics'] == [] for w in words))
        self.import_words(period='2026-10-30')
        words = self.request('GET', '/title-keywords').json()['items']
        self.assertTrue(all(w['evidence_count'] == 2 for w in words))
        evidence = self.request('GET', '/title-keywords/2/evidence').json()['items'][0]
        self.assertEqual((evidence['original_phrase'], evidence['position']), ('lengan panjang', 0))
        self.assertEqual(self.request('POST', f'/title-imports/{identity}/confirm', json={'selections': [], 'enable': True}).status_code, 200)

    def test_note_only_import_uses_system_metadata_and_preserves_note_through_review_and_evidence(self):
        uploaded = self.request('POST','/title-imports',files={'file':('words.csv',KEYWORDS.encode())},
            data={'kind':'keywords','note':'  女装长袖词语整理\n第一批  '})
        self.assertEqual(uploaded.status_code,200,uploaded.text)
        preview = uploaded.json()
        self.confirm_new_categories(preview)
        self.assertEqual(preview['note'],'女装长袖词语整理\n第一批')
        self.request('POST',f"/title-tasks/{preview['id']}/start",json={})
        self.run_worker()
        task = self.request('GET',f"/title-tasks/{preview['id']}").json()
        self.assertEqual(task['state'],'review',task)
        self.assertEqual(task['source_name'],'手动导入')
        self.assertEqual(task['period_basis'],'import_date')
        self.assertEqual(task['note'],preview['note'])
        self.assertEqual(self.request('GET','/title-tasks').json()['items'][0]['note'],preview['note'])
        confirmed = self.request('POST',f"/title-imports/{preview['id']}/confirm",json={'selections':[word['index'] for word in task['result']['words']],'enable':True})
        self.assertEqual(confirmed.status_code,200,confirmed.text)
        evidence = self.request('GET','/title-keywords/1/evidence').json()['items'][0]
        self.assertEqual((evidence['note'],evidence['period_basis']), (preview['note'],'import_date'))
        self.assertEqual(self.request('GET',f"/title-tasks/{preview['id']}",uid=5).status_code,404)

    def test_task_list_paginates_with_market_user_and_kind_scope(self):
        self.enable()
        with self.sessions() as db:
            market = get_market(db, 1)
            other_market = get_market(db, 2, create=True)
            records = []
            for index in range(35):
                task = TitleTask(market_id=market.id, user_id=1 if index%2 else 2, kind='import',
                    request_key=f'page-import-{index}', request_hash=f'page-import-{index}', state='completed',
                    snapshot={'kind':'titles' if index%2 else 'keywords','filename':f'{index}.csv','note':f'第 {index} 批'},
                    result={'word_total':index},progress={'total':50,'completed':50})
                db.add(task)
                records.append(task)
            db.add(TitleTask(market_id=market.id,user_id=1,kind='batch',request_key='page-batch',request_hash='page-batch',snapshot={}))
            db.add(TitleTask(market_id=other_market.id,user_id=5,kind='import',request_key='page-other',request_hash='page-other',snapshot={'kind':'titles'}))
            db.flush()
            identities = [task.id for task in reversed(records)]
            member_ids = [task.id for task in reversed(records) if task.user_id==2]
            db.commit()
        first = self.request('GET','/title-tasks',params={'page':1,'page_size':20}).json()
        second = self.request('GET','/title-tasks',params={'page':2,'page_size':20}).json()
        self.assertEqual((first['total'],first['page'],first['page_size']),(35,1,20))
        self.assertEqual([task['id'] for task in first['items']],identities[:20])
        self.assertEqual([task['id'] for task in second['items']],identities[20:])
        self.assertEqual(first['items'][0]['import_kind'],'keywords')
        self.assertEqual(first['items'][0]['note'],'第 34 批')
        default = self.request('GET','/title-tasks').json()
        self.assertEqual((len(default['items']),default['page_size']),(30,30))
        expanded = self.request('GET','/title-tasks',params={'page_size':50}).json()
        self.assertEqual([task['id'] for task in expanded['items']],identities)
        clamped = self.request('GET','/title-tasks',params={'page':99,'page_size':20}).json()
        self.assertEqual((clamped['page'],[task['id'] for task in clamped['items']]),(2,identities[20:]))
        member = self.request('GET','/title-tasks',uid=2,params={'page_size':10,'page':2}).json()
        self.assertEqual((member['total'],[task['id'] for task in member['items']]),(18,member_ids[10:]))
        batch = self.request('GET','/title-tasks',params={'kind':'batch'}).json()
        self.assertEqual((batch['total'],batch['items'][0]['import_kind']),(1,None))
        other = self.request('GET','/title-tasks',uid=5).json()
        self.assertEqual(other['total'],1)
        self.assertNotIn(other['items'][0]['id'],identities)
        for params in ({'page':0},{'page_size':0},{'page_size':1001}):
            self.assertEqual(self.request('GET','/title-tasks',params=params).status_code,422)

    def test_empty_task_list_has_consistent_pagination_metadata(self):
        for enabled in (False,True):
            if enabled:
                self.enable()
            response = self.request('GET','/title-tasks',params={'page':3,'page_size':50})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(response.json(),{'items':[],'total':0,'page':1,'page_size':50})

    def test_note_and_later_upload_date_do_not_create_a_second_task_for_the_same_file(self):
        with patch('app.title_api.datetime',wraps=datetime) as clock:
            clock.now.return_value=datetime(2026,10,30,23,59)
            first = self.request('POST','/title-imports',files={'file':('first.csv',KEYWORDS.encode())},data={'kind':'keywords','note':'原备注'})
            self.assertEqual(first.status_code,200,first.text)
            self.assertEqual(str(clock.now.call_args.args[0]),'Asia/Hong_Kong')
            clock.now.return_value=datetime(2026,10,31,0,1)
            repeated = self.request('POST','/title-imports',files={'file':('renamed.csv',KEYWORDS.encode())},data={'kind':'keywords','note':'新备注'})
        self.assertEqual(repeated.status_code,200,repeated.text)
        self.assertEqual((repeated.json()['id'],repeated.json()['duplicate'],repeated.json()['note']), (first.json()['id'],True,'原备注'))
        detail = self.request('GET',f"/title-tasks/{first.json()['id']}").json()
        self.assertEqual(detail['period'],'2026-10-30')
        self.assertEqual(detail['filename'],'first.csv')
        with self.sessions() as db:
            self.assertEqual(db.scalar(select(func.count(TitleTask.id))),1)

    def test_note_is_optional_bounded_and_does_not_change_same_day_evidence_deduplication(self):
        for content,note in ((KEYWORDS,''),(KEYWORDS+'\n','第二份文件')):
            uploaded = self.request('POST','/title-imports',files={'file':('words.csv',content.encode())},data={'kind':'keywords','note':note})
            self.assertEqual(uploaded.status_code,200,uploaded.text)
            identity = uploaded.json()['id']
            self.confirm_new_categories(uploaded.json())
            self.request('POST',f'/title-tasks/{identity}/start',json={});self.run_worker()
            task = self.request('GET',f'/title-tasks/{identity}').json()
            result = self.request('POST',f'/title-imports/{identity}/confirm',json={'selections':[word['index'] for word in task['result']['words']],'enable':True})
            self.assertEqual(result.status_code,200,result.text)
        words = self.request('GET','/title-keywords').json()['items']
        self.assertTrue(all(word['evidence_count']==1 for word in words))
        self.assertEqual(len(words),3)
        missing_note = self.request('POST','/title-imports',files={'file':('optional.csv',(KEYWORDS+'\n\n').encode())},data={'kind':'keywords'})
        self.assertEqual(missing_note.status_code,200,missing_note.text)
        self.assertEqual(missing_note.json()['note'],'')
        too_long = self.request('POST','/title-imports',files={'file':('words.csv',KEYWORDS.encode())},data={'kind':'keywords','note':'字'*501})
        self.assertEqual(too_long.status_code,422,too_long.text)

    def test_legacy_import_metadata_is_validated_but_unused_metrics_are_ignored(self):
        invalid = self.request('POST','/title-imports',files={'file':('words.csv',KEYWORDS.encode())},
            data={'kind':'keywords','source_name':'历史采集','period':'2026-02-30'})
        self.assertEqual(invalid.status_code,400,invalid.text)
        unused_metrics = self.request('POST','/title-imports',files={'file':('titles.csv','原标题,商品品类,搜索热度\nBaju wanita,上衣,100\n'.encode())},data={'kind':'titles'})
        self.assertEqual(unused_metrics.status_code,200,unused_metrics.text)
        self.assertNotIn('搜索热度',unused_metrics.json()['preview'][0]['values'])
        valid_metrics = self.request('POST','/title-imports',files={'file':('titles.csv','原标题,商品品类,搜索热度,搜索来源,搜索周期\nBaju wanita,上衣,100,真实搜索数据,2026-10\n'.encode())},data={'kind':'titles','note':'搜索数据样本'})
        self.assertEqual(valid_metrics.status_code,200,valid_metrics.text)
        self.assertEqual(valid_metrics.json()['preview'][0]['values'],{'原标题':'Baju wanita','商品品类':'上衣'})

    def test_import_accepts_10000_rows_and_rejects_10001_rows(self):
        for kind, headers, values in (
            ('titles', ['商品标题', '商品分类'], ['Baju wanita', '上衣']),
            ('keywords', ['马来语词语', '中文释义', '产品分类', '标签'], ['baju', '上衣', '上衣', '品类']),
        ):
            book = Workbook()
            book.active.append(headers)
            for _ in range(10000):
                book.active.append(values)
            csv_content = (','.join(headers) + '\n' + (','.join(values) + '\n') * 10000).encode()
            for count in (10000, 10001):
                xlsx = io.BytesIO()
                book.save(xlsx)
                for filename, content in (('rows.csv', csv_content), ('rows.xlsx', xlsx.getvalue())):
                    with self.subTest(kind=kind, count=count, filename=filename):
                        if count == 10000:
                            self.assertEqual(len(parse_file(content, filename, kind)), 10000)
                        else:
                            with self.assertRaises(HTTPException) as oversized:
                                parse_file(content, filename, kind)
                            self.assertIn('单次最多导入 10000 行', oversized.exception.detail)
                book.active.append(values)
                csv_content += (','.join(values) + '\n').encode()

    def test_import_accepts_20mb_and_rejects_one_byte_over(self):
        header = '商品标题,商品分类,备注,补充\n'.encode()
        prefix = 'Baju wanita,上衣,'.encode() + b'x' * 1990 + b','
        row_size, longer_rows = divmod(20 * 1024 * 1024 - len(header), 10000)
        short_row = prefix + b'y' * (row_size - len(prefix) - 1) + b'\n'
        long_row = short_row[:-1] + b'y\n'
        content = header + short_row * (10000 - longer_rows) + long_row * longer_rows
        self.assertEqual(len(content), 20 * 1024 * 1024)
        self.assertEqual(len(parse_file(content, 'limit.csv', 'titles')), 10000)
        with self.assertRaises(HTTPException) as oversized:
            parse_file(content + b'\n', 'oversized.csv', 'titles')
        self.assertIn('超过 20MB', oversized.exception.detail)

    def test_xlsx_and_import_columns_validation(self):
        book = Workbook(); sheet = book.active
        sheet.append(['商品标题', '商品分类']); sheet.append(['Gaun berlengan panjang bercorak bunga', '连衣裙'])
        content = io.BytesIO(); book.save(content)
        rows = parse_file(content.getvalue(), 'sample.xlsx', 'titles')
        self.assertEqual(rows[0]['row'], 2)
        self.assertEqual(rows[0]['values']['商品品类'], '连衣裙')
        current = parse_file('商品标题,商品分类\nBaju wanita,上衣\n'.encode(), 'current.csv', 'titles')
        legacy = parse_file('原标题,商品品类\nBaju wanita,上衣\n'.encode(), 'legacy.csv', 'titles')
        self.assertEqual(current, legacy)
        template = self.request('GET', '/title-import-template', params={'kind': 'titles'})
        self.assertEqual(template.status_code, 200, template.text)
        self.assertEqual(template.content.decode('utf-8-sig').splitlines()[0].split(','), ['商品标题', '商品分类', '商品ID', '链接'])
        with self.assertRaises(HTTPException) as missing:
            parse_file('商品ID\n10001\n'.encode(), 'missing.csv', 'titles')
        self.assertIn('缺少列：商品标题、商品分类', missing.exception.detail)
        for headers in ('商品标题,原标题,商品分类', '商品标题,商品分类,商品品类'):
            with self.subTest(headers=headers), self.assertRaises(HTTPException) as duplicate:
                parse_file(f'{headers}\nBaju wanita,上衣,上衣\n'.encode(), 'ambiguous.csv', 'titles')
            self.assertIn('表头不能重复', duplicate.exception.detail)
        for text in ('原标题,商品品类\nBaju wanita,',
                     '商品标题,商品分类\n,上衣'):
            with self.assertRaises(HTTPException):
                parse_file(text.encode(), 'sample.csv', 'titles')

    def test_keyword_import_column_names_and_legacy_compatibility(self):
        current = parse_file('马来语词语,中文释义,产品分类,标签\ngaun,连衣裙,连衣裙,品类\n'.encode(), 'current.csv', 'keywords')
        legacy = parse_file('马来语词语,中文释义,适用品类,属性分类\ngaun,连衣裙,连衣裙,品类\n'.encode(), 'legacy.csv', 'keywords')
        self.assertEqual(current, legacy)
        book = Workbook()
        book.active.append(['马来语词语', '中文释义', '产品分类', '标签'])
        book.active.append(['gaun', '连衣裙', '连衣裙', '品类'])
        content = io.BytesIO()
        book.save(content)
        self.assertEqual(parse_file(content.getvalue(), 'current.xlsx', 'keywords'), current)
        with self.assertRaises(HTTPException) as missing:
            parse_file('马来语词语,中文释义\ngaun,连衣裙\n'.encode(), 'missing.csv', 'keywords')
        self.assertIn('缺少列：产品分类、标签', missing.exception.detail)
        for headers in ('马来语词语,中文释义,产品分类,适用品类,标签', '马来语词语,中文释义,产品分类,标签,属性分类'):
            with self.subTest(headers=headers), self.assertRaises(HTTPException) as duplicate:
                parse_file(f'{headers}\ngaun,连衣裙,连衣裙,品类,品类\n'.encode(), 'ambiguous.csv', 'keywords')
            self.assertIn('表头不能重复', duplicate.exception.detail)

    def test_both_import_templates_only_offer_product_id_and_link_and_ignore_metric_columns(self):
        for kind,headers,values,stored in (
            ('titles','商品标题,商品分类','Baju wanita,上衣',{'原标题':'Baju wanita','商品品类':'上衣'}),
            ('keywords','马来语词语,中文释义,产品分类,标签','baju,上衣,上衣,品类',{'马来语词语':'baju','中文释义':'上衣','适用品类':'上衣','属性分类':'品类'}),
        ):
            with self.subTest(kind=kind):
                template = self.request('GET','/title-import-template',params={'kind':kind})
                self.assertEqual(template.status_code,200,template.text)
                self.assertEqual(template.content.decode('utf-8-sig').splitlines()[0].split(','),headers.split(',')+['商品ID','链接'])
                content = headers+',商品ID,链接,销量,搜索热度,竞争度,搜索曝光,搜索点击,搜索转化,搜索来源,搜索周期\n'+values+',P001,https://example.com/product/P001,nan,不需要的热度,-10,不解析,不解析,不解析,,\n'
                parsed = parse_file(content.encode(),'legacy.csv',kind)
                self.assertEqual(parsed[0]['values'],{**stored,'商品ID':'P001','链接':'https://example.com/product/P001'})
                self.assertEqual(parsed[0]['metrics'],{})

    def test_ignored_metrics_do_not_enter_new_evidence_or_change_historical_data(self):
        self.import_words()
        historical = {'销量':100,'搜索热度':50,'source':'历史来源','period':'2026-10'}
        with self.sessions() as db:
            original = db.scalar(select(TitleEvidence).order_by(TitleEvidence.id))
            original.metrics=historical
            original_id=original.id
            db.commit()
        imported = self.import_words(content='马来语词语,中文释义,适用品类,属性分类,商品ID,链接,销量,搜索热度\ngaun,连衣裙,连衣裙,品类,P001,https://example.com/P001,无效销量,无效热度\n')
        with self.sessions() as db:
            self.assertEqual(db.get(TitleEvidence,original_id).metrics,historical)
            current = db.scalar(select(TitleEvidence).where(TitleEvidence.import_id==imported))
            self.assertEqual(current.metrics,{})
            self.assertEqual((current.source_product,current.source_url),('P001','https://example.com/P001'))

    def test_extraction_cannot_invent_original_words(self):
        row = {'row': 2, 'values': {'原标题': 'Baju-baju lengan panjang', '商品品类': '上衣'}, 'metrics': {}}
        with patch('app.title_import.ask_json', new=AsyncMock(return_value={'language': 'ms', 'keywords': [
            {'phrase': 'cotton', 'translation': '棉', 'tag': '材质'}]})):
            with self.assertRaises(ValueError):
                asyncio.run(extract_row(row, 'titles', 'key'))
        with patch('app.title_import.ask_json', new=AsyncMock(return_value={'language': 'ms', 'keywords': [
            {'phrase': 'Baju-baju', 'translation': '衣服', 'tag': '品类'},
            {'phrase': 'lengan panjang', 'translation': '长袖', 'tag': '结构设计'}]})):
            words = asyncio.run(extract_row(row, 'titles', 'key'))
        self.assertEqual([w['position'] for w in words], [0, 10])
        self.assertEqual(words[0]['phrase'], 'Baju-baju')

    def test_single_generation_cache_and_actual_adoption(self):
        self.enable(); self.profile(); self.import_words()
        model = self.model()
        with patch('app.title_service.ask_json', new=model):
            response = self.request('POST', '/templates/1/generate-draft-title', uid=2,
                json={'method': 'library', 'image_url': 'https://img.example/1.png', 'request_key': 'single-record-001'})
            self.assertEqual(response.status_code, 200, response.text)
            generated = response.json()
            self.assertEqual(generated['title'], TITLE)
            self.assertEqual(generated['language'], 'ms')
            repeat = self.request('POST', '/templates/1/generate-draft-title', uid=2,
                json={'method': 'library', 'image_url': 'https://img.example/1.png', 'request_key': 'single-record-001'})
            self.assertEqual(repeat.json()['generation_id'], generated['generation_id'])
            self.assertEqual(model.await_count, 2)
        self.assertEqual(self.request('GET', '/title-statistics').json()['adopted'], 0)
        response = self.request('POST', '/drafts/from-material-assets', uid=2,
            json={'template_id': 1, 'material_asset_ids': [1], 'title': TITLE, 'generation_id': generated['generation_id']})
        self.assertEqual(response.status_code, 200, response.text)
        draft = response.json()
        self.assertEqual(self.request('GET', '/title-statistics').json()['adopted'], 1)
        for _ in range(2):
            response = self.request('PUT', f"/drafts/{draft['id']}", uid=2, json={'title': TITLE})
            self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.request('GET', '/title-statistics').json()['adopted'], 1)

    def test_template_requirements_reach_malay_generation_and_change_request_identity(self):
        self.enable(); self.profile(); self.import_words()
        self.request('PUT', '/templates/1', json={'title_template': 'Place category first'})
        payload = {'method': 'library', 'image_url': 'https://img.example/1.png', 'request_key': 'template-style-001',
                   'additional_requirements': 'Emphasize visible flowers'}
        model = self.model()
        with patch('app.title_service.ask_json', new=model):
            generated = self.request('POST', '/templates/1/generate-draft-title', uid=2, json=payload)
            self.assertEqual(generated.status_code, 200, generated.text)
            composition = next(call.args[2] for call in model.await_args_list if 'selected' in call.args[2])
            self.assertEqual(composition['template_requirements'], 'Place category first')
            self.assertEqual(composition['style'], 'Emphasize visible flowers')
            self.assertNotIn('template_requirements', model.await_args_list[0].args[2]['manual'])
            self.request('PUT', '/templates/1', json={'title_template': 'Use a natural short title'})
            stale = self.request('POST', '/templates/1/generate-draft-title', uid=2, json=payload)
            self.assertEqual(stale.status_code, 409, stale.text)
            self.assertEqual(model.await_count, 2)
            fresh = self.request('POST', '/templates/1/generate-draft-title', uid=2,
                                 json={**payload, 'request_key': 'template-style-002'})
            self.assertEqual(fresh.status_code, 200, fresh.text)
            self.assertEqual(model.await_args.args[2]['template_requirements'], 'Use a natural short title')
            self.assertEqual(model.await_count, 3)  # 表达要求变化不重复识别同一商品事实。

    def test_title_import_review_exposes_original_and_merge_suggestion(self):
        self.import_words()
        upload = self.upload(content='商品标题,商品分类,商品ID\nGaun Lengan Panjang Corak Bunga,连衣裙,10001\n', kind='titles')
        self.request('POST', f"/title-tasks/{upload['id']}/start", json={})
        with patch('app.title_import.ask_json', new=AsyncMock(return_value={'language': 'ms', 'keywords': [
            {'phrase': 'Lengan Panjang', 'translation': '长袖', 'tag': '结构设计'}]})):
            self.run_worker()
        task = self.request('GET', f"/title-tasks/{upload['id']}").json()
        word = task['result']['words'][0]
        self.assertEqual(word['original_title'], TITLE)
        self.assertEqual(word['existing_code'], 'KW000002')
        self.assertEqual(word['position'], 5)
        self.request('POST', f"/title-imports/{upload['id']}/confirm", json={'selections': [word['index']], 'enable': True})
        keyword = next(w for w in self.request('GET', '/title-keywords').json()['items'] if w['id'] == 2)
        self.assertEqual(keyword['title_frequency'], 1)

    def test_manual_edit_preserves_facts_and_does_not_count_unadopted_ai_title(self):
        self.enable(); self.profile(); self.import_words()
        with patch('app.title_service.ask_json', new=self.model()):
            generated = self.request('POST', '/templates/1/generate-draft-title', uid=2,
                json={'method': 'library', 'image_url': 'https://img.example/1.png'}).json()
        payload = {'template_id': 1, 'material_asset_ids': [1], 'title': TITLE+' Kapas', 'generation_id': generated['generation_id']}
        self.assertEqual(self.request('POST', '/drafts/from-material-assets', uid=2, json=payload).status_code, 400)
        payload['title'] = 'Gaun Corak Bunga Lengan Panjang'
        response = self.request('POST', '/drafts/from-material-assets', uid=2, json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        statistics = self.request('GET', '/title-statistics').json()
        self.assertEqual(statistics['manual'], 1)
        self.assertEqual(statistics['adopted'], 0)
        self.assertTrue(all(w['adopted_count'] == 0 for w in self.request('GET', '/title-keywords').json()['items']))

    def test_changed_template_facts_require_regeneration_before_adoption(self):
        self.enable(); self.profile()
        with patch('app.title_service.ask_json', new=self.model()):
            generated = self.request('POST', '/templates/1/generate-draft-title', uid=2,
                json={'method': 'library', 'image_url': 'https://img.example/1.png'}).json()
        self.profile(attributes={'结构设计': '短袖'})
        response = self.request('POST', '/drafts/from-material-assets', uid=2, json={
            'template_id': 1, 'material_asset_ids': [1], 'title': TITLE, 'generation_id': generated['generation_id']})
        self.assertEqual(response.status_code, 400)
        self.assertIn('属性已改变', response.json()['detail'])

    def test_attempt_limit_persists_and_facts_reject_invented_material(self):
        self.enable(); self.profile(); self.import_words()
        model = self.model('Gaun Lengan Panjang Corak Bunga Kapas')
        self.request('PUT', '/templates/1', json={'title_template': 'Include cotton and write in English'})
        payload = {'method': 'library', 'image_url': 'https://img.example/1.png', 'request_key': 'bad-material-001'}
        with patch('app.title_service.ask_json', new=model):
            self.assertEqual(self.request('POST', '/templates/1/generate-draft-title', uid=2, json=payload).status_code, 502)
            self.assertEqual(model.await_count, 5)
            self.assertEqual(self.request('POST', '/templates/1/generate-draft-title', uid=2, json=payload).status_code, 502)
            self.assertEqual(model.await_count, 5)
        with self.sessions() as db:
            record = db.scalar(select(TitleRecord))
            self.assertEqual(record.attempts, 4)
            self.assertIsNone(record.title)

    def test_visual_material_is_ignored_and_manual_conflicts_block(self):
        self.enable(); self.profile()
        malformed = copy.deepcopy(FACTS)
        malformed['attributes'].append({'tag': '材质', 'value': '棉', 'ms': 'kapas', 'evidence': 'looks cotton', 'source': 'image'})
        with patch('app.title_service.ask_json', new=AsyncMock(side_effect=[malformed, {'title': TITLE}])):
            response = self.request('POST', '/templates/1/generate-draft-title', uid=2, json={'method': 'library', 'image_url': 'https://img.example/1.png'})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertNotIn('材质', [f['tag'] for f in response.json()['facts']['attributes']])
        conflict = {**FACTS, 'conflicts': ['sleeve mismatch']}
        with patch('app.title_service.ask_json', new=AsyncMock(return_value=conflict)):
            response = self.request('POST', '/templates/1/generate-draft-title', uid=2, json={'method': 'library', 'image_url': 'https://img.example/2.png'})
            self.assertEqual(response.status_code, 502)
            self.assertIn('冲突', response.json()['detail'])

    def test_batch_snapshots_cancel_retry_and_scope(self):
        self.enable(); self.profile(); self.import_words()
        self.request('PUT', '/templates/1', json={'title_template': 'Original category-first style'})
        payload = {'template_id': 1, 'request_key': 'batch-example-01', 'groups': [
            {'key': 'one', 'source': {'source_type': 'material', 'id': 1}},
            {'key': 'two', 'source': {'source_type': 'material', 'id': 2}}]}
        response = self.request('POST', '/title-generation-jobs', uid=2, json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        identity = response.json()['id']
        self.assertEqual(self.request('POST', '/title-generation-jobs', uid=2, json=payload).json()['id'], identity)
        self.assertEqual(self.request('POST', '/title-generation-jobs', uid=2, json={**payload, 'platform': 'Shopee'}).json()['id'], identity)
        self.assertEqual(self.request('GET', f'/title-tasks/{identity}', uid=3).status_code, 404)
        self.assertEqual(self.request('GET', f'/title-tasks/{identity}', uid=5).status_code, 404)
        self.assertEqual(self.request('POST', f'/title-tasks/{identity}/cancel', uid=2, json={}).json()['state'], 'cancelled')
        self.assertIsNone(claim_task(self.sessions))
        self.assertEqual(self.request('POST', f'/title-tasks/{identity}/retry', uid=2, json={}).status_code, 200)
        # 修改词库及模板要求不改变已提交任务的快照。
        self.request('PUT', '/templates/1', json={'title_template': 'New writing style'})
        self.request('PATCH', '/title-keywords/1', json={'status': 'disabled'})
        with patch('app.title_service.ask_json', new=self.model()) as model:
            self.run_worker()
        compositions = [call.args[2] for call in model.await_args_list if 'selected' in call.args[2]]
        self.assertTrue(compositions)
        self.assertTrue(all(data['template_requirements'] == 'Original category-first style' for data in compositions))
        task = self.request('GET', f'/title-tasks/{identity}', uid=2).json()
        self.assertEqual(task['state'], 'completed', task)
        self.assertEqual(task['progress']['completed'], 2)
        self.assertTrue(task['result']['items']['two']['warnings'])
        self.assertEqual(len(task['result']['items']), 2)

    def test_expired_task_lease_is_reclaimed_and_cancel_fences_late_writes(self):
        identity = self.upload()['id']
        self.request('POST', f'/title-tasks/{identity}/start', json={})
        claimed_id, owner = claim_task(self.sessions)
        with self.sessions() as db:
            task = db.get(TitleTask, identity)
            task.lease_expires_at = datetime.utcnow() - timedelta(seconds=1)
            db.commit()
        second_id, next_owner = claim_task(self.sessions)
        self.assertEqual(claimed_id, second_id)
        self.assertNotEqual(owner, next_owner)
        asyncio.run(run_task(identity, owner, session_factory=self.sessions))
        self.request('POST', f'/title-tasks/{identity}/cancel', json={})
        asyncio.run(run_task(identity, next_owner, session_factory=self.sessions))
        self.assertEqual(self.request('GET', f'/title-tasks/{identity}').json()['state'], 'cancelled')
        self.assertEqual(self.request('GET', '/title-keywords').json()['total'], 0)

    def test_supply_new_jobs_keep_prompt_engine_and_english_fallback_when_library_is_configured(self):
        from app.models import SupplyJob
        self.request('PUT', '/templates/1', json={'title_template': 'Place product category first'})
        strategy = self.create_strategy()
        old = self.submit(strategy)
        self.request('POST', f'/supply/jobs/{old}/cancel', uid=4, json={})
        self.profile(fallback_title=None, fallback_approved=False)
        category = self.dictionary_items('category')[0]
        self.request('PATCH', f"/title-dictionaries/{category['id']}", json={'enabled': False})
        # 没有马来语兜底、词库分类停用时，货盘仍按原提示词流程提交。
        new = self.submit(key='new-prompt-only')
        with self.sessions() as db:
            for identity in (old, new):
                snap = db.get(SupplyJob, identity).snapshot
                self.assertEqual(snap['title_engine'], 'v1')
                self.assertIsNone(snap['title_profile'])
                self.assertIsNone(snap['title_market_id'])
                self.assertEqual(snap['template']['title_template'], 'Place product category first')
                self.assertEqual(snap['template']['supply_fallback_title'], db.get(ProductTemplate, 1).supply_fallback_title)
        template = next(row for row in self.request('GET', '/supply/templates').json() if row['id']==1)
        self.assertEqual(template['title_engine'], 'v1')
        self.assertNotEqual(template['supply_fallback_title'], '')

    def test_supply_v2_completes_with_malay_fallback_and_records_adoption(self):
        from app.models import SupplyJob
        from app.supply_worker import claim_job, run_job
        self.enable(); self.profile()
        identity = self.submit(key='new-malay-fallback')
        # 模拟上线前已经提交的词库货盘任务；新任务不会再创建此引擎快照。
        with self.sessions() as db:
            job = db.get(SupplyJob, identity)
            market = get_market(db, 1)
            job.snapshot = {**job.snapshot, 'title_engine': 'v2', 'title_market_id': market.id,
                            'title_profile': profile_snapshot(db, market, 1, require_fallback=True)}
            db.commit()
        # 没有公司密钥时不调用模型，使用审核过的马来语兜底。
        upload = {'key': 'supply-exports/company/1/jobs/1/1/'+'a'*32+'.xlsx', 'sha256': 'b'*64, 'size': 10}
        with patch('app.supply_worker.upload_supply_export', return_value=upload), \
             patch('app.title_service.ask_json', new=AsyncMock()) as model:
            claimed = claim_job(self.sessions)
            asyncio.run(run_job(*claimed, session_factory=self.sessions))
            model.assert_not_awaited()
        with self.sessions() as db:
            job = db.get(SupplyJob, identity)
            self.assertEqual(job.state, 'completed', job.error)
            links = job.checkpoint['links']
            self.assertTrue(all(link['title'].startswith('Gaun Wanita') for link in links))
            self.assertTrue(all(link['title_source'] == 'fallback' for link in links))
            self.assertEqual(len({link['title'] for link in links}), 10)
            records = db.scalars(select(TitleRecord)).all()
            self.assertEqual(len(records), 10)
            self.assertTrue(all(r.adopted_at and r.source == 'fallback' for r in records))


class TitlePureTests(unittest.TestCase):
    def test_custom_tag_semantics_are_supplied_to_extraction_with_original_evidence(self):
        row = {'row': 2, 'values': {'原标题': 'Gaun Lengan Panjang Corak Bunga', '商品品类': '连衣裙'}, 'metrics': {}}
        answer = {'language': 'ms', 'keywords': [{'phrase': 'Lengan Panjang', 'translation': '长袖', 'tag': '袖型'}]}
        with patch('app.title_import.ask_json', new=AsyncMock(return_value=answer)) as model:
            words = asyncio.run(extract_row(row, 'titles', 'test-key', allowed_tags=['袖型'], tag_semantics={'袖型': '结构设计'}))
        self.assertEqual(model.await_args.args[2]['tag_semantics'], {'袖型': '结构设计'})
        self.assertEqual((words[0]['tag'], words[0]['position'], words[0]['phrase']), ('袖型', 5, 'Lengan Panjang'))

    def test_similarity_ignores_sequence_case_punctuation_and_order(self):
        self.assertEqual(normalized_title('Gaun, Bunga - L01'), normalized_title('GAUN BUNGA - L02'))
        self.assertTrue(similarity(TITLE, 'Corak Bunga Gaun Lengan Panjang - L02'))
        self.assertFalse(similarity(TITLE, 'Seluar Lelaki Potongan Longgar'))
        for bad in ('Floral Long Sleeve Everyday Dress', '中文马来语商品标题' * 4, TITLE+'\nnew line'):
            with self.assertRaises(ProviderError):
                validate_malay_title(bad)
        self.assertEqual(validate_malay_title(TITLE), TITLE)

    def test_dictionary_migration_backfills_existing_names_without_changing_words_or_tasks(self):
        from app.title_models import TitleDictionaryEntry, TitleTemplateProfile
        from app.title_service import DICTIONARY_TAGS
        from sqlalchemy.schema import CreateTable
        from sqlalchemy.dialects import mysql
        engine = create_engine('sqlite://')
        with engine.begin() as connection:
            config = alembic_config(connection)
            command.upgrade(config, '20261029_50')
            connection.execute(TitleMarket.__table__.insert(), dict(id=7, company_id=7, platform='TikTok', site='MY', language='ms'))
            original_categories = ['连衣裙', 'Ｔｏｐｓ', 'tops']
            connection.execute(TitleKeyword.__table__.insert(), dict(id=8, market_id=7, code='KW000001', phrase='lengan panjang',
                translation='长袖', tag='结构设计', identity_hash='a'*64, categories=original_categories))
            connection.execute(TitleTemplateProfile.__table__.insert(), dict(market_id=7, template_id=9, category='连衣裙', attributes={'材质': '棉'}))
            snapshot = {'profile': {'category': '连衣裙'}, 'groups': [], 'marker': '原快照'}
            connection.execute(TitleTask.__table__.insert(), dict(id=10, market_id=7, user_id=1, kind='batch', state='completed',
                request_key='b'*64, request_hash='c'*64, snapshot=snapshot))
            command.upgrade(config, 'head')
            rows = connection.execute(select(TitleDictionaryEntry.__table__)).mappings().all()
            self.assertEqual(len([r for r in rows if r['kind']=='tag']), len(DICTIONARY_TAGS))
            self.assertTrue(all(r['gender'] is None for r in rows if r['kind']=='category'))
            self.assertEqual({r['name'] for r in rows if r['kind']=='category'}, {'连衣裙', 'Ｔｏｐｓ'})
            self.assertTrue(all(r['built_in'] and r['enabled'] for r in rows if r['kind']=='tag'))
            self.assertEqual(connection.scalar(select(TitleKeyword.__table__.c.categories)), original_categories)
            self.assertEqual(connection.scalar(select(TitleTask.__table__.c.snapshot)), snapshot)
            self.assertEqual(connection.scalar(select(TitleTemplateProfile.__table__.c.attributes)), {'材质': '棉'})
            sql = str(CreateTable(TitleDictionaryEntry.__table__).compile(dialect=mysql.dialect()))
            self.assertIn('uq_title_dictionary_name', sql)
        engine.dispose()

    def test_import_enable_migration_only_updates_pending_words_from_confirmed_imports(self):
        engine = create_engine('sqlite://')
        old_time = datetime(2026, 10, 1)
        with engine.begin() as connection:
            config = alembic_config(connection)
            command.upgrade(config, '20261030_54')
            for identity in (7, 8):
                connection.execute(TitleMarket.__table__.insert(), dict(id=identity, company_id=identity,
                    platform='TikTok', site='MY', language='ms', revision=10, updated_at=old_time))
            for identity, state, market in ((10, 'completed', 7), (11, 'review', 7), (12, 'completed', 8)):
                connection.execute(TitleTask.__table__.insert(), dict(id=identity, market_id=market, user_id=1,
                    kind='import', state=state, request_key=f'import-{identity}', request_hash='a'*64,
                    snapshot={'marker': '保留原始任务'}, result={'marker': '保留入库结果'}))
            for identity in range(1, 8):
                connection.execute(TitleKeyword.__table__.insert(), dict(id=identity, market_id=7,
                    code=f'KW{identity:06d}', phrase=f'baju {identity}', translation='上衣', tag='品类',
                    identity_hash=digest([identity]), categories=['上衣'], status={2: 'disabled', 7: 'enabled'}.get(identity, 'pending'),
                    merged_into_id=1 if identity == 3 else None, updated_at=old_time))
            for identity, task_id, market in ((1, 10, 7), (2, 10, 7), (3, 10, 7), (5, 11, 7), (6, 12, 8), (7, 10, 7)):
                connection.execute(TitleEvidence.__table__.insert(), dict(market_id=market, keyword_id=identity,
                    import_id=task_id, observation_hash=digest(['evidence', identity]), original_title='Baju wanita',
                    original_phrase='baju', source_name='历史导入', source_product=str(identity), period='2026-10-01', source_kind='titles'))
            command.upgrade(config, 'head')
            rows = connection.execute(select(TitleKeyword.__table__)).mappings().all()
            expected = {1: 'enabled', 2: 'disabled', 3: 'pending', 4: 'pending', 5: 'pending', 6: 'pending', 7: 'enabled'}
            self.assertEqual({row['id']: row['status'] for row in rows}, expected)
            self.assertTrue(all(row['updated_at'] == old_time for row in rows if row['id'] != 1))
            self.assertEqual(connection.scalar(select(TitleMarket.__table__.c.revision).where(TitleMarket.__table__.c.id == 7)), 11)
            self.assertEqual(connection.scalar(select(TitleMarket.__table__.c.revision).where(TitleMarket.__table__.c.id == 8)), 10)
            self.assertEqual(connection.scalar(select(func.count(TitleEvidence.id))), 6)
            self.assertTrue(all(result == {'marker': '保留入库结果'} for result in connection.scalars(select(TitleTask.__table__.c.result))))
            command.downgrade(config, '20261030_54')
            command.upgrade(config, 'head')
            self.assertEqual(dict(connection.execute(select(TitleKeyword.__table__.c.id, TitleKeyword.__table__.c.status)).all()), expected)
            self.assertEqual(connection.scalar(select(TitleMarket.__table__.c.revision).where(TitleMarket.__table__.c.id == 7)), 11)
        engine.dispose()

    def test_import_enable_migration_renders_for_mysql_and_mariadb(self):
        from types import SimpleNamespace
        for dialect in ('mysql', 'mariadb'):
            with self.subTest(dialect=dialect):
                output = io.StringIO()
                config = alembic_config()
                config.output_buffer = output
                with patch('app.config.get_settings', return_value=SimpleNamespace(database_url=f'{dialect}+pymysql://test@localhost/test')):
                    command.upgrade(config, '20261030_54:20261030_55', sql=True)
                sql = output.getvalue()
                self.assertIn('UPDATE title_keywords', sql)
                self.assertIn('UPDATE title_markets', sql)
                self.assertIn("title_tasks.state = 'completed'", sql)

    def test_migration_adds_tables_without_rewriting_history(self):
        engine = create_engine('sqlite://')
        with engine.begin() as connection:
            command.upgrade(alembic_config(connection), 'head')
            self.assertTrue({'title_keywords', 'title_evidence', 'title_records', 'title_tasks'}.issubset(inspect(connection).get_table_names()))
            self.assertIn('title_template', {c['name'] for c in inspect(connection).get_columns('product_templates')})
        engine.dispose()

    def test_gender_migration_preserves_existing_words_and_snapshots_and_can_downgrade(self):
        from sqlalchemy import MetaData, Table
        engine = create_engine('sqlite://')
        with engine.begin() as connection:
            config = alembic_config(connection)
            command.upgrade(config, '20261030_51')
            connection.execute(TitleMarket.__table__.insert(), dict(id=7, company_id=7, platform='TikTok', site='MY', language='ms'))
            table = Table('title_dictionary_entries', MetaData(), autoload_with=connection)
            base = dict(market_id=7, description='', enabled=True, updated_at=datetime.utcnow())
            for name in ('品类', '性别'):
                from app.title_dictionary import name_key
                connection.execute(table.insert(), dict(**base, kind='tag', name=name, name_key=name_key(name), semantic_tag=name, built_in=True))
            connection.execute(table.insert(), dict(**base, kind='category', name='女装', name_key=name_key('女装'), built_in=False))
            connection.execute(TitleKeyword.__table__.insert(), dict(id=8, market_id=7, code='KW000001', phrase='wanita',
                translation='女', tag='性别', identity_hash='a'*64, categories=['女装']))
            snapshot = {'profile': {'category': '女装', 'attributes': {'性别': '女'}}, 'marker': '已提交快照'}
            connection.execute(TitleTask.__table__.insert(), dict(id=9, market_id=7, user_id=1, kind='batch', state='queued',
                request_key='b'*64, request_hash='c'*64, snapshot=snapshot))
            command.upgrade(config, 'head')
            current = Table('title_dictionary_entries', MetaData(), autoload_with=connection)
            rows = connection.execute(select(current)).mappings().all()
            self.assertEqual([(r['name'], r['gender']) for r in rows], [('女装', None)])
            self.assertEqual(connection.scalar(select(TitleKeyword.__table__.c.tag)), '性别')
            self.assertEqual(connection.scalar(select(TitleTask.__table__.c.snapshot)), snapshot)
            command.downgrade(config, '20261030_51')
            self.assertNotIn('gender', {column['name'] for column in inspect(connection).get_columns('title_dictionary_entries')})
            restored = Table('title_dictionary_entries', MetaData(), autoload_with=connection)
            self.assertTrue({'品类', '性别'}.issubset(set(connection.scalars(select(restored.c.name).where(restored.c.kind == 'tag')))))
            self.assertEqual(connection.scalar(select(TitleTask.__table__.c.snapshot)), snapshot)
            command.upgrade(config, 'head')
        engine.dispose()

    def test_site_name_migration_preserves_category_gender_references_and_historical_snapshot(self):
        from app.title_dictionary import name_key
        from sqlalchemy import MetaData, Table
        engine = create_engine('sqlite://')
        with engine.begin() as connection:
            config = alembic_config(connection)
            command.upgrade(config, '20261030_52')
            connection.execute(TitleMarket.__table__.insert(), dict(id=7, company_id=7, platform='TikTok', site='MY', language='ms'))
            old = Table('title_dictionary_entries', MetaData(), autoload_with=connection)
            connection.execute(old.insert(), dict(id=8, market_id=7, kind='category', name='上衣', name_key=name_key('上衣'),
                description='原说明', gender='女', built_in=False, enabled=True, updated_at=datetime.utcnow()))
            snapshot = {'profile': {'category': '上衣'}, 'marker': '原快照'}
            connection.execute(TitleTask.__table__.insert(), dict(id=9, market_id=7, user_id=1, kind='batch', state='queued',
                request_key='b'*64, request_hash='c'*64, snapshot=snapshot))
            command.upgrade(config, 'head')
            current = Table('title_dictionary_entries', MetaData(), autoload_with=connection)
            row = connection.execute(select(current)).mappings().one()
            self.assertEqual((row['id'], row['name'], row['gender'], row['description'], row['localized_name']), (8, '上衣', '女', '原说明', None))
            self.assertEqual(connection.scalar(select(TitleTask.__table__.c.snapshot)), snapshot)
            command.downgrade(config, '20261030_52')
            self.assertNotIn('localized_name', {column['name'] for column in inspect(connection).get_columns('title_dictionary_entries')})
            self.assertEqual(connection.scalar(select(TitleTask.__table__.c.snapshot)), snapshot)
            command.upgrade(config, 'head')
        engine.dispose()

    def test_template_configuration_migration_uses_only_confirmed_template_facts_and_preserves_conflicts(self):
        from sqlalchemy import MetaData, Table
        from app.title_dictionary import name_key
        engine=create_engine('sqlite://')
        with engine.begin() as connection:
            config=alembic_config(connection)
            command.upgrade(config,'20261030_53')
            connection.execute(TitleMarket.__table__.insert(),dict(id=7,company_id=7,platform='TikTok',site='MY',language='ms'))
            old_profiles=Table('title_template_profiles',MetaData(),autoload_with=connection)
            source=[{'性别':'女','年龄':'成人','面料':'棉'},
                    {'性别':'儿童','材质':'棉','面料':'涤纶'},
                    {'结构设计':'长袖'}, {'性别':'儿童','年龄':'成人'}]
            for index,attributes in enumerate(source,1):
                connection.execute(old_profiles.insert(),dict(id=index,market_id=7,template_id=index,category='上衣',attributes=attributes,
                    category_names=[],locked_keyword_ids=[],forbidden_keyword_ids=[],fallback_approved=False,updated_at=datetime.utcnow()))
            dictionary=Table('title_dictionary_entries',MetaData(),autoload_with=connection)
            for name,kind,semantic,built_in in [('上衣','category',None,False),('年龄','tag','年龄',True),('面料','tag','面料',True),('材质','tag','材质',True),('布料','tag','面料',False)]:
                connection.execute(dictionary.insert(),dict(market_id=7,name=name,name_key=name_key(name),kind=kind,semantic_tag=semantic,
                    built_in=built_in,enabled=True,description='',gender='女' if kind=='category' else None,updated_at=datetime.utcnow()))
            snapshot={'profile':{'category':'上衣','attributes':{'面料':'棉'}},'marker':'旧任务'}
            connection.execute(TitleTask.__table__.insert(),dict(id=9,market_id=7,user_id=1,kind='batch',state='queued',request_key='b'*64,request_hash='c'*64,snapshot=snapshot))
            command.upgrade(config,'head')
            current=Table('title_template_profiles',MetaData(),autoload_with=connection)
            rows=connection.execute(select(current).order_by(current.c.id)).mappings().all()
            self.assertEqual((rows[0]['gender'],rows[0]['age_group'],rows[0]['attributes']),('女','成人',{'材质':'棉'}))
            self.assertEqual((rows[1]['gender'],rows[1]['age_group'],rows[1]['attributes']),(None,'儿童',{'材质':'棉','面料':'涤纶'}))
            self.assertIsNone(rows[2]['gender'])
            self.assertIsNone(rows[2]['age_group'])
            self.assertEqual(rows[3]['attributes'],{'性别':'儿童','年龄':'成人'})
            self.assertIsNone(rows[3]['age_group'])
            names=set(connection.scalars(select(dictionary.c.name).where(dictionary.c.kind=='tag')))
            self.assertNotIn('面料',names);self.assertNotIn('年龄',names)
            self.assertEqual(connection.scalar(select(dictionary.c.semantic_tag).where(dictionary.c.name=='布料')),'材质')
            self.assertEqual(connection.scalar(select(TitleTask.__table__.c.snapshot)),snapshot)
            command.downgrade(config,'20261030_53')
            columns={col['name'] for col in inspect(connection).get_columns('title_template_profiles')}
            self.assertNotIn('age_group',columns);self.assertNotIn('gender',columns)
            self.assertEqual(connection.scalar(select(TitleTask.__table__.c.snapshot)),snapshot)
        engine.dispose()
