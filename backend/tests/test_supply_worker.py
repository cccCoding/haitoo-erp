"""货盘租约恢复、最终图片检查、兜底与文件完成的行为测试。"""
import asyncio
import copy
import hashlib
from datetime import datetime, timedelta
from io import BytesIO
import random
import unittest
from collections import Counter
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook
from sqlalchemy import select

from app.credentials import encrypt_secret
from app.models import (Company, MaterialAsset, MaterialStyle, ProductTemplate, SupplyJob,
                        SupplyJobMaterial, TiktokCategoryCatalog, User, UserTemplateWhiteImage)
from app.supply_common import (digest, eligible_materials, normalize_title,
                              template_snapshot, template_version, unique_title)
if __package__:
    from .test_supply_catalog import SupplyFixture
else:
    from test_supply_catalog import SupplyFixture


class SupplyStyleSamplingTests(unittest.TestCase):
    def candidates(self, counts):
        items = []
        for style_id, count in counts:
            for _ in range(count):
                identity = len(items) + 1
                items.append({"id": identity, "sku": f"SKU{identity}", "url": f"https://img/{identity}",
                              "style_id": style_id, "name": "同名风格"})
        return items

    def test_complete_styles_are_kept_separate_and_ungrouped_are_not_used(self):
        from app.supply_worker import style_sample
        candidates = self.candidates([(1, 6), (2, 7), (None, 10)])
        original = copy.deepcopy(candidates)
        chosen_styles = set()
        for seed in range(20):
            group = style_sample(candidates, 5, random.Random(seed))
            self.assertEqual(len(group), 5)
            self.assertEqual(len({item['style_id'] for item in group}), 1)
            self.assertIsNotNone(group[0]['style_id'])
            chosen_styles.add(group[0]['style_id'])
        self.assertEqual(chosen_styles, {1, 2})
        self.assertEqual(candidates, original)
        group[0]['sku'] = 'changed'
        self.assertEqual(candidates, original)

    def test_shortfall_uses_largest_styles_before_ungrouped(self):
        from app.supply_worker import style_sample
        for counts, expected in [([(1, 3), (2, 2), (3, 1), (None, 9)], {1: 3, 2: 2}),
                                 ([(1, 2), (2, 1), (None, 9)], {1: 2, 2: 1, None: 2})]:
            for seed in range(10):
                group = style_sample(self.candidates(counts), 5, random.Random(seed))
                self.assertEqual(Counter(item['style_id'] for item in group), expected)

    def test_remaining_groups_are_recounted_after_url_conflicts(self):
        from app.supply_worker import style_sample
        candidates = self.candidates([(1, 3), (2, 2), (3, 2), (None, 6)])
        candidates[3]['url'] = candidates[0]['url']
        for seed in range(10):
            group = style_sample(candidates, 5, random.Random(seed))
            self.assertEqual(Counter(item['style_id'] for item in group), {1: 3, 3: 2})

    def test_raw_group_size_does_not_bypass_deduplication(self):
        from app.supply_worker import style_sample
        for field in ('id', 'sku', 'url'):
            candidates = self.candidates([(1, 6), (2, 5)])
            for item in candidates[:6]:
                item[field] = candidates[0][field]
            group = style_sample(candidates, 5, random.Random(3))
            self.assertEqual({item['style_id'] for item in group}, {2})
            for unique_field in ('id', 'sku', 'url'):
                self.assertEqual(len({item[unique_field] for item in group}), 5)

    def test_conflicts_can_fall_back_to_original_sampler(self):
        from app import supply_worker as worker
        candidates = [
            {'id': 1, 'sku': 'a', 'url': 'x', 'style_id': 1},
            {'id': 2, 'sku': 'b', 'url': 'y', 'style_id': 1},
            {'id': 3, 'sku': 'c', 'url': 'x', 'style_id': 2},
            {'id': 4, 'sku': 'd', 'url': 'y', 'style_id': 2},
            {'id': 5, 'sku': 'a', 'url': 'z', 'style_id': None},
        ]
        with patch.object(worker, 'distinct_sample', wraps=worker.distinct_sample) as sample:
            group = worker.style_sample(candidates, 3, random.Random(4))
        self.assertTrue(any(call.args[0] is candidates for call in sample.call_args_list))
        for field in ('id', 'sku', 'url'):
            self.assertEqual(len({item[field] for item in group}), 3)

    def test_legacy_and_unassigned_snapshots_keep_original_rng_sequence(self):
        from app.supply_worker import distinct_sample, style_sample
        for legacy in (False, True):
            candidates = self.candidates([(None, 12)])
            if legacy:
                for item in candidates:
                    item.pop('style_id')
            expected_rng, actual_rng = random.Random('seed'), random.Random('seed')
            for _ in range(10):
                self.assertEqual(style_sample(candidates, 5, actual_rng),
                                 distinct_sample(candidates, 5, expected_rng))
            self.assertEqual(actual_rng.getstate(), expected_rng.getstate())

    def test_styled_sampling_is_seeded_and_reports_true_shortfalls(self):
        from app.supply_worker import style_sample
        candidates = self.candidates([(1, 3), (2, 2), (3, 2)])
        self.assertEqual(style_sample(candidates, 5, random.Random('seed')),
                         style_sample(candidates, 5, random.Random('seed')))
        self.assertLess(len(style_sample(candidates[:3], 5, random.Random(0))), 5)


class SupplyWorkerTests(SupplyFixture, unittest.TestCase):
    def test_each_generated_link_obeys_snapshot_filter_after_flags_change(self):
        for mode in ("exclude", "include", "only"):
            with self.sessions() as db:
                for asset in db.scalars(select(MaterialAsset).where(MaterialAsset.company_id == 1)):
                    asset.is_infringing = asset.id <= 10
                db.commit()
            identity = self.submit(self.create_strategy(infringement_filter=mode), key=f"filter-{mode}-job")
            # 提交后反转标记，执行仍须使用原候选和标记快照。
            with self.sessions() as db:
                for asset in db.scalars(select(MaterialAsset).where(MaterialAsset.company_id == 1)):
                    asset.is_infringing = not asset.is_infringing
                db.commit()
            self.run_job(identity)
            checkpoint = self.assert_completed(identity)
            for link in checkpoint["links"]:
                self.assertEqual(len(link["materials"]), 5)
                for item in link["materials"]:
                    self.assertEqual(item["is_infringing"], item["id"] <= 10)
                    if mode != "include":
                        self.assertEqual(item["is_infringing"], mode == "only")

    def setUp(self):
        super().setUp()
        from app import supply_worker
        self.worker = supply_worker
        settings = supply_worker.get_settings().model_copy(update={"supply_heartbeat_seconds": 3600})
        self.enterContext(patch("app.supply_worker.get_settings", return_value=settings))
        self.enterContext(patch("app.supply_worker.RETRY_DELAYS", (0, 0, 0)))
        self.title = self.enterContext(patch("app.supply_worker.request_title", new_callable=AsyncMock,
                                            return_value="Printed Long Sleeve Nursing Maxi Dress for Daily Wear"))
        self.image = self.image_mock
        self.contents = []
        def upload(company, job, epoch, filename, content):
            self.contents.append(content)
            return {"key": f"supply-exports/company/{company}/jobs/{job}/{epoch}/" + "a" * 32 + ".xlsx",
                    "sha256": hashlib.sha256(content).hexdigest(), "size": len(content)}
        self.upload = self.enterContext(patch("app.supply_worker.upload_supply_export", side_effect=upload))
        self.delete = self.enterContext(patch("app.supply_worker.delete_supply_export"))
        self.enterContext(patch("app.supply_worker.list_supply_exports", return_value=[]))
        with self.sessions() as db:
            db.get(Company, 1).supply_ai_secret_encrypted = encrypt_secret("test-deepseek-key")
            db.commit()

    def assign_styles(self):
        with self.sessions() as db:
            for identity in (1, 2):
                db.add(UserTemplateWhiteImage(id=identity, company_id=1, template_id=1, user_id=2,
                    name=f'白底图{identity}', image_url=f'https://img.example/white{identity}.png'))
                db.add(MaterialStyle(id=identity, company_id=1, template_id=1, created_by=2,
                    source_white_image_id=identity, name='同名风格'))
            for asset in db.scalars(select(MaterialAsset).where(MaterialAsset.company_id == 1)):
                asset.style_id = 1 if asset.id <= 10 else 2
            db.commit()

    def make_job(self, uid=4, *, candidates=None, link_count=10, key=None):
        with self.sessions() as db:
            template = template_snapshot(db.get(ProductTemplate, 1))
            version = template_version(template)
            parameters = self.strategy_input(link_count=link_count)
            parameters.update(stock=999, cod="Y", title_requirements="")
            source = eligible_materials(db, 1, 1, version) if candidates is None else candidates
            catalog = db.get(TiktokCategoryCatalog, 1)
            job = SupplyJob(company_id=1, user_id=uid, agent_id=1, agent_name="原代理", template_id=1,
                category_catalog_id=1, idempotency_key=key or f"job-{random.random()}", request_hash="a" * 64,
                category_blob=catalog.template_blob,
                snapshot={"strategy": parameters, "template": template, "template_version": version,
                    "catalog_options": copy.deepcopy(catalog.parsed_options), "catalog_type": catalog.template_type,
                    "candidates": copy.deepcopy(source), "seed": "fixed-seed", "user_name": "用户", "internal_usage": False,
                    "expected_rows": 5 * link_count * 2})
            db.add(job)
            db.flush()
            db.add_all(SupplyJobMaterial(company_id=1, job_id=job.id, material_id=item["id"],
                version_hash=item["version_hash"], link_index=-1) for item in source)
            db.commit()
            return job.id

    def run_job(self, identity):
        claimed = self.worker.claim_job(self.sessions, owner="test-worker")
        self.assertIsNotNone(claimed)
        self.assertEqual(claimed[0], identity)
        asyncio.run(self.worker.run_job(*claimed, session_factory=self.sessions))

    def assert_completed(self, identity):
        with self.sessions() as db:
            job = db.get(SupplyJob, identity)
            self.assertEqual(job.state, "completed", job.error)
            self.assertIsNone(job.export_blob)
            self.assertEqual(job.output_size, len(self.contents[-1]))
            self.assertIsNone(job.lease_owner)
            self.assertEqual(db.query(SupplyJobMaterial).filter_by(job_id=identity, successful=True).count(), 50)
            return copy.deepcopy(job.checkpoint)

    def export_rows(self):
        book = load_workbook(BytesIO(self.contents[-1]), read_only=True, data_only=False)
        sheet = book["Template"]
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator)
        fields = {str(value): index for index, value in enumerate(headers) if value}
        rows = [dict((field, row[index]) for field, index in fields.items())
                for row in sheet.iter_rows(min_row=7, max_row=106, values_only=True)]
        book.close()
        return rows

    def test_lease_reclaim_and_fencing_prevent_old_worker_from_committing(self):
        identity = self.make_job()
        first = self.worker.claim_job(self.sessions, owner="worker-a")
        self.assertIsNone(self.worker.claim_job(self.sessions, owner="worker-b"))
        with self.sessions() as db:
            db.get(SupplyJob, identity).lease_expires_at = datetime.utcnow() - timedelta(seconds=1)
            db.commit()
        second = self.worker.claim_job(self.sessions, owner="worker-b")
        self.assertEqual(second[0], identity)
        self.assertGreater(second[1], first[1])
        with self.assertRaises(self.worker.LostLease):
            self.worker.mutate(*first, lambda db, job: setattr(job, "state", "completed"), session_factory=self.sessions)
        self.worker.mutate(*second, lambda db, job: setattr(job, "stage", "new-owner"), session_factory=self.sessions)
        with self.sessions() as db:
            self.assertEqual(db.get(SupplyJob, identity).stage, "new-owner")

    def test_company_concurrency_and_ineligible_accounts(self):
        first, second, third = self.make_job(), self.make_job(uid=7), self.make_job(uid=2)
        self.assertEqual(self.worker.claim_job(self.sessions, owner="worker-a")[0], first)
        self.assertEqual(self.worker.claim_job(self.sessions, owner="worker-b")[0], second)
        self.assertIsNone(self.worker.claim_job(self.sessions, owner="worker-c"))
        with self.sessions() as db:
            db.get(User, 4).supply_expires_at = datetime.utcnow() - timedelta(seconds=1)
            db.get(User, 7).is_active = False
            db.get(SupplyJob, first).export_blob = b"unfinished-export"
            db.get(SupplyJob, second).export_blob = b"unfinished-export"
            db.commit()
        self.worker.cancel_ineligible_jobs(self.sessions)
        with self.sessions() as db:
            self.assertEqual(db.get(SupplyJob, first).state, "cancelled")
            self.assertEqual(db.get(SupplyJob, second).state, "cancelled")
            self.assertIsNone(db.get(SupplyJob, first).export_blob)
            self.assertIsNone(db.get(SupplyJob, second).export_blob)
        self.assertEqual(self.worker.claim_job(self.sessions, owner="worker-c")[0], third)

    def test_success_is_snapshot_based_with_unique_titles_and_deduplicated_images(self):
        identity = self.make_job()
        with self.sessions() as db:
            db.get(ProductTemplate, 1).supply_public = False
            db.get(ProductTemplate, 1).sku_specifications = {"size": {"options": ["XL"]}}
            for material in db.scalars(select(MaterialAsset).where(MaterialAsset.company_id == 1)):
                material.supply_listed = False
            db.delete(db.get(TiktokCategoryCatalog, 1))
            db.commit()
        completed_at = datetime.utcnow()
        with patch("app.supply_worker.now", return_value=completed_at):
            self.run_job(identity)
        state = self.assert_completed(identity)
        with self.sessions() as db:
            self.assertEqual(db.get(SupplyJob, identity).output_expires_at, completed_at + timedelta(days=7))
        self.assertEqual(self.title.call_count, 10)
        urls = [call.args[0] for call in self.image.call_args_list]
        self.assertEqual(len(urls), len(set(urls)))
        self.image.assert_not_called()
        rows = self.export_rows()
        self.assertEqual(len({normalize_title(row["product_name"]) for row in rows}), 10)
        self.assertEqual({row["product_name"] for row in rows}, {
            f"Printed Long Sleeve Nursing Maxi Dress for Daily Wear - L{index:02d}" for index in range(1, 11)})
        self.assertTrue(all(row["seller_sku"].endswith(("-S", "-M")) for row in rows))
        self.assertTrue(all(25 <= len(link["title"]) <= 255 for link in state["links"]))

    def test_export_keeps_each_link_in_snapshot_style_after_style_and_source_changes(self):
        self.assign_styles()
        identity = self.submit()
        with self.sessions() as db:
            original = copy.deepcopy(db.get(SupplyJob, identity).snapshot['candidates'])
            self.assertEqual({item['style_id'] for item in original}, {1, 2})
            for item in original:
                self.assertEqual(item['version_hash'], digest({key: value for key, value in item.items()
                                                              if key != 'version_hash'}))
            # 提交后改风格、删除来源白底图、下架素材，仍然只使用原快照。
            for asset in db.scalars(select(MaterialAsset).where(MaterialAsset.company_id == 1)):
                asset.style_id = 2 if asset.id <= 10 else None
            for white in db.scalars(select(UserTemplateWhiteImage)):
                white.deleted_at = datetime.utcnow()
            db.commit()
            current = eligible_materials(db, 1, 1, original[0]['template_version'])
            self.assertNotEqual(original[0]['version_hash'], current[0]['version_hash'])
            for asset in db.scalars(select(MaterialAsset).where(MaterialAsset.company_id == 1)):
                asset.supply_listed = False
            db.commit()
        self.run_job(identity)
        state = self.assert_completed(identity)
        snapshot_styles = {item['sku']: item['style_id'] for item in original}
        rows = self.export_rows()
        for link in state['links']:
            self.assertEqual(len({item['style_id'] for item in link['materials']}), 1)
            exported = [row for row in rows if row['product_name'] == link['title']]
            self.assertEqual(len(exported), 10)
            self.assertEqual({row['seller_sku'].rsplit('-', 1)[0] for row in exported},
                             {item['sku'] for item in link['materials']})
            self.assertEqual(len({snapshot_styles[row['seller_sku'].rsplit('-', 1)[0]] for row in exported}), 1)
            self.assertEqual({row['property_1_image'] for row in exported},
                             {item['url'] for item in link['materials']})
        self.assertLess(len({item['id'] for link in state['links'] for item in link['materials']}), 50)

    def test_legacy_snapshot_uses_original_selection_even_when_current_assets_have_styles(self):
        self.assign_styles()
        with self.sessions() as db:
            candidates = eligible_materials(db, 1, 1, template_version(template_snapshot(db.get(ProductTemplate, 1))))
        for item in candidates:
            item.pop('style_id')
            item['version_hash'] = digest({key: value for key, value in item.items() if key != 'version_hash'})
        identity = self.make_job(candidates=candidates)
        self.run_job(identity)
        state = self.assert_completed(identity)
        rng = random.Random('fixed-seed')
        for link in state['links']:
            self.assertEqual(link['materials'], self.worker.distinct_sample(candidates, 5, rng))

    def test_existing_link_checkpoint_is_exported_without_resampling(self):
        self.assign_styles()
        identity = self.make_job()
        with self.sessions() as db:
            job = db.get(SupplyJob, identity)
            candidates = job.snapshot['candidates']
            # 已保存的历史混合风格链接也不能被新算法重排。
            group = [candidates[index] for index in (0, 10, 1, 11, 2)]
            links = [{'index': index, 'materials': copy.deepcopy(group), 'title': None, 'replaced': False}
                     for index in range(10)]
            job.checkpoint = {'links': links}
            db.commit()
        with patch.object(self.worker, 'style_sample', side_effect=AssertionError('不应重新选图')):
            self.run_job(identity)
        state = self.assert_completed(identity)
        for link in state['links']:
            self.assertEqual(link['materials'], group)

    def test_title_failure_has_durable_four_call_budget_then_fallback(self):
        identity = self.make_job()
        self.title.side_effect = ValueError("provider failed")
        self.run_job(identity)
        state = self.assert_completed(identity)
        self.assertEqual(self.title.call_count, 40)
        self.assertEqual(set(state["title_attempts"].values()), {4})
        self.assertEqual([link["title"] for link in state["links"]], [
            f"Printed Long Sleeve Nursing Maxi Dress - L{index:02d}" for index in range(1, 11)])
        self.assertEqual(len({normalize_title(link["title"]) for link in state["links"]}), 10)










    def test_upload_retry_reuses_excel_without_regenerating_or_rechecking(self):
        self.assign_styles()
        identity = self.make_job()
        upload = self.upload.side_effect
        self.upload.side_effect = ValueError("storage outage")
        self.run_job(identity)
        first_titles, first_images = self.title.call_count, self.image.call_count
        with self.sessions() as db:
            job = db.get(SupplyJob, identity)
            self.assertEqual(job.state, "retry_wait", job.error)
            self.assertIsNotNone(job.export_blob)
            self.assertEqual(job.upload_attempts, 1)
            first_bytes = bytes(job.export_blob)
            first_links = copy.deepcopy(job.checkpoint['links'])
            job.available_at = datetime.utcnow() - timedelta(seconds=1)
            db.commit()
        self.upload.side_effect = upload
        with patch.object(self.worker, 'style_sample', side_effect=AssertionError('重试不能重新选图')):
            self.run_job(identity)
        state = self.assert_completed(identity)
        self.assertEqual(state['links'], first_links)
        self.assertEqual(self.contents[-1], first_bytes)
        self.assertEqual(self.title.call_count, first_titles)
        self.assertEqual(self.image.call_count, first_images)
        with self.sessions() as db:
            self.assertEqual(db.get(SupplyJob, identity).upload_attempts, 2)

    def test_four_upload_failures_end_task_without_success_statistics(self):
        identity = self.make_job()
        self.upload.side_effect = ValueError("storage outage")
        for attempt in range(4):
            self.run_job(identity)
            with self.sessions() as db:
                job = db.get(SupplyJob, identity)
                self.assertEqual(job.upload_attempts, attempt + 1)
                self.assertEqual(job.state, "failed" if attempt == 3 else "retry_wait")
                self.assertEqual(db.query(SupplyJobMaterial).filter_by(job_id=identity, successful=True).count(), 0)
        self.assertEqual(self.upload.call_count, 4)
        self.assertEqual(self.title.call_count, 10)
        urls = [call.args[0] for call in self.image.call_args_list]
        self.assertEqual(len(urls), len(set(urls)))
        with self.sessions() as db:
            self.assertIsNone(db.get(SupplyJob, identity).export_blob)

    def test_user_cancellation_after_upload_prevents_completion_and_deletes_orphan(self):
        identity = self.make_job()
        upload = self.upload.side_effect
        def cancel_after_upload(company, job_id, epoch, filename, content):
            result = upload(company, job_id, epoch, filename, content)
            response = self.request("POST", f"/supply/jobs/{job_id}/cancel", uid=4)
            self.assertEqual(response.status_code, 200, response.text)
            return result
        self.upload.side_effect = cancel_after_upload
        self.run_job(identity)
        with self.sessions() as db:
            job = db.get(SupplyJob, identity)
            self.assertEqual(job.state, "cancelled")
            self.assertIsNone(job.output_key)
            self.assertEqual(db.query(SupplyJobMaterial).filter_by(job_id=identity, successful=True).count(), 0)
        self.delete.assert_called_once()


    def test_file_cleanup_keeps_task_metadata_and_skips_live_referenced_files(self):
        expired, retained = self.make_job(), self.make_job(uid=7)
        expired_key = f"supply-exports/company/1/jobs/{expired}/1/" + "a" * 32 + ".xlsx"
        retained_key = f"supply-exports/company/1/jobs/{retained}/1/" + "b" * 32 + ".xlsx"
        orphan_key = "supply-exports/company/1/jobs/999/1/" + "c" * 32 + ".xlsx"
        with self.sessions() as db:
            for identity, key, delta in ((expired, expired_key, -1), (retained, retained_key, 1)):
                job = db.get(SupplyJob, identity)
                job.state, job.output_key = "completed", key
                job.output_expires_at = datetime.utcnow() + timedelta(days=delta)
            db.commit()
        with patch("app.supply_worker.list_supply_exports", return_value=[retained_key, orphan_key]):
            self.worker.cleanup_exports(self.sessions)
        self.assertEqual({call.args[0] for call in self.delete.call_args_list}, {expired_key, orphan_key})
        with self.sessions() as db:
            job = db.get(SupplyJob, expired)
            self.assertEqual(job.state, "completed")
            self.assertIsNotNone(job.output_deleted_at)
            self.assertTrue(job.snapshot)
            self.assertIsNone(db.get(SupplyJob, retained).output_deleted_at)

    def test_stable_title_codes_survive_normalization_and_truncation(self):
        titles = [unique_title("Ａ" * 255, "Printed Long Sleeve Nursing Maxi Dress", index) for index in range(100)]
        self.assertEqual(len({normalize_title(title) for title in titles}), 100)
        self.assertTrue(all(len(title) <= 255 for title in titles))
        self.assertTrue(titles[0].endswith(" - L01"))
        self.assertTrue(titles[9].endswith(" - L10"))
        self.assertTrue(titles[98].endswith(" - L99"))
        self.assertTrue(titles[99].endswith(" - L100"))


if __name__ == "__main__":
    unittest.main()
