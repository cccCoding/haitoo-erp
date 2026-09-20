from io import BytesIO
from types import SimpleNamespace
import asyncio
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import main
from app.database import Base
from app.models import ProductDraft, ProductTemplate, Role, TiktokCategoryCatalog, User
from app.schemas import ShopeeDraftExportInput
from app.shopee_export import build_workbook, parse_listing_options


FIELDS = [
    "ps_category", "ps_product_name", "ps_product_description", "ps_sku_parent_short",
    "ps_dangerous_goods", "et_title_variation_integration_no", "et_title_variation_1",
    "et_title_option_for_variation_1", "et_title_image_per_variation", "et_title_variation_2",
    "et_title_option_for_variation_2", "ps_price", "ps_stock", "ps_sku_short",
    "et_title_size_chart", "ps_item_cover_image", "ps_item_image_1", "ps_item_image_2",
    "ps_weight", "ps_length", "ps_width", "ps_height", "channel_id.2000", "channel_id.2001",
]


def shopee_template_bytes() -> bytes:
    workbook = Workbook()
    workbook.active.title = "Guidance"
    sheet = workbook.create_sheet("Template")
    workbook.create_sheet("Upload sample")
    category_sheet = workbook.create_sheet("Pre-order DTS Range")
    workbook.create_sheet("Size chart template list")
    workbook.create_sheet("HiddenShopBrand").sheet_state = "hidden"
    workbook.create_sheet("HiddenTax").sheet_state = "hidden"
    for column, field in enumerate(FIELDS, start=1):
        sheet.cell(1, column).value = f"{field}|0|0"
        sheet.cell(3, column).value = {
            "channel_id.2000": "Doorstep Delivery",
            "channel_id.2001": "Sea Shipping",
        }.get(field, field)
    sheet["A2"] = "basic"
    sheet.freeze_panes = "A5"
    category_sheet["A7"] = "100350-Women Clothes/Tops/Tanks & Camisoles"
    category_sheet["B7"] = "100350"
    category_sheet["C7"] = "3 - 30"
    output = BytesIO()
    workbook.save(output)

    # 复现真实 Shopee 模板中的非标准窗格值，确保解析器会先规范化。
    broken = BytesIO()
    with ZipFile(BytesIO(output.getvalue()), "r") as source, ZipFile(broken, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/"):
                content = content.replace(b'activePane="bottomLeft"', b'activePane="bottom_left"')
            target.writestr(info, content)
    return broken.getvalue()


class ShopeeExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.template_bytes = shopee_template_bytes()
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.session_factory = sessionmaker(bind=self.engine)

    def tearDown(self) -> None:
        self.engine.dispose()

    def test_parser_normalizes_real_shopee_pane_and_reads_options(self) -> None:
        options = parse_listing_options(self.template_bytes, "shopee_basic")
        self.assertEqual(options["template_type"], "shopee_basic")
        self.assertEqual(options["categories"], [{
            "id": "100350", "name": "Women Clothes/Tops/Tanks & Camisoles", "pre_order_dts_range": "3 - 30",
        }])
        self.assertEqual([item["field"] for item in options["shipping_channels"]], ["channel_id.2000", "channel_id.2001"])

    def test_builder_maps_color_size_images_and_shipping(self) -> None:
        template = SimpleNamespace(
            sku_specifications={"size": {"options": ["S", "M"]}},
            package_weight=0.28, package_length=30, package_width=16, package_height=2,
        )
        result = build_workbook(
            template=template, category_id="100350", shipping_channels=["channel_id.2000"], dangerous_goods="No",
            products=[{
                "draft_id": 1, "title": "Unique Shopee product title", "description": "Detailed product description for Shopee.",
                "image_urls": ["https://img.example/cover.jpg", "https://img.example/gallery.jpg"],
                "sku_images": [{"image_url": "https://img.example/color.jpg", "sku": "Y1AA000001"}],
                "size_chart_url": "https://img.example/size.jpg", "price": 12.5, "quantity": 88,
            }], template_bytes=self.template_bytes,
        )
        workbook = load_workbook(BytesIO(result), data_only=False)
        sheet = workbook["Template"]
        fields = {str(sheet.cell(1, column).value).split("|", 1)[0]: column for column in range(1, sheet.max_column + 1)}
        self.assertEqual(sheet.cell(7, fields["et_title_variation_1"]).value, "Color")
        self.assertEqual(sheet.cell(7, fields["et_title_variation_2"]).value, "Size")
        self.assertEqual(sheet.cell(7, fields["et_title_option_for_variation_2"]).value, "S")
        self.assertEqual(sheet.cell(8, fields["ps_sku_short"]).value, "Y1AA000001-M")
        self.assertEqual(sheet.cell(7, fields["channel_id.2000"]).value, "On")
        self.assertEqual(sheet.cell(7, fields["channel_id.2001"]).value, "Off")
        self.assertEqual(workbook.sheetnames[0], "Guidance")

    def test_catalog_upload_uses_shopee_parser(self) -> None:
        with self.session_factory() as db:
            admin = User(id=1, company_id=1, email="admin@example.com", name="Admin", password_hash="x", role=Role.COMPANY_ADMIN)
            db.add(admin)
            db.commit()
            result = asyncio.run(main.create_tiktok_category_catalog(
                "Shopee 女装",
                UploadFile(filename="shopee-basic.xlsx", file=BytesIO(self.template_bytes)),
                "shopee_basic",
                user=admin,
                db=db,
            ))
            self.assertEqual(result["template_type"], "shopee_basic")
            self.assertEqual(result["template_version"], "basic")
            self.assertEqual(result["category_count"], 1)
            self.assertEqual(result["options"]["shipping_channels"][0]["field"], "channel_id.2000")

    def test_route_exports_shopee_catalog(self) -> None:
        options = parse_listing_options(self.template_bytes, "shopee_basic")
        with self.session_factory() as db:
            db.add_all([
                User(id=1, company_id=1, email="operator@example.com", name="Operator", password_hash="x", role=Role.MEMBER),
                ProductTemplate(
                    id=1, company_id=1, name="Y1", product_description="Template product description for Shopee.",
                    size_chart_url="https://img.example/size.jpg", package_weight=0.28,
                    package_length=30, package_width=16, package_height=2,
                    sku_specifications={"size": {"options": ["S", "M"]}},
                ),
                ProductDraft(
                    id=1, company_id=1, template_id=1, title="Unique Shopee product title",
                    product_description="Detailed product description for Shopee.",
                    image_urls=["https://img.example/cover.jpg"],
                    sku_items=[{"image_url": "https://img.example/color.jpg", "sku": "Y1AA000001"}],
                    created_by=1, updated_by=1,
                ),
                TiktokCategoryCatalog(
                    id=1, company_id=1, name="Shopee basic", source_filename="shopee.xlsx",
                    template_type="shopee_basic", template_version="basic", template_blob=self.template_bytes,
                    parsed_options=options, created_by=1,
                ),
            ])
            db.commit()
            response = main.export_drafts_to_shopee(ShopeeDraftExportInput(
                draft_ids=[1], category_catalog_id=1, category_id="100350", default_price=10,
                default_quantity=99, shipping_channels=["channel_id.2000"], dangerous_goods="No",
            ), user=db.get(User, 1), db=db)
            self.assertEqual(db.get(ProductDraft, 1).export_count, 1)
        exported = load_workbook(BytesIO(response.body), data_only=False)["Template"]
        self.assertEqual(exported["A7"].value, "100350")


if __name__ == "__main__":
    unittest.main()
