import asyncio
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
import unittest
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from fastapi import HTTPException, UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import main
from app.database import Base
from app.database import get_db
from app.models import Company, MaterialAsset, ProductLibraryOrder, ProductLibraryOrderProduct, ProductLibraryProduct, ProductLibrarySource, ProductTemplate, Role, User
from app.schemas import ProductLibraryBatchTemplateInput
from app.security import create_access_token
from app.product_library import ALLOWED_PLATFORMS, ALLOWED_SITES


HEADERS = ["店铺名称", "站点", "平台", "订单编号", "数量", "下单时间", "标题", "平台SKU", "产品图片链接", "产品ID"]


def xlsx(rows, headers=HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(name) for name in headers])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def row(*, sku="M06LFSKA9RPG7B6A-M", product_id="1736307142216353080", order="586284685874856962",
        shop="KK Cantik", quantity=1, ordered_at="2026-09-29 10:00:00", title="First title", image="https://example.com/image.jpg"):
    return {"店铺名称": shop, "站点": "马来西亚", "平台": "TikTok", "订单编号": order,
            "数量": quantity, "下单时间": ordered_at, "标题": title, "平台SKU": sku,
            "产品图片链接": image, "产品ID": product_id}


class ProductLibraryTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        with self.sessions() as db:
            db.add_all([
                Company(id=1, name="First"), Company(id=2, name="Second"),
                User(id=1, company_id=1, email="a@example.com", name="A", password_hash="x", role=Role.COMPANY_ADMIN),
                User(id=2, company_id=2, email="b@example.com", name="B", password_hash="x", role=Role.COMPANY_ADMIN),
                ProductTemplate(id=1, company_id=1, name="M06L"),
                ProductTemplate(id=2, company_id=1, name="M05L"),
            ])
            db.commit()

    def tearDown(self):
        self.engine.dispose()

    def import_bytes(self, data, user_id=1):
        with self.sessions() as db:
            return asyncio.run(main.import_product_library(
                file=UploadFile(filename="orders.xlsx", file=BytesIO(data)), user=db.get(User, user_id), db=db,
            ))

    def list_items(self, user_id=1, **filters):
        with self.sessions() as db:
            return main.list_product_library(
                page=1, page_size=20, platform=filters.get("platform"), site=filters.get("site"),
                shop_name=filters.get("shop_name"), sku=filters.get("sku"), template_id=filters.get("template_id"),
                unmatched=filters.get("unmatched", False), source_ids=filters.get("source_ids"),
                user=db.get(User, user_id), db=db,
            )

    def test_sku_search_matches_part_of_sku_and_treats_wildcards_as_text(self):
        self.import_bytes(xlsx([
            row(sku="SKU_100-S", product_id="product-1", order="order-1"),
            row(sku="SKUX100-S", product_id="product-2", order="order-2"),
        ]))

        self.assertEqual(self.list_items(sku="_100")["total"], 1)
        self.assertEqual(self.list_items(sku="  sku_100  ")["items"][0]["sku"], "SKU_100")
        self.assertEqual(self.list_items(sku="missing")["total"], 0)

    def test_import_accepts_files_above_previous_size_limits(self):
        for megabytes, compression in ((12, ZIP_STORED), (51, ZIP_DEFLATED)):
            with self.subTest(megabytes=megabytes):
                output = BytesIO(xlsx([row(order=f"order-{megabytes}")]))
                with ZipFile(output, "a", compression=compression) as archive:
                    with archive.open("extra-data.bin", "w") as entry:
                        for _ in range(megabytes):
                            entry.write(b"x" * 1024 * 1024)
                result = self.import_bytes(output.getvalue())
                self.assertEqual(result["created_orders"], 1)

    def test_import_rejects_files_above_new_size_limits(self):
        with self.assertRaisesRegex(HTTPException, "30MB"):
            self.import_bytes(b"x" * (30 * 1024 * 1024 + 1))
        output = BytesIO(xlsx([row()]))
        with ZipFile(output, "a", compression=ZIP_DEFLATED) as archive:
            with archive.open("extra-data.bin", "w") as entry:
                for _ in range(151):
                    entry.write(b"x" * 1024 * 1024)
        with self.assertRaisesRegex(HTTPException, "150MB"):
            self.import_bytes(output.getvalue())
        with self.sessions() as db:
            for model in (ProductLibrarySource, ProductLibraryProduct, ProductLibraryOrder):
                self.assertEqual(db.scalar(select(func.count()).select_from(model)), 0)

    def test_import_groups_sizes_deduplicates_orders_and_filters_by_shop(self):
        first = row()
        continuation = row(sku="M06LFSKA9RPG7B6A-L")
        for name in ("店铺名称", "站点", "平台", "订单编号", "下单时间"):
            continuation[name] = None
        second_sku = row(sku="M05LFSPE7Y65-L", product_id=first["产品ID"])
        second_order = row(order="586284685874856963", ordered_at="2026-09-30 11:00:00")
        other_shop = row(shop="Other Shop")
        data = xlsx([first, continuation, second_sku, second_order, other_shop], headers=list(reversed(HEADERS)))

        self.assertEqual(self.import_bytes(data), {"created_shops": 2, "created_products": 3, "updated_products": 0, "created_orders": 3})
        self.assertEqual(self.import_bytes(data), {"created_shops": 0, "created_products": 0, "updated_products": 0, "created_orders": 0})
        listing = self.list_items()
        self.assertEqual(listing["total"], 2)
        self.assertEqual({item["sku"]: (item["order_count"], item["sales_quantity"]) for item in listing["items"]}, {
            "M06LFSKA9RPG7B6A": (3, 4), "M05LFSPE7Y65": (1, 1),
        })
        self.assertEqual(self.list_items(shop_name="Other Shop")["total"], 1)
        self.assertEqual(self.list_items(platform="Shopee")["total"], 0)
        self.assertEqual(self.list_items(user_id=2)["total"], 0)
        with self.sessions() as db:
            order = db.scalar(select(ProductLibraryOrder).where(ProductLibraryOrder.order_number == first["订单编号"]).order_by(ProductLibraryOrder.id))
            self.assertEqual(order.ordered_at, datetime(2026, 9, 29, 2, 0, 0))
            self.assertEqual(db.scalar(select(func.count()).select_from(ProductLibraryOrderProduct)), 4)
            company_sku_orders = db.scalar(select(func.count(func.distinct(ProductLibraryOrderProduct.order_id)))
                .join(ProductLibraryProduct, ProductLibraryProduct.id == ProductLibraryOrderProduct.product_id)
                .where(ProductLibraryProduct.company_id == 1, ProductLibraryProduct.sku == "M06LFSKA9RPG7B6A"))
            self.assertEqual(company_sku_orders, 3)

    def test_shop_multi_filter_uses_full_source_identity(self):
        first = row(shop="Same Shop", product_id="p1", order="o1")
        second = row(shop="Same Shop", product_id="p2", order="o2", sku="M05LFSPE7Y65-S")
        second["站点"] = "泰国"
        third = row(shop="Other Shop", product_id="p3", order="o3", sku="OTHERAA123456-S")
        self.import_bytes(xlsx([first, second, third]))
        with self.sessions() as db:
            db.add(User(id=3, company_id=1, email="operator@example.com", name="Operator",
                        password_hash="x", role=Role.MEMBER))
            db.scalar(select(ProductLibrarySource).where(ProductLibrarySource.company_id == 1,
                ProductLibrarySource.site == "马来西亚", ProductLibrarySource.shop_name == "Same Shop")).assigned_user_id = 3
            db.commit()
            user = db.get(User, 1)
            options = main.product_library_filters(user=user, db=db)["shops"]
            self.assertEqual([option["label"] for option in options], [
                "TikTok-泰国-Same Shop-未分配", "TikTok-马来西亚-Other Shop-未分配", "TikTok-马来西亚-Same Shop-Operator",
            ])
            ids = {option["label"]: option["id"] for option in options}
        first_id = ids["TikTok-马来西亚-Same Shop-Operator"]
        second_id = ids["TikTok-泰国-Same Shop-未分配"]
        self.assertEqual(self.list_items(source_ids=[first_id])["total"], 1)
        self.assertEqual(self.list_items(source_ids=[first_id, second_id])["total"], 2)
        self.assertEqual(self.list_items(source_ids=[second_id])["items"][0]["shop_data"][0]["source_id"], second_id)
        self.assertEqual(self.list_items(source_ids=[999999])["total"], 0)

        def override_db():
            with self.sessions() as db:
                yield db

        main.app.dependency_overrides[get_db] = override_db
        try:
            with self.sessions() as db:
                token = create_access_token(db.get(User, 1))
            response = TestClient(main.app).get("/product-library", params=[
                ("source_ids", first_id), ("source_ids", second_id),
            ], headers={"Authorization": f"Bearer {token}"})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["total"], 2)
        finally:
            main.app.dependency_overrides.clear()

    def test_product_list_sorts_by_sales_before_pagination(self):
        records = [
            row(sku=sku, product_id=product_id, order=f"{product_id}-{index}")
            for product_id, sku, count in (
                ("p1", "M06LFSKA9RPG7B6A-S", 3),
                ("p2", "M05LFSPE7Y65-S", 2),
                ("p3", "OTHERAA123456-S", 2),
            ) for index in range(count)
        ]
        self.import_bytes(xlsx(records))
        with self.sessions() as db:
            source_id = db.scalar(select(ProductLibraryProduct.source_id).where(ProductLibraryProduct.external_product_id == "p1"))
            db.add(ProductLibraryProduct(company_id=1, source_id=source_id, external_product_id="zero",
                                         sku="ZEROAA123456", title="Zero", image_url=""))
            db.commit()
            user = db.get(User, 1)
            def listing(page):
                return main.list_product_library(page=page, page_size=2, platform=None, site=None,
                    shop_name=None, sku=None, template_id=None, unmatched=False, user=user, db=db)
            first, second = listing(1), listing(2)
        self.assertEqual(first["total"], 4)
        self.assertEqual([(item["shop_data"][0]["product_id"], item["order_count"]) for item in first["items"]],
                         [("p1", 3), ("p2", 2)])
        self.assertEqual([(item["shop_data"][0]["product_id"], item["order_count"]) for item in second["items"]],
                         [("p3", 2), ("zero", 0)])

    def test_product_list_looks_up_material_creator_with_company_scope(self):
        self.import_bytes(xlsx([
            row(product_id="matched", order="matched-order"),
            row(sku="OTHERAA123456-S", product_id="foreign", order="foreign-order"),
        ]))
        created_at = datetime(2026, 9, 20, 8, 30)
        with self.sessions() as db:
            db.add_all([
                MaterialAsset(company_id=1, url="https://example.com/a.jpg", name="A",
                              sku="M06LFSKA9RPG7B6A", claimed_by=1, created_at=created_at),
                MaterialAsset(company_id=2, url="https://example.com/b.jpg", name="B",
                              sku="OTHERAA123456", claimed_by=2, created_at=created_at),
            ])
            db.commit()
        items = {item["shop_data"][0]["product_id"]: item for item in self.list_items()["items"]}
        self.assertEqual(items["matched"]["material_created_by_name"], "A")
        self.assertEqual(items["matched"]["material_created_at"], main.timestamp_ms(created_at))
        self.assertIsNone(items["foreign"]["material_created_by_name"])
        self.assertIsNone(items["foreign"]["material_created_at"])

    def test_reimport_updates_nonempty_product_details(self):
        self.import_bytes(xlsx([row()]))
        revised = row(title="Updated title", image="https://example.com/new.jpg")
        self.assertEqual(self.import_bytes(xlsx([revised])), {"created_shops": 0, "created_products": 0, "updated_products": 1, "created_orders": 0})
        item = self.list_items()["items"][0]
        self.assertEqual((item["title"], item["image_url"]), ("Updated title", "https://example.com/new.jpg"))
        self.assertEqual(item["order_count"], 1)

    def test_required_product_details_reject_entire_import(self):
        for field in ("标题", "产品图片链接"):
            for value in (None, "", "   "):
                with self.subTest(field=field, value=value):
                    invalid = row(order="second-order")
                    invalid[field] = value
                    with self.assertRaisesRegex(HTTPException, f"第 3 行缺少「{field}」"):
                        self.import_bytes(xlsx([row(), invalid]))
                    with self.sessions() as db:
                        for model in (ProductLibrarySource, ProductLibraryProduct, ProductLibraryOrder, ProductLibraryOrderProduct):
                            self.assertEqual(db.scalar(select(func.count()).select_from(model)), 0)

    def test_all_allowed_sites_and_platforms_create_separate_shops(self):
        rows = []
        for site in ALLOWED_SITES:
            for platform in ALLOWED_PLATFORMS:
                record = row()
                record.update({"站点": site, "平台": platform})
                rows.append(record)
        result = self.import_bytes(xlsx(rows))
        self.assertEqual(result, {"created_shops": 18, "created_products": 18,
                                  "updated_products": 0, "created_orders": 18})
        self.assertEqual(self.import_bytes(xlsx(rows)), {"created_shops": 0, "created_products": 0,
                                                       "updated_products": 0, "created_orders": 0})

    def test_invalid_site_and_platform_return_allowed_values_to_frontend(self):
        def override_db():
            with self.sessions() as db:
                yield db

        main.app.dependency_overrides[get_db] = override_db
        try:
            with self.sessions() as db:
                token = create_access_token(db.get(User, 1))
            client = TestClient(main.app)
            for field, invalid_value, allowed in (("站点", "马来", ALLOWED_SITES),
                                                  ("平台", "tiktok", ALLOWED_PLATFORMS)):
                with self.subTest(field=field):
                    invalid = row(order="second-order")
                    invalid[field] = invalid_value
                    response = client.post("/product-library/import",
                        headers={"Authorization": f"Bearer {token}"},
                        files={"file": ("orders.xlsx", xlsx([row(), invalid]),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
                    self.assertEqual(response.status_code, 400)
                    self.assertEqual(response.json()["detail"],
                        f"第 3 行「{field}」值「{invalid_value}」无效，可填值：{'、'.join(allowed)}")
                    self.assertEqual(self.list_items()["total"], 0)
        finally:
            main.app.dependency_overrides.clear()

    def test_shop_sku_order_details_are_deduplicated_sorted_and_paginated(self):
        start = datetime(2026, 9, 1, 10, 0, 0)
        rows = [row(order=f"order-{index:02d}", ordered_at=(start + timedelta(hours=index)).strftime("%Y-%m-%d %H:%M:%S"))
                for index in range(25)]
        rows += [row(order="order-00", product_id="another-product", ordered_at=rows[0]["下单时间"]),
                 row(order="another-shop-order", shop="Another Shop"),
                 row(order="another-sku-order", sku="M05LFSPE7Y65-L")]
        self.import_bytes(xlsx(rows))
        product = next(item for item in self.list_items()["items"] if item["sku"] == "M06LFSKA9RPG7B6A")
        self.assertEqual((product["order_count"], product["sales_quantity"]), (26, 27))
        with self.sessions() as db:
            user = db.get(User, 1)
            source_id = db.scalar(select(ProductLibrarySource.id).where(ProductLibrarySource.shop_name == "KK Cantik"))
            first = main.list_product_library_orders(product["id"], page=1, source_ids=[source_id], user=user, db=db)
            second = main.list_product_library_orders(product["id"], page=2, source_ids=[source_id], user=user, db=db)
            self.assertEqual((first["total"], first["order_count"], first["page_size"], len(first["items"])), (26, 25, 20, 20))
            self.assertEqual([item["order_number"] for item in first["items"]],
                             [f"order-{index:02d}" for index in range(24, 4, -1)])
            self.assertEqual([item["order_number"] for item in second["items"]],
                             [f"order-{index:02d}" for index in range(4, -1, -1)] + ["order-00"])
            self.assertEqual(first["items"][0]["quantity"], 1)
            self.assertEqual(first["items"][0]["ordered_at"], main.timestamp_ms(start + timedelta(hours=24) - timedelta(hours=8)))
            with self.assertRaisesRegex(HTTPException, "无权查看"):
                main.list_product_library_orders(product["id"], page=1, user=db.get(User, 2), db=db)

    def test_template_filter_and_batch_assignment_accept_all_company_products(self):
        unmatched = row(sku="OTHERAA123456-S", product_id="unmatched-product", order="unmatched-order")
        self.import_bytes(xlsx([row(), unmatched]))
        self.assertEqual(self.list_items(template_id=1)["total"], 1)
        self.assertEqual(self.list_items(unmatched=True)["total"], 1)
        unmatched_id = self.list_items(unmatched=True)["items"][0]["id"]
        matched_id = self.list_items(template_id=1)["items"][0]["id"]

        with self.sessions() as db:
            user = db.get(User, 1)
            with self.assertRaisesRegex(HTTPException, "产品不能重复选择"):
                main.set_product_library_templates(ProductLibraryBatchTemplateInput(
                    product_ids=[unmatched_id, unmatched_id], template_id=2,
                ), user=user, db=db)
            self.assertEqual(main.set_product_library_templates(ProductLibraryBatchTemplateInput(
                product_ids=[unmatched_id, matched_id], template_id=2,
            ), user=user, db=db)["updated"], 2)
            self.assertEqual(db.get(ProductLibraryProduct, matched_id).template_id, 2)
        self.assertEqual(self.list_items(unmatched=True)["total"], 0)
        self.assertEqual({item["id"] for item in self.list_items(template_id=2)["items"]}, {unmatched_id, matched_id})
        self.import_bytes(xlsx([unmatched]))
        self.assertEqual(self.list_items(template_id=2)["total"], 2)

        with self.sessions() as db:
            source_id = db.get(ProductLibraryProduct, unmatched_id).source_id
            orphan = ProductLibraryProduct(company_id=1, source_id=source_id, external_product_id="orphan", sku="ORPHANAA123456", template_id=999, title="Orphan", image_url="")
            db.add(orphan); db.commit()
            orphan_id = orphan.id
        self.assertEqual(self.list_items(unmatched=True)["items"][0]["id"], orphan_id)
        with self.sessions() as db:
            self.assertEqual(main.set_product_library_templates(ProductLibraryBatchTemplateInput(
                product_ids=[orphan_id], template_id=2,
            ), user=db.get(User, 1), db=db)["updated"], 1)
        self.assertEqual(self.list_items(unmatched=True)["total"], 0)

        with self.sessions() as db:
            db.add(ProductLibraryProduct(company_id=2, source_id=999, external_product_id="other", sku="OTHERAA123456", title="Other", image_url=""))
            db.commit()
            other_id = db.scalar(select(ProductLibraryProduct.id).where(ProductLibraryProduct.company_id == 2))
            with self.assertRaisesRegex(HTTPException, "无权设置"):
                main.set_product_library_templates(ProductLibraryBatchTemplateInput(
                    product_ids=[other_id], template_id=2,
                ), user=db.get(User, 1), db=db)

    def test_batch_assignment_enforces_100_item_limit(self):
        with self.sessions() as db:
            db.add_all(ProductLibraryProduct(
                company_id=1, source_id=1, external_product_id=f"product-{index}", sku=f"SKU-{index}", title="", image_url="",
            ) for index in range(101))
            db.commit()
            ids = db.scalars(select(ProductLibraryProduct.id).where(ProductLibraryProduct.company_id == 1).order_by(ProductLibraryProduct.id)).all()
            result = main.set_product_library_templates(ProductLibraryBatchTemplateInput(
                product_ids=ids[:100], template_id=1,
            ), user=db.get(User, 1), db=db)
            self.assertEqual(result["updated"], 100)
            self.assertIsNone(db.get(ProductLibraryProduct, ids[100]).template_id)
            with self.assertRaisesRegex(ValueError, "at most 100"):
                ProductLibraryBatchTemplateInput(product_ids=ids, template_id=1)

    def test_date_is_required_and_conflicting_order_dates_roll_back(self):
        with self.assertRaisesRegex(HTTPException, "下单时间"):
            self.import_bytes(xlsx([row()], headers=[name for name in HEADERS if name != "下单时间"]))
        self.assertEqual(self.list_items()["total"], 0)
        conflict = row(sku="M05LFSPE7Y65-L", ordered_at="2026-09-30 10:00:00")
        with self.assertRaisesRegex(HTTPException, "下单时间"):
            self.import_bytes(xlsx([row(), conflict]))
        self.assertEqual(self.list_items()["total"], 0)

    def test_template_has_required_headers_and_text_columns(self):
        with self.sessions() as db:
            response = main.download_product_library_template(user=db.get(User, 1))
        workbook = load_workbook(BytesIO(response.body))
        sheet = workbook.active
        self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
        self.assertTrue(all(sheet[f"{column}2"].number_format == "@" for column in ("D", "H", "J")))
        self.assertEqual(sheet["F2"].number_format, "yyyy-mm-dd hh:mm:ss")
        self.assertEqual(sheet["A2"].value, None)
        validations = {str(v.sqref): v for v in sheet.data_validations.dataValidation}
        for column, allowed in (("B", ALLOWED_SITES), ("C", ALLOWED_PLATFORMS)):
            validation = validations[f"{column}2:{column}50001"]
            self.assertEqual(validation.formula1, '"' + ','.join(allowed) + '"')
            self.assertTrue(validation.showErrorMessage)
            self.assertEqual(validation.errorStyle, "stop")
        for column in ("G", "I"):
            self.assertIn("必填", sheet[f"{column}1"].comment.text)
        for index, header in enumerate(HEADERS, start=1):
            sheet.cell(2, index).value = row()[header]
        output = BytesIO(); workbook.save(output)
        self.assertEqual(self.import_bytes(output.getvalue())["created_products"], 1)

    def test_only_admin_can_import_and_view_product_rankings(self):
        def override_db():
            with self.sessions() as db:
                yield db

        main.app.dependency_overrides[get_db] = override_db
        try:
            with self.sessions() as db:
                db.get(User, 1).role = Role.COMPANY_ADMIN
                db.get(User, 2).role = Role.COMPANY_ADMIN
                db.add(User(id=3, company_id=1, email="member@example.com", name="Member",
                            password_hash="x", role=Role.MEMBER))
                db.commit()
                token_a = create_access_token(db.get(User, 1))
                token_b = create_access_token(db.get(User, 2))
                token_member = create_access_token(db.get(User, 3))
            client = TestClient(main.app)
            self.assertEqual(client.get("/product-library").status_code, 403)
            member_headers = {"Authorization": f"Bearer {token_member}"}
            self.assertEqual(client.get("/product-library", headers=member_headers).json()["total"], 0)
            self.assertEqual(client.get("/product-library/filters", headers=member_headers).status_code, 200)
            self.assertEqual(client.get("/product-library/import-template", headers=member_headers).status_code, 403)
            self.assertEqual(client.post("/product-library/import", headers=member_headers,
                files={"file": ("orders.xlsx", xlsx([row()]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).status_code, 403)
            result = client.post(
                "/product-library/import", headers={"Authorization": f"Bearer {token_a}"},
                files={"file": ("orders.xlsx", xlsx([row(sku="UNKNOWNAA123456-S")]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
            self.assertEqual(result.status_code, 200, result.text)
            unmatched_list = client.get("/product-library?unmatched=true", headers={"Authorization": f"Bearer {token_a}"}).json()
            self.assertEqual(unmatched_list["total"], 1)
            sku_list = client.get("/product-library", params={"sku": "unknownaa"}, headers={"Authorization": f"Bearer {token_a}"}).json()
            self.assertEqual(sku_list["total"], 1)
            self.assertEqual(client.get("/product-library", params={"sku": "missing"}, headers={"Authorization": f"Bearer {token_a}"}).json()["total"], 0)
            product_id = unmatched_list["items"][0]["id"]
            self.assertEqual(client.get(f"/product-library/{product_id}/orders", headers=member_headers).status_code, 404)
            self.assertEqual(client.post("/product-library/templates/batch", headers=member_headers,
                json={"product_ids": [product_id], "template_id": 1}).status_code, 403)
            assigned = client.post("/product-library/templates/batch", headers={"Authorization": f"Bearer {token_a}"},
                json={"product_ids": [product_id], "template_id": 1})
            self.assertEqual(assigned.status_code, 200, assigned.text)
            self.assertEqual(client.get("/product-library?template_id=1", headers={"Authorization": f"Bearer {token_a}"}).json()["total"], 1)
            self.assertEqual(client.get(f"/product-library/{product_id}/orders", headers={"Authorization": f"Bearer {token_a}"}).json()["total"], 1)
            self.assertEqual(client.get(f"/product-library/{product_id}/orders?page=0", headers={"Authorization": f"Bearer {token_a}"}).status_code, 422)
            self.assertEqual(client.get(f"/product-library/{product_id}/orders", headers={"Authorization": f"Bearer {token_b}"}).status_code, 404)
            self.assertEqual(client.post("/product-library/templates/batch", headers={"Authorization": f"Bearer {token_a}"},
                json={"product_ids": list(range(1, 102)), "template_id": 1}).status_code, 422)
            self.assertEqual(client.get("/product-library", headers={"Authorization": f"Bearer {token_b}"}).json()["total"], 0)
        finally:
            main.app.dependency_overrides.clear()


if __name__ == "__main__":
    unittest.main()
