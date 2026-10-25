"""财务导入、隔离、核算与时间范围的业务回归。"""
from datetime import datetime, timedelta, timezone
from io import BytesIO
from decimal import Decimal
from unittest import TestCase
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from alembic import command
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, inspect, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.finance import router
from app.finance_import import parse_workbook
from app.models import (Company, FinanceBill, FinanceErpItem, FinanceImport, FinanceOrder,
    FinanceShop, FinanceShopSkuCost, FinanceSkuCost, FinanceSkuMapping, FinanceTemplateCost, ProductTemplate, Role, User)
from app.schema_version import alembic_config
from app.security import create_access_token

BILL_HEADERS = ['交易类型', '订单ID/调整单ID', '货币', '结算总金额', '总收入', '结算单 ID', 'SKU ID']
ERP_HEADERS = ['订单编号', '平台SKU', '商品数量', '下单时间']
PAID = '2026-10-20 12:00:00'


def workbook(headers, rows, second_sheet=False):
    w = Workbook()
    w.active.title = '数据'
    w.active.append(headers)
    for row in rows:
        w.active.append(row)
    if second_sheet:
        s = w.create_sheet('另一张订单表')
        s.append(headers)
        for row in rows:
            s.append(row)
    b = BytesIO()
    w.save(b)
    w.close()
    return b.getvalue()


def bill(order='o1', settlement='s1', sku='sku1', amount='20', revenue='30', currency='MYR', kind='订单'):
    return [kind, order, currency, amount, revenue, settlement, sku]


def erp(order='o1', sku='Y1AA001-S', quantity=1, paid=PAID):
    return [order, sku, quantity, paid]


class FinanceTests(TestCase):
    def setUp(self):
        self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        with self.sessions() as db:
            db.add_all([Company(id=1, name='测试甲'), Company(id=2, name='测试乙')])
            for uid, cid, role in [(1, 1, Role.COMPANY_ADMIN), (2, 1, Role.MEMBER), (3, 1, Role.TEAM_LEADER),
                                   (4, 2, Role.COMPANY_ADMIN), (5, None, Role.SUPER_ADMIN)]:
                db.add(User(id=uid, company_id=cid, role=role, name=f'用户{uid}', email=f'u{uid}@example.com', password_hash='unused'))
            db.add_all([ProductTemplate(id=1, company_id=1, name='Y1'),
                ProductTemplate(id=2, company_id=1, name='Y'), ProductTemplate(id=3, company_id=2, name='Z9')])
            db.commit()
            self.tokens = {uid: create_access_token(db.get(User, uid)) for uid in range(1, 6)}
        self.app = FastAPI()
        self.app.include_router(router)
        def database():
            with self.sessions() as db:
                yield db
        self.app.dependency_overrides[get_db] = database
        self.client = TestClient(self.app)
        self.shop = self.request('POST', '/finance/shops', json={'name': '测试店铺'}).json()['id']

    def tearDown(self):
        self.client.close()
        self.engine.dispose()

    def request(self, method, url, uid=1, **kwargs):
        return self.client.request(method, url, headers={'Authorization': f'Bearer {self.tokens[uid]}'}, **kwargs)

    def upload(self, kind, rows, shop=None, uid=1, headers=None, data=None, **kwargs):
        response = self.request('POST', f'/finance/shops/{shop or self.shop}/imports/{kind}', uid=uid,
            files={'file': ('test.xlsx', data or workbook(headers or (BILL_HEADERS if kind == 'platform' else ERP_HEADERS), rows),
                            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}, **kwargs)
        return response

    def query(self, shops=None, **kwargs):
        params = {'shop_ids': shops or [self.shop], 'paid_from': '2026-10-01T00:00:00+08:00',
                  'paid_to': '2026-10-25T23:59:59+08:00', **kwargs}
        return self.request('GET', '/finance/results', params=params)

    def ok(self, response):
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def settings(self, cost='5', rate='1.63'):
        self.ok(self.request('PUT', '/finance/sku-costs', json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['Y1AA001']], 'unit_cost':cost}))
        self.ok(self.request('PUT', '/finance/exchange-rates', json={'items': [{'shop_id': self.shop, 'currency': 'MYR', 'rate': rate}]}))

    def test_only_company_admin_and_all_shop_ids_validated(self):
        for uid in (2, 3, 5):
            self.assertEqual(self.request('GET', '/finance/shops', uid=uid).status_code, 403)
            self.assertEqual(self.upload('platform', [bill()], uid=uid).status_code, 403)
        other = self.ok(self.request('POST', '/finance/shops', uid=4, json={'name': '其他店铺'}))['id']
        self.assertEqual(self.query(shops=[self.shop, other]).status_code, 404)
        self.assertEqual(self.upload('erp', [erp()], shop=other).status_code, 404)
        self.assertEqual(self.request('GET', '/finance/pending-payment-time', params={'shop_ids': [self.shop, other]}).status_code, 404)
        self.assertEqual(self.request('POST', '/finance/shops', json={'name': ' 测试店铺 '}).status_code, 409)
        self.assertEqual(self.request('GET', '/finance/results').status_code, 422)
        self.assertEqual(self.request('GET', '/finance/pending-payment-time').status_code, 422)

    def test_bill_dedup_multiple_settlements_and_erp_replacement(self):
        self.ok(self.upload('platform', [bill(), bill(sku='sku2', amount='4', revenue='8'), bill(kind='调整')] ))
        repeat = self.ok(self.upload('platform', [bill(), bill(sku='sku2', amount='4', revenue='8')]))
        self.assertEqual(repeat['unchanged'], 2)
        self.ok(self.upload('platform', [bill(settlement='s2', amount='-2', revenue='0')]))
        self.ok(self.upload('erp', [erp(quantity=2), erp(None, quantity=1, paid=None), erp(None, sku='Y1AA001-M', paid=None)]))
        self.assertEqual(self.ok(self.upload('erp', [erp(quantity=2), erp(None, quantity=1, paid=None), erp(None, sku='Y1AA001-M', paid=None)]))['unchanged'], 1)
        self.settings()
        result = self.ok(self.query())
        self.assertEqual(result['total']['order_count'], 1)
        self.assertEqual((result['total']['sales'], result['total']['cost'], result['total']['profit']), ('61.94', '20.00', '15.86'))
        self.assertEqual(result['skus']['items'][0]['quantity'], 4)
        self.assertEqual(len(result['orders']['items'][0]['bills']), 3)
        self.assertEqual(len(result['orders']['items'][0]['bills'][0]['source_details']), 1)
        self.ok(self.upload('erp', [erp(quantity=1)]))
        self.assertEqual(self.ok(self.query())['total']['cost'], '5.00')
        with self.sessions() as db:
            self.assertEqual(len(db.scalars(select(FinanceErpItem)).all()), 1)
        self.ok(self.upload('platform', [bill(amount='10')]))
        self.assertEqual(self.ok(self.query())['orders']['items'][0]['settlement'], '19.56')

    def test_shop_rename_keeps_financial_data_and_validates_company_and_type(self):
        self.ok(self.upload('erp',[erp()]))
        self.ok(self.upload('platform',[bill()]))
        self.settings()
        baseline=self.ok(self.query())
        def update(name, shop_type='cross_border', shop=None, uid=1):
            return self.request('PATCH',f'/finance/shops/{shop or self.shop}',uid=uid,
                json={'name':name,'shop_type':shop_type})
        renamed=self.ok(update('  改名后的店铺  '))
        self.assertEqual((renamed['id'],renamed['name'],renamed['has_platform_data']),
                         (self.shop,'改名后的店铺',True))
        result=self.ok(self.query())
        self.assertEqual(result['total'],baseline['total'])
        for row in (result['summaries'][0],result['orders']['items'][0],result['skus']['items'][0]):
            self.assertEqual(row['shop_name'],'改名后的店铺')
        self.assertEqual(self.ok(update('改名后的店铺'))['name'],'改名后的店铺')
        self.assertEqual(update('不能一起保存','local').status_code,409)
        other=self.ok(self.request('POST','/finance/shops',json={'name':'另一店'}))['id']
        self.assertEqual(update('另一店').status_code,409)
        self.assertEqual(update('  ').status_code,400)
        self.assertEqual(update('x'*121).status_code,422)
        self.assertEqual(update('空店改名','local',shop=other).status_code,200)
        self.assertEqual(update('外公司',uid=4).status_code,404)
        for uid in (2,3,5):
            self.assertEqual(update('越权',uid=uid).status_code,403)
        shops=self.ok(self.request('GET','/finance/shops'))
        self.assertEqual(next(s for s in shops if s['id']==self.shop)['name'],'改名后的店铺')
        # 失败的编辑不能变更已有订单、成本或店铺类型。
        self.assertEqual(self.ok(self.query())['total'],baseline['total'])

    def test_local_bills_require_only_five_columns_and_replace_by_order(self):
        self.ok(self.request('PATCH',f'/finance/shops/{self.shop}/type',json={'shop_type':'local'}))
        headers = BILL_HEADERS[:5]
        records = [bill(amount='0.10', revenue='0.20')[:5], bill(amount='0.20', revenue='0.30')[:5],
            bill(order='o2',amount='5',revenue='6')[:5]]
        self.assertEqual(self.ok(self.upload('platform',records,headers=headers))['created'],2)
        self.assertEqual(self.ok(self.upload('platform',records,headers=headers))['unchanged'],2)
        with self.sessions() as db:
            groups = db.scalars(select(FinanceBill).order_by(FinanceBill.id)).all()
            self.assertEqual(len(groups),2)
            self.assertEqual((groups[0].settlement,groups[0].revenue),(Decimal('0.30'),Decimal('0.50')))
            self.assertEqual((groups[0].settlement_id,groups[0].sku_id),('',''))
            self.assertEqual(len(groups[0].source_details),2)
        updated = self.ok(self.upload('platform',[bill(amount='10',revenue='11')[:5]],headers=headers))
        self.assertEqual(updated['updated'],1)
        with self.sessions() as db:
            self.assertEqual([b.settlement for b in db.scalars(select(FinanceBill).order_by(FinanceBill.id))],
                [Decimal('10'),Decimal('5')])
        # 本土店忽略结算单和 SKU ID，即使多余列存在也不能形成另一条账单。
        self.assertEqual(self.ok(self.upload('platform',[bill(sku=None,settlement=None,amount='10',revenue='11')]))['unchanged'],1)
        conflict=self.upload('platform',[bill()[:5],bill(currency='CNY')[:5]],headers=headers)
        self.assertEqual(conflict.status_code,400)
        self.assertEqual(self.request('PATCH',f'/finance/shops/{self.shop}/type',json={'shop_type':'cross_border'}).status_code,409)

    def test_shop_types_and_cross_border_sku_validation_and_deduplication(self):
        shops=self.ok(self.request('GET','/finance/shops'))
        self.assertEqual(shops[0]['shop_type'],'cross_border')
        self.assertFalse(shops[0]['has_platform_data'])
        for headers,record in ((BILL_HEADERS[:-1],bill()[:-1]),(BILL_HEADERS[:-2],bill()[:-2])):
            self.assertEqual(self.upload('platform',[record],headers=headers).status_code,400)
        self.assertEqual(self.upload('platform',[bill(sku=None)]).status_code,400)
        self.assertEqual(self.upload('platform',[bill(),bill(amount='99')]).status_code,400)
        first=self.ok(self.upload('platform',[bill(),bill(),bill(sku='second',amount='4')]))
        self.assertEqual(first['created'],2)
        self.assertEqual(self.ok(self.upload('platform',[bill(sku='second',amount='9')]))['updated'],1)
        with self.sessions() as db:
            self.assertEqual([b.settlement for b in db.scalars(select(FinanceBill).order_by(FinanceBill.id))],[Decimal('20'),Decimal('9')])
        self.assertTrue(self.ok(self.request('GET','/finance/shops'))[0]['has_platform_data'])
        self.assertEqual(self.request('PATCH',f'/finance/shops/{self.shop}/type',json={'shop_type':'local'}).status_code,409)
        other=self.ok(self.request('POST','/finance/shops',uid=4,json={'name':'其他公司','shop_type':'local'}))['id']
        self.assertEqual(self.request('PATCH',f'/finance/shops/{other}/type',json={'shop_type':'cross_border'}).status_code,404)
        self.assertEqual(self.request('PATCH',f'/finance/shops/{self.shop}/type',uid=2,json={'shop_type':'local'}).status_code,403)
        self.assertEqual(self.request('POST','/finance/shops',json={'name':'无效','shop_type':'invalid'}).status_code,422)

    def test_legacy_unidentified_cross_border_group_requires_matching_full_detail(self):
        self.ok(self.upload('erp',[erp()]))
        with self.sessions() as db:
            order=db.scalar(select(FinanceOrder).where(FinanceOrder.order_number=='o1'))
            db.add(FinanceBill(order_id=order.id,settlement_id='s1',sku_id='',currency='MYR',
                settlement=Decimal('24'),revenue=Decimal('38'),source_details=[{'sku_id':None,'currency':'MYR','settlement':'24','revenue':'38'}]))
            db.commit()
        self.assertEqual(self.upload('platform',[bill()]).status_code,409)
        restored=self.ok(self.upload('platform',[bill(),bill(sku='second',amount='4',revenue='8')]))
        self.assertEqual(restored['created'],2)
        with self.sessions() as db:
            self.assertEqual({b.sku_id for b in db.scalars(select(FinanceBill))},{'sku1','second'})

    def test_missing_time_separate_and_backfill_enters_correct_period(self):
        self.ok(self.upload('platform', [bill(), bill(order='old', settlement='s2')]))
        self.assertEqual(self.ok(self.query())['total']['order_count'], 0)
        missing = self.ok(self.request('GET', '/finance/pending-payment-time', params={'shop_ids': [self.shop]}))
        self.assertEqual(missing['total'], 2)
        self.ok(self.upload('erp', [erp(), erp(order='old', paid='2026-08-01 12:00:00')]))
        self.assertEqual(self.ok(self.query())['total']['order_count'], 1)
        self.assertEqual(self.ok(self.request('GET', '/finance/pending-payment-time', params={'shop_ids': [self.shop]}))['total'], 0)

    def test_missing_settings_zero_cost_negative_settlement_and_history_recalculation(self):
        self.ok(self.upload('erp', [erp(), erp(order='erp-only')]))
        self.ok(self.upload('platform', [bill(amount='-2')]))
        result = self.ok(self.query())
        self.assertIsNone(result['total']['sales'])
        self.assertIsNone(result['total']['settlement'])
        self.assertIsNone(result['total']['profit'])
        self.assertEqual(result['total']['erp_only_count'], 1)
        self.assertEqual(result['total']['pending_count'], 1)
        self.assertIn('缺汇率', next(row for row in result['orders']['items'] if row['has_bill'])['statuses'])
        self.settings(cost='0')
        result = self.ok(self.query())
        self.assertEqual(result['total']['cost'], '0.00')
        self.assertEqual(result['total']['profit'], '-3.26')
        self.settings(cost='3', rate='2')
        self.assertEqual(self.ok(self.query())['total']['profit'], '-7.00')
        with self.sessions() as db:
            db.delete(db.get(ProductTemplate, 1)); db.commit()
        result = self.ok(self.query())
        self.assertEqual(result['total']['profit'], '-7.00')
        self.assertEqual(result['total']['pending_count'], 0)
        self.assertNotIn('templates', self.ok(self.request('GET', '/finance/settings')))
        self.assertNotIn('template_id', result['skus']['items'][0])

    def test_order_status_filters_split_calculated_and_pending_before_pagination(self):
        self.ok(self.upload('erp',[
            erp(order='done',sku='READY-S'), erp(order='cost',sku='UNKNOWN-S'),
            erp(order='rate',sku='READY-S'), erp(order='erp-only',sku='READY-S'),
            erp(order='both',sku='UNKNOWN-S'),
        ]))
        self.ok(self.upload('platform',[bill(order='done',currency='CNY'),bill(order='cost',currency='CNY'),
            bill(order='rate'),bill(order='both'),bill(order='no-erp',currency='CNY')]))
        self.ok(self.request('PUT','/finance/sku-costs',json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['READY']],'unit_cost':'0'}))
        with self.sessions() as db:
            db.scalar(select(FinanceOrder).where(FinanceOrder.order_number=='no-erp')).paid_at=datetime(2026,10,20,4)
            db.commit()
        baseline=self.ok(self.query())
        expected={'all':{'done','cost','rate','erp-only','both','no-erp'},'calculated':{'done'},
            'pending':{'cost','rate','erp-only','both','no-erp'}}
        for status, numbers in expected.items():
            with self.subTest(status=status):
                result=self.ok(self.query(order_status=status))
                self.assertEqual({r['order_number'] for r in result['orders']['items']},numbers)
                self.assertEqual(result['orders']['total'],len(numbers))
                self.assertEqual(result['total'],baseline['total'])
                self.assertEqual(result['skus'],baseline['skus'])
        page=self.ok(self.query(order_status='pending',page_size=1,order_page=2))
        self.assertEqual(page['orders']['total'],5)
        self.assertEqual(len(page['orders']['items']),1)
        self.assertNotEqual(page['orders']['items'][0]['order_number'],'done')
        for status in ('invalid','missing_bill','missing_erp','missing_cost','missing_rate'):
            self.assertEqual(self.query(order_status=status).status_code,422)
        self.ok(self.request('PUT','/finance/exchange-rates',json={'items':[{'shop_id':self.shop,'currency':'MYR','rate':'1'}]}))
        self.assertEqual({r['order_number'] for r in self.ok(self.query(order_status='calculated'))['orders']['items']},
            {'done','rate'})
        pending=self.ok(self.query(order_status='pending'))['orders']['items']
        self.assertEqual({r['order_number'] for r in pending},{'cost','erp-only','both','no-erp'})
        self.assertIn('缺后台账单',next(r for r in pending if r['order_number']=='erp-only')['statuses'])

    def test_exact_time_boundaries_90_days_and_timezone(self):
        self.ok(self.upload('platform', [bill()]))
        self.ok(self.upload('erp', [erp()]))
        response = self.query(paid_from='2026-10-20T04:00:00Z', paid_to='2026-10-20T04:00:01Z')
        self.assertEqual(self.ok(response)['total']['order_count'], 1)
        self.assertEqual(self.ok(self.query(paid_from='2026-10-20T11:59:59+08:00', paid_to='2026-10-20T12:00:00+08:00'))['total']['order_count'], 1)
        self.assertEqual(self.ok(self.query(paid_from='2026-10-20T12:00:01+08:00', paid_to='2026-10-20T12:00:02+08:00'))['total']['order_count'], 0)
        start = datetime(2026, 1, 1, tzinfo=timezone.utc)
        self.assertEqual(self.query(paid_from=start.isoformat(), paid_to=(start+timedelta(days=90)).isoformat()).status_code, 200)
        self.assertEqual(self.query(paid_from=start.isoformat(), paid_to=(start+timedelta(days=90,seconds=1)).isoformat()).status_code, 400)
        for start, end in [('2026-10-20T00:00:00', '2026-10-21T00:00:00'), ('2026-10-21T00:00:00Z','2026-10-20T00:00:00Z'), ('2026-10-20T00:00:00Z','2026-10-20T00:00:00Z')]:
            self.assertEqual(self.query(paid_from=start,paid_to=end).status_code, 400)

    def test_multishop_same_number_pagination_and_missing_rates_partial_summary(self):
        other = self.ok(self.request('POST', '/finance/shops', json={'name': '另一店铺'}))['id']
        for shop in (self.shop, other):
            self.ok(self.upload('erp', [erp()], shop=shop))
            self.ok(self.upload('platform', [bill()], shop=shop))
        self.settings()
        result = self.ok(self.query(shops=[self.shop,other],page_size=1))
        self.assertEqual(len(result['orders']['items']), 1)
        self.assertEqual(result['orders']['total'], 2)
        self.assertEqual(result['total']['order_count'], 2)
        self.assertEqual(result['total']['converted_count'], 1)
        self.assertEqual(result['total']['calculated_count'], 1)
        self.assertEqual(result['total']['sales'], '48.90')
        self.assertEqual(result['total']['settlement'], '32.60')
        self.assertEqual([s['settlement'] for s in result['summaries']],['32.60',None])
        self.assertEqual(result['total']['profit'], '27.60')
        self.assertEqual(len(result['skus']['items']), 2)
        self.assertEqual([r['order_count'] for r in result['skus']['items']], [1,1])

    def test_sku_page_sizes_are_independent_and_preserve_summary(self):
        self.ok(self.upload('erp', [erp(order=f'o{i//50}', sku=f'SKU{i:04d}-S') for i in range(1005)]))
        self.ok(self.upload('platform', [bill(order=f'o{i}', currency='CNY') for i in range(21)]))
        baseline = self.ok(self.query())
        self.assertEqual(len(baseline['skus']['items']), 20)
        for size in (20, 50, 100, 500, 1000):
            with self.subTest(size=size):
                result = self.ok(self.query(sku_page_size=size))
                self.assertEqual(len(result['skus']['items']), size)
                self.assertEqual(result['skus']['total'], 1005)
                self.assertEqual(len(result['orders']['items']), 20)
                self.assertEqual(result['total'], baseline['total'])
        second = self.ok(self.query(sku_page_size=500, sku_page=2))
        self.assertEqual(second['skus']['items'][0]['base_sku'], 'SKU0500')
        last = self.ok(self.query(sku_page_size=1000, sku_page=2))
        self.assertEqual([s['base_sku'] for s in last['skus']['items']], [f'SKU{i:04d}' for i in range(1000,1005)])
        for size in (0, 1001):
            self.assertEqual(self.query(sku_page_size=size).status_code, 422)

    def test_order_and_pending_page_sizes_and_boundaries(self):
        with self.sessions() as db:
            for paid in (datetime(2026,10,20,4), None):
                orders = [FinanceOrder(company_id=1, shop_id=self.shop,
                    order_number=f'{"paid" if paid else "pending"}{i:04d}', paid_at=paid) for i in range(1005)]
                db.add_all(orders)
                db.flush()
                db.add_all([FinanceBill(order_id=o.id, settlement_id='s', sku_id='s', currency='CNY',
                    settlement=Decimal('20'), revenue=Decimal('30')) for o in orders])
            db.commit()
        baseline = self.ok(self.query())
        self.assertEqual(len(baseline['orders']['items']), 20)
        def pending(**params):
            return self.request('GET', '/finance/pending-payment-time', params={'shop_ids':[self.shop], **params})
        self.assertEqual(len(self.ok(pending())['items']), 20)
        for size in (20, 50, 100, 500, 1000):
            with self.subTest(size=size):
                result = self.ok(self.query(page_size=size))
                self.assertEqual(len(result['orders']['items']), size)
                self.assertEqual(result['total'], baseline['total'])
                self.assertEqual(len(self.ok(pending(page_size=size))['items']), size)
        for size, page, expected in ((500,2,500),(1000,2,5)):
            self.assertEqual(len(self.ok(self.query(page_size=size,order_page=page))['orders']['items']),expected)
            self.assertEqual(len(self.ok(pending(page_size=size,page=page))['items']),expected)
        for size in (0,1001):
            self.assertEqual(self.query(page_size=size).status_code,422)
            self.assertEqual(pending(page_size=size).status_code,422)

    def test_case_sensitive_prefix_and_selected_sku_costs(self):
        self.ok(self.upload('erp', [erp(sku='Y1AA001-S'), erp(order='o2',sku='y1aa001-S'),
            erp(order='o3',sku='AY1AA001-S'), erp(order='o4',sku='Y1AA002-M')]))
        self.assertEqual([s['base_sku'] for s in self.ok(self.query(sku_search='Y1'))['skus']['items']], ['Y1AA001','Y1AA002'])
        self.assertEqual([s['base_sku'] for s in self.ok(self.query(sku_search='y1'))['skus']['items']], ['y1aa001'])
        self.assertEqual(self.ok(self.query(sku_search='AA'))['skus']['total'],0)
        self.ok(self.request('PUT','/finance/sku-costs',json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['Y1AA001','Y1AA002']],'unit_cost':'18.25'}))
        costs={s['base_sku']:s['unit_cost'] for s in self.ok(self.query())['skus']['items']}
        self.assertEqual(costs,{'Y1AA001':'18.25','Y1AA002':'18.25','y1aa001':None,'AY1AA001':None})
        self.ok(self.request('PUT','/finance/sku-costs',json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['y1aa001']],'unit_cost':'0'}))
        costs={s['base_sku']:s['unit_cost'] for s in self.ok(self.query())['skus']['items']}
        self.assertEqual((costs['Y1AA001'],costs['y1aa001']),('18.25','0.00'))
        self.assertEqual(self.request('PUT','/finance/sku-costs',json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['Y1AA001','Y1aa001']],'unit_cost':'99'}).status_code,400)
        self.assertEqual(self.request('PUT','/finance/sku-costs',uid=4,json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['Y1AA001']],'unit_cost':'99'}).status_code,404)
        self.assertEqual(self.request('POST','/finance/sku-mappings',json={'base_skus':['Y1AA001'],'mode':'auto'}).status_code,404)
        self.assertEqual(self.request('PUT','/finance/template-costs',json={'items':[]}).status_code,404)

    def test_shop_sku_costs_are_independent_and_batches_validate_all_pairs(self):
        other = self.ok(self.request('POST', '/finance/shops', json={'name':'第二店'}))['id']
        foreign = self.ok(self.request('POST', '/finance/shops', uid=4, json={'name':'外公司'}))['id']
        for shop in (self.shop, other):
            self.ok(self.upload('erp', [erp(quantity=2), erp(order='o2', quantity=3)], shop=shop))
            self.ok(self.upload('platform', [bill(currency='CNY'), bill(order='o2', currency='CNY')], shop=shop))
        def save(items, value):
            return self.request('PUT','/finance/sku-costs',json={'items':items,'unit_cost':value})
        first = {'shop_id':self.shop,'base_sku':'Y1AA001'}
        second = {'shop_id':other,'base_sku':'Y1AA001'}
        self.ok(save([first], '5'))
        result = self.ok(self.query(shops=[self.shop,other],sku_page_size=1))
        self.assertEqual(result['skus']['total'],2)
        self.assertEqual(result['total']['cost'],'25.00')
        self.assertEqual(result['total']['calculated_count'],2)
        self.assertEqual(self.ok(self.query(shops=[other],sku_cost_status='unset'))['skus']['total'],1)
        self.ok(save([second], '8'))
        result = self.ok(self.query(shops=[self.shop,other]))
        self.assertEqual({r['shop_id']:r['total_cost'] for r in result['skus']['items']},
                         {self.shop:'25.00',other:'40.00'})
        self.assertEqual((result['total']['cost'],result['total']['profit']),('65.00','15.00'))
        # 越权和不存在的配对均导致整批失败，不能先保存合法项。
        self.assertEqual(save([first,{'shop_id':foreign,'base_sku':'Y1AA001'}], '99').status_code,404)
        self.assertEqual(save([first,{'shop_id':other,'base_sku':'unknown'}], '99').status_code,400)
        self.assertEqual(self.ok(self.query())['total']['cost'],'25.00')
        self.assertEqual(self.ok(save([first, second, first], '0'))['updated'],2)
        self.assertEqual(self.ok(self.query(shops=[self.shop,other]))['total']['cost'],'0.00')
        # 新店导入相同 SKU 不会继承旧成本，旧公司级留档也不能作为 fallback。
        new = self.ok(self.request('POST','/finance/shops',json={'name':'新店'}))['id']
        with self.sessions() as db:
            db.add(FinanceSkuCost(company_id=1,base_sku='Y1AA001',unit_cost=Decimal('99')))
            db.commit()
        self.ok(self.upload('erp',[erp()],shop=new))
        self.assertIsNone(self.ok(self.query(shops=[new]))['skus']['items'][0]['unit_cost'])
        self.assertEqual(self.request('PUT','/finance/sku-costs',json={'base_skus':['Y1AA001'],'unit_cost':'5'}).status_code,422)

    def test_same_order_case_variants_and_missing_cost_block_profit(self):
        self.ok(self.upload('erp',[erp(sku='Y1AA001-S',quantity=2),erp(None,sku='y1aa001-S',paid=None)]))
        self.ok(self.upload('platform',[bill()]))
        self.settings()
        result=self.ok(self.query())
        self.assertIsNone(result['orders']['items'][0]['cost'])
        self.assertEqual(result['total']['calculated_count'],0)
        self.assertEqual(result['orders']['items'][0]['statuses'],['缺成本'])
        self.assertEqual({s['base_sku']:s['total_cost'] for s in result['skus']['items']},
            {'Y1AA001':'10.00','y1aa001':None})
        configured = self.ok(self.query(sku_cost_status='set'))
        self.assertEqual([s['base_sku'] for s in configured['skus']['items']],['Y1AA001'])
        self.assertEqual(configured['total'],result['total'])
        self.assertEqual([s['base_sku'] for s in self.ok(self.query(sku_cost_status='unset'))['skus']['items']],['y1aa001'])
        self.assertEqual(self.ok(self.query(sku_cost_status='unset',sku_search='Y1'))['skus']['total'],0)
        self.assertEqual(self.query(sku_cost_status='invalid').status_code,422)
        self.ok(self.request('PUT','/finance/sku-costs',json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['y1aa001']],'unit_cost':'0'}))
        result=self.ok(self.query())
        self.assertEqual((result['total']['cost'],result['total']['profit']),('10.00','22.60'))
        self.assertEqual({s['base_sku']:s['total_cost'] for s in result['skus']['items']},
            {'Y1AA001':'10.00','y1aa001':'0.00'})
        self.assertEqual(self.ok(self.query(sku_cost_status='set'))['skus']['total'],2)
        self.assertEqual(self.ok(self.query(sku_cost_status='unset'))['skus']['total'],0)
        self.ok(self.request('PUT','/finance/sku-costs',json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['Y1AA001']],'unit_cost':'9'}))
        self.assertEqual(self.ok(self.query())['total']['cost'],'18.00')

    def test_multicurrency_round_at_order_level_and_cost_once(self):
        self.ok(self.upload('erp',[erp()]))
        self.ok(self.upload('platform',[bill(amount='.105',revenue='.105',currency='CNY'),
            bill(sku='sku2',amount='.105',revenue='.105',currency='CNY'), bill(settlement='s2',sku='sku3',amount='1',revenue='2')]))
        self.settings(cost='0.01',rate='1.63')
        row=self.ok(self.query())['orders']['items'][0]
        self.assertEqual((row['settlement'],row['sales'],row['cost'],row['profit']),('1.84','3.47','0.01','1.83'))
        self.assertEqual(self.ok(self.query())['total']['settlement'],'1.84')

    def test_settlement_summary_counts_converted_orders_without_cost_and_rounds_each_order(self):
        numbers=['a','b','negative','zero','no-rate','erp-only']
        self.ok(self.upload('erp',[erp(order=n) for n in numbers]))
        self.ok(self.upload('platform',[
            bill(order='a',currency='CNY',amount='.105'),bill(order='b',currency='CNY',amount='.105'),
            bill(order='negative',currency='CNY',amount='-.10'),bill(order='zero',currency='CNY',amount='0'),
            bill(order='no-rate',amount='100')]))
        result=self.ok(self.query(page_size=1))
        self.assertEqual(result['total']['settlement'],'0.12')
        self.assertEqual(result['summaries'][0]['settlement'],'0.12')
        self.assertEqual(result['total']['converted_count'],4)
        self.assertEqual(result['total']['calculated_count'],0)
        self.assertIsNone(result['total']['cost'])
        self.assertIsNone(result['total']['profit'])
        self.assertEqual(self.ok(self.query(page_size=100))['total'],result['total'])
        empty=self.ok(self.query(paid_from='2026-10-21T00:00:00Z',paid_to='2026-10-22T00:00:00Z'))
        self.assertEqual(empty['total']['settlement'],'0.00')

    def test_invalid_import_rolls_back_and_numeric_identifiers_rejected(self):
        for rows in ([erp(),erp(order='o2',paid=None)], [erp(quantity=0)], [erp(quantity=-1)],
                     [erp(),erp(None,paid='2026-10-21 12:00:00')], [erp(order=586775103209767953)],
                     [erp(paid='=NOW()')], [erp(None,paid=None)]):
            self.assertEqual(self.upload('erp',rows).status_code,400)
        self.assertEqual(self.upload('platform',[bill(amount='NaN')]).status_code,400)
        self.assertEqual(self.upload('platform',[bill(),bill(currency='CNY')]).status_code,400)
        with self.sessions() as db:
            self.assertEqual(len(db.scalars(select(FinanceOrder)).all()),0)
            self.assertEqual(len(db.scalars(select(FinanceImport)).all()),0)
        for data in (b'garbage', b'PKbad'):
            self.assertEqual(self.upload('erp',[],data=data).status_code,400)
        with patch('app.finance_import.MAX_ROWS',1):
            self.assertEqual(self.upload('erp',[erp(),erp(order='o2')]).status_code,400)
        self.assertEqual(self.upload('erp',[erp()],headers=['订单编号','平台SKU','商品数量']).status_code,400)

    def test_required_columns_allow_aliases_extra_columns_and_reordered_headers(self):
        headers=['备注','总收入','订单 ID','结算单 ID','货币','SKU ID','交易类型','结算总金额']
        self.ok(self.upload('platform', [['无关内容','30','o1','s1','CNY','sku1','订单','20']],headers=headers))
        headers=['商品数量','下单时间','平台 sku','其他列','订单号']
        self.ok(self.upload('erp', [[2,PAID,'Y1AA001-S','=unused()','o1']],headers=headers))
        result=self.ok(self.query())
        self.assertEqual(result['total']['sales'],'30.00')
        self.assertEqual(result['skus']['items'][0]['quantity'],2)
        response=self.upload('erp', [[1,'Y1AA001-S','o2']],headers=['商品数量','平台 SKU','订单号'])
        self.assertEqual(response.status_code,400)
        self.assertIn('缺少：下单时间',response.json()['detail'])
        response=self.upload('erp', [['o2','o2','Y1AA001-S',1,PAID]],
            headers=['订单编号','订单 ID','平台 SKU','商品数量','下单时间'])
        self.assertEqual(response.status_code,400)
        self.assertIn('重复必需列：订单编号',response.json()['detail'])

    def test_erp_requires_order_time_and_never_substitutes_payment_time(self):
        headers=['订单编号','平台SKU','商品数量','付款时间','下单时间']
        self.ok(self.upload('erp',[
            ['o1','Y1AA001-S',1,'2026-10-20 12:00:00','2026-10-19 09:00:00'],
            [None,'Y1AA001-M',2,None,None],
        ],headers=headers))
        self.ok(self.upload('platform',[bill(currency='CNY')]))
        ordered=self.ok(self.query(paid_from='2026-10-19T09:00:00+08:00',paid_to='2026-10-19T09:00:01+08:00'))
        self.assertEqual(ordered['total']['order_count'],1)
        self.assertEqual(datetime.fromisoformat(ordered['orders']['items'][0]['paid_at'].replace('Z','+00:00')),
            datetime(2026,10,19,1,tzinfo=timezone.utc))
        self.assertEqual(ordered['skus']['items'][0]['quantity'],3)
        paid=self.ok(self.query(paid_from='2026-10-20T12:00:00+08:00',paid_to='2026-10-20T12:00:01+08:00'))
        self.assertEqual(paid['orders']['total'],0)
        missing=self.upload('erp',[erp(order='o2')],headers=['订单编号','平台SKU','商品数量','付款时间'])
        self.assertEqual(missing.status_code,400)
        self.assertIn('缺少：下单时间',missing.json()['detail'])
        with self.sessions() as db:
            self.assertIsNone(db.scalar(select(FinanceOrder).where(FinanceOrder.order_number=='o2')))

    def test_header_after_title_row(self):
        w=Workbook()
        w.active.append(['导出说明，仅供参考'])
        w.active.append(ERP_HEADERS)
        w.active.append(erp())
        output=BytesIO(); w.save(output); w.close()
        parsed=parse_workbook(output.getvalue(),'erp')
        self.assertEqual(parsed['orders']['o1']['items'],{'Y1AA001-S':('Y1AA001',1)})

    def test_multiple_sheets_and_wrong_dimensions(self):
        data=workbook(BILL_HEADERS,[bill()],second_sheet=True)
        choice=self.upload('platform',[],data=data)
        self.assertEqual(choice.status_code,409)
        self.assertEqual(choice.json()['detail']['sheets'],['数据','另一张订单表'])
        response=self.request('POST',f'/finance/shops/{self.shop}/imports/platform',data={'sheet_name':'另一张订单表'},files={'file':('test.xlsx',data)})
        self.ok(response)
        # 构造 dimension 声称只有表头、实际存在多条订单的导出文件。
        data=workbook(ERP_HEADERS,[erp(),erp(order='o2')]); output=BytesIO()
        with ZipFile(BytesIO(data)) as source, ZipFile(output,'w',ZIP_DEFLATED) as target:
            for name in source.namelist():
                content=source.read(name)
                if name=='xl/worksheets/sheet1.xml':
                    content=content.replace(b'ref="A1:D3"',b'ref="A1:D1"')
                target.writestr(name,content)
        self.assertEqual(len(parse_workbook(output.getvalue(),'erp')['orders']),2)

    def test_setting_validation_atomicity(self):
        self.ok(self.upload('erp',[erp()]))
        for value in ('-1','NaN','Infinity','0.001','',' ','10000000000000000'):
            response=self.request('PUT','/finance/sku-costs',json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['Y1AA001']],'unit_cost':value})
            self.assertEqual(response.status_code,400)
            self.assertIsNone(self.ok(self.query())['skus']['items'][0]['unit_cost'])
        self.assertEqual(self.request('PUT','/finance/sku-costs',json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['Y1AA001','unknown']],'unit_cost':'2'}).status_code,400)
        self.assertIsNone(self.ok(self.query())['skus']['items'][0]['unit_cost'])
        self.assertEqual(self.request('PUT','/finance/sku-costs',json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in []],'unit_cost':'2'}).status_code,422)
        for uid in (2,3,5):
            self.assertEqual(self.request('PUT','/finance/sku-costs',uid=uid,json={'items':[{'shop_id':self.shop,'base_sku':sku} for sku in ['Y1AA001']],'unit_cost':'2'}).status_code,403)
        for currency, value in [('MYR','0'),('MYR','-2'),('CNY','2')]:
            self.assertEqual(self.request('PUT','/finance/exchange-rates',json={'items':[{'shop_id':self.shop,'currency':currency,'rate':value}]}).status_code,400)


class FinanceMigrationTests(TestCase):
    def test_shop_cost_migration_preserves_prices_and_does_not_invent_missing_costs(self):
        engine = create_engine('sqlite://')
        with engine.begin() as connection:
            config = alembic_config(connection)
            command.upgrade(config,'20261026_48')
            connection.execute(FinanceShop.__table__.insert(),[
                {'id':1,'company_id':1,'name':'一店'}, {'id':2,'company_id':1,'name':'二店'},
                {'id':3,'company_id':2,'name':'其他公司'}, {'id':4,'company_id':1,'name':'空店'}])
            connection.execute(FinanceOrder.__table__.insert(),[
                {'id':shop,'company_id':2 if shop==3 else 1,'shop_id':shop,'order_number':'same'} for shop in (1,2,3)])
            connection.execute(FinanceErpItem.__table__.insert(),[
                {'order_id':shop,'platform_sku':sku+'-S','base_sku':sku,'quantity':2}
                for shop in (1,2,3) for sku in ('Abc','abc','missing')])
            connection.execute(FinanceSkuCost.__table__.insert(),[
                {'company_id':1,'base_sku':'Abc','unit_cost':Decimal('18.25')},
                {'company_id':1,'base_sku':'abc','unit_cost':Decimal('0')},
                {'company_id':1,'base_sku':'unused','unit_cost':Decimal('99')},
                {'company_id':2,'base_sku':'Abc','unit_cost':Decimal('3')}])
            command.upgrade(config,'head')
            costs = {(shop,sku):cost for shop,sku,cost in connection.execute(select(
                FinanceShopSkuCost.shop_id,FinanceShopSkuCost.base_sku,FinanceShopSkuCost.unit_cost))}
            self.assertEqual(costs,{(1,'Abc'):Decimal('18.25'),(2,'Abc'):Decimal('18.25'),
                (1,'abc'):Decimal('0'),(2,'abc'):Decimal('0'),(3,'Abc'):Decimal('3')})
            connection.execute(FinanceShopSkuCost.__table__.update().where(FinanceShopSkuCost.shop_id==1).values(unit_cost=Decimal('7')))
            self.assertEqual(connection.scalar(select(FinanceShopSkuCost.unit_cost).where(
                FinanceShopSkuCost.shop_id==2,FinanceShopSkuCost.base_sku=='Abc')),Decimal('18.25'))
            command.downgrade(config,'20261026_48')
            self.assertNotIn('finance_shop_sku_costs',inspect(connection).get_table_names())
            self.assertEqual(len(connection.execute(select(FinanceSkuCost.id)).all()),4)
        engine.dispose()

    def test_shop_types_restore_cross_border_skus_and_keep_unidentified_history(self):
        engine=create_engine('sqlite://')
        with engine.begin() as connection:
            config=alembic_config(connection)
            command.upgrade(config,'20261026_47')
            connection.execute(FinanceShop.__table__.insert(),[{'id':1,'company_id':1,'name':'原有店铺'}])
            connection.execute(FinanceBill.__table__.insert(),[
                {'id':1,'order_id':1,'settlement_id':'s1','sku_id':'','currency':'MYR','settlement':Decimal('3'),'revenue':Decimal('6'),
                 'source_details':[{'sku_id':'A','currency':'MYR','settlement':'1','revenue':'2'},
                                   {'sku_id':'a','currency':'MYR','settlement':'2','revenue':'4'}]},
                {'id':2,'order_id':2,'settlement_id':'s2','sku_id':'','currency':'MYR','settlement':Decimal('5'),'revenue':Decimal('7'),
                 'source_details':[{'sku_id':None,'currency':'MYR','settlement':'5','revenue':'7'}]},
            ])
            command.upgrade(config,'head')
            self.assertEqual(connection.scalar(select(FinanceShop.shop_type)),'cross_border')
            restored=connection.execute(select(FinanceBill.__table__).order_by(FinanceBill.id)).mappings().all()
            self.assertEqual({r['sku_id'] for r in restored},{'A','a',''})
            self.assertEqual(sum(r['settlement'] for r in restored),Decimal('8'))
            self.assertEqual(sum(len(r['source_details']) for r in restored),3)
            self.assertEqual(next(r for r in restored if r['sku_id']=='a')['source_details'][0]['settlement_id'],'s1')
            constraint=next(c for c in inspect(connection).get_unique_constraints('finance_bills') if c['name']=='uq_finance_bill')
            self.assertEqual(constraint['column_names'],['order_id','settlement_id','sku_id'])
            command.downgrade(config,'20261026_47')
            rows=connection.execute(select(FinanceBill.__table__)).mappings().all()
            self.assertEqual(len(rows),2)
            self.assertEqual(sum(r['settlement'] for r in rows),Decimal('8'))
            self.assertNotIn('shop_type',{c['name'] for c in inspect(connection).get_columns('finance_shops')})
        engine.dispose()

    def test_settlement_upgrade_merges_legacy_skus_preserves_totals_and_can_restore(self):
        engine=create_engine('sqlite://')
        with engine.begin() as connection:
            config=alembic_config(connection)
            command.upgrade(config,'20261025_46')
            records=[
                {'id':1,'order_id':1,'settlement_id':'s1','sku_id':'A','currency':'MYR','settlement':Decimal('1.111111'),'revenue':Decimal('2.222222')},
                {'id':2,'order_id':1,'settlement_id':'s1','sku_id':'B','currency':'MYR','settlement':Decimal('-0.111111'),'revenue':Decimal('0.777778')},
                {'id':3,'order_id':1,'settlement_id':'s2','sku_id':'A','currency':'CNY','settlement':Decimal('0'),'revenue':Decimal('0')},
                {'id':4,'order_id':2,'settlement_id':'s1','sku_id':'A','currency':'MYR','settlement':Decimal('7'),'revenue':Decimal('8')},
                {'id':5,'order_id':3,'settlement_id':'s1','sku_id':'A','currency':'MYR','settlement':Decimal('-3'),'revenue':Decimal('0')},
            ]
            connection.execute(FinanceBill.__table__.insert(),records)
            columns=[FinanceBill.order_id,FinanceBill.settlement_id,FinanceBill.sku_id,
                FinanceBill.currency,FinanceBill.settlement,FinanceBill.revenue]
            before=connection.execute(select(*columns).order_by(FinanceBill.id)).all()
            command.upgrade(config,'20261026_47')
            merged=connection.execute(select(FinanceBill.__table__).order_by(FinanceBill.id)).mappings().all()
            self.assertEqual(len(merged),4)
            self.assertEqual((merged[0]['settlement'],merged[0]['revenue']),(Decimal('1'),Decimal('3')))
            self.assertEqual([r['sku_id'] for r in merged[0]['source_details']],['A','B'])
            self.assertEqual(sum(r['settlement'] for r in merged),sum(r[4] for r in before))
            constraint=next(c for c in inspect(connection).get_unique_constraints('finance_bills') if c['name']=='uq_finance_bill')
            self.assertEqual(constraint['column_names'],['order_id','settlement_id'])
            with self.assertRaises(IntegrityError):
                with connection.begin_nested():
                    connection.execute(FinanceBill.__table__.insert().values(order_id=1,settlement_id='s1',
                        sku_id='different',currency='MYR',settlement=Decimal('1'),revenue=Decimal('2')))
            command.downgrade(config,'20261025_46')
            restored=connection.execute(select(*columns)).all()
            self.assertEqual(sorted(restored),sorted(before))
            self.assertNotIn('source_details',{c['name'] for c in inspect(connection).get_columns('finance_bills')})
        engine.dispose()

    def test_settlement_upgrade_rejects_mixed_currencies_before_changing_legacy_rows(self):
        engine=create_engine('sqlite://')
        with engine.begin() as connection:
            config=alembic_config(connection)
            command.upgrade(config,'20261025_46')
            connection.execute(FinanceBill.__table__.insert(),[
                {'order_id':1,'settlement_id':'s1','sku_id':'A','currency':'MYR','settlement':Decimal('1'),'revenue':Decimal('2')},
                {'order_id':1,'settlement_id':'s1','sku_id':'B','currency':'CNY','settlement':Decimal('3'),'revenue':Decimal('4')},
            ])
            with self.assertRaisesRegex(RuntimeError,'不同货币'):
                command.upgrade(config,'20261026_47')
            self.assertEqual(len(connection.execute(select(FinanceBill.id)).all()),2)
            self.assertNotIn('source_details',{c['name'] for c in inspect(connection).get_columns('finance_bills')})
        engine.dispose()

    def test_upgrade_downgrade_preserves_existing_company(self):
        engine=create_engine('sqlite://')
        with engine.begin() as connection:
            config=alembic_config(connection)
            command.upgrade(config,'20261025_44')
            connection.execute(Company.__table__.insert().values(id=1,name='已有公司',is_active=True,created_at=datetime.utcnow()))
            command.upgrade(config,'head')
            finance_tables={t.name for t in Base.metadata.tables.values() if t.name.startswith('finance_')}
            self.assertEqual(len(finance_tables),10)
            self.assertTrue(finance_tables.issubset(inspect(connection).get_table_names()))
            command.downgrade(config,'20261025_44')
            self.assertFalse(finance_tables.intersection(inspect(connection).get_table_names()))
            self.assertEqual(connection.scalar(select(Company.name)), '已有公司')
        engine.dispose()

    def test_existing_costs_backfill_raw_case_and_remain_independent(self):
        engine=create_engine('sqlite://')
        with engine.begin() as connection:
            config=alembic_config(connection)
            command.upgrade(config,'20261025_45')
            connection.execute(Company.__table__.insert(),[{'id':1,'name':'甲'},{'id':2,'name':'乙'}])
            connection.execute(ProductTemplate.__table__.insert(),[{'id':1,'company_id':1,'name':'OLD'},{'id':2,'company_id':2,'name':'OTHER'}])
            connection.execute(FinanceShop.__table__.insert(),[{'id':1,'company_id':1,'name':'甲店'},{'id':2,'company_id':2,'name':'乙店'}])
            connection.execute(FinanceOrder.__table__.insert(),[{'id':1,'company_id':1,'shop_id':1,'order_number':'o1'},{'id':2,'company_id':2,'shop_id':2,'order_number':'o2'}])
            connection.execute(FinanceErpItem.__table__.insert(),[{'id':1,'order_id':1,'platform_sku':'Abc-S','base_sku':'ABC','quantity':2},{'id':2,'order_id':2,'platform_sku':'ABC-M','base_sku':'ABC','quantity':1}])
            connection.execute(FinanceSkuMapping.__table__.insert(),[{'company_id':1,'base_sku':'ABC','template_id':1},{'company_id':2,'base_sku':'ABC','template_id':2},{'company_id':1,'base_sku':'DELETED','template_id':999}])
            connection.execute(FinanceTemplateCost.__table__.insert(),[{'company_id':1,'template_id':1,'unit_cost':Decimal('18.25')},{'company_id':2,'template_id':2,'unit_cost':Decimal('0')},{'company_id':1,'template_id':999,'unit_cost':Decimal('9')}])
            command.upgrade(config,'head')
            costs={(company,sku):value for company,sku,value in connection.execute(select(FinanceSkuCost.company_id,FinanceSkuCost.base_sku,FinanceSkuCost.unit_cost))}
            self.assertEqual(costs,{(1,'ABC'):Decimal('18.25'),(1,'Abc'):Decimal('18.25'),(2,'ABC'):Decimal('0')})
            self.assertEqual(connection.scalar(select(FinanceErpItem.base_sku).where(FinanceErpItem.id==1)),'Abc')
            connection.execute(ProductTemplate.__table__.delete())
            self.assertEqual(connection.scalar(select(FinanceSkuCost.unit_cost).where(FinanceSkuCost.company_id==1,FinanceSkuCost.base_sku=='Abc')),Decimal('18.25'))
            command.downgrade(config,'20261025_45')
            self.assertNotIn('finance_sku_costs',inspect(connection).get_table_names())
            self.assertEqual(connection.scalar(select(FinanceTemplateCost.unit_cost).where(FinanceTemplateCost.company_id==1,FinanceTemplateCost.template_id==1)),Decimal('18.25'))
        engine.dispose()
